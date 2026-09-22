# -*- coding: utf-8 -*-
"""生成泰隆银行 2026 年四季度参访与高端小局执行 PPT。

相对可执行版 V2 的调整：
- 单场费用硬顶 5000 元
- 不再回补已过窗口的 8 月、9 月上旬场次
- 10 月起步，11–12 月加密参访与高端小局
- 目标改为存量客户满意、新客户获取
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK = RGBColor(0x0A, 0x1E, 0x36)
INK_2 = RGBColor(0x12, 0x2E, 0x4A)
TEAL = RGBColor(0x15, 0x4A, 0x5C)
AMBER = RGBColor(0xC9, 0x8A, 0x2A)
AMBER_LT = RGBColor(0xE8, 0xC9, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xE6, 0xEE, 0xF3)
MUTED = RGBColor(0x96, 0xAE, 0xBB)
OK = RGBColor(0x4E, 0xA3, 0x78)

FONT = "微软雅黑"
SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]
PAGE = 0
FOOTER = "泰隆银行客户参访与高端小局 · 2026年10—12月 · 单场不超过5000元"


def _font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def slide(bg=INK):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid()
    r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line:
        sp.line.color.rgb = line
        sp.line.width = Pt(1.1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def rrect(s, x, y, w, h, color, line=None, adj=0.08):
    sp = rect(s, x, y, w, h, color, MSO_SHAPE.ROUNDED_RECTANGLE, line)
    try:
        sp.adjustments[0] = adj
    except Exception:
        pass
    return sp


def txt(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(1)
    tf.margin_bottom = Pt(1)
    for i, spec in enumerate(lines):
        text, size, color, bold = spec[0], spec[1], spec[2], spec[3]
        space_after = spec[4] if len(spec) > 4 else 3
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        _font(run, size, color, bold)
    return tb


def header(s, title, subtitle=None, tag=None):
    rect(s, 0, 0, Inches(0.12), SH, AMBER)
    txt(s, Inches(0.42), Inches(0.2), Inches(10.4), Inches(0.46),
        [(title, 24, WHITE, True)])
    if subtitle:
        txt(s, Inches(0.44), Inches(0.66), Inches(10.3), Inches(0.32),
            [(subtitle, 13, AMBER_LT, False)])
    if tag:
        chip = rrect(s, Inches(11.15), Inches(0.26), Inches(1.85), Inches(0.38), TEAL, AMBER)
        tf = chip.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = tag
        _font(r, 11, AMBER_LT, True)
    rect(s, Inches(0.42), Inches(1.05), Inches(12.45), Pt(1.25), AMBER)


def footer(s, n):
    txt(s, Inches(0.42), Inches(7.12), Inches(10.6), Inches(0.26),
        [(FOOTER, 10, MUTED, False)])
    txt(s, Inches(11.7), Inches(7.12), Inches(1.2), Inches(0.26),
        [(str(n), 10, MUTED, False)], align=PP_ALIGN.RIGHT)


def new(title=None, subtitle=None, tag=None):
    global PAGE
    PAGE += 1
    s = slide()
    if title:
        header(s, title, subtitle, tag)
    footer(s, PAGE)
    return s


def card(s, x, y, w, h, title, bullets, accent=AMBER, body_size=13):
    rrect(s, x, y, w, h, TEAL)
    rect(s, x, y, Inches(0.08), h, accent)
    txt(s, x + Inches(0.2), y + Inches(0.12), w - Inches(0.32), Inches(0.36),
        [(title, 15, AMBER_LT, True)])
    lines = [("• " + b, body_size, OFF, False, 4) for b in bullets]
    txt(s, x + Inches(0.2), y + Inches(0.52), w - Inches(0.34), h - Inches(0.64), lines)


def draw_table(s, x0, y0, col_widths_in, row_h_in, headers, rows,
               header_bg=AMBER, header_fg=INK, zebra=(TEAL, INK_2),
               first_col_color=AMBER_LT, body_size=12, header_size=12,
               gap_in=0.05):
    widths = [Inches(w) for w in col_widths_in]
    xs = [x0]
    for w in widths[:-1]:
        xs.append(xs[-1] + w)
    gap = Inches(gap_in)
    row_h = Inches(row_h_in)
    for j, h in enumerate(headers):
        rrect(s, xs[j], y0, widths[j] - gap, row_h, header_bg)
        txt(s, xs[j] + Inches(0.05), y0, widths[j] - gap - Inches(0.08), row_h,
            [(h, header_size, header_fg, True)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i, row in enumerate(rows):
        y = y0 + row_h + Inches(0.045) + i * (row_h + Inches(0.04))
        bg = zebra[0] if i % 2 == 0 else zebra[1]
        for j, val in enumerate(row):
            rrect(s, xs[j], y, widths[j] - gap, row_h, bg)
            color = first_col_color if j == 0 else WHITE
            txt(s, xs[j] + Inches(0.06), y, widths[j] - gap - Inches(0.1), row_h,
                [(str(val), body_size, color, j == 0)],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# 1 封面
# ---------------------------------------------------------------------------
s = new()
rect(s, Inches(8.55), 0, Inches(4.8), SH, INK_2)
rect(s, 0, 0, SW, Pt(5), AMBER)
rect(s, Inches(0.62), Inches(1.45), Inches(0.12), Inches(1.35), AMBER)
txt(s, Inches(0.95), Inches(1.35), Inches(7.2), Inches(0.36),
    [("2026年10—12月执行方案", 15, AMBER_LT, False)])
txt(s, Inches(0.92), Inches(1.85), Inches(7.3), Inches(2.3),
    [("让老客户愿意再来", 32, WHITE, True, 6),
     ("让新客户进得来", 32, WHITE, True, 10),
     ("泰隆银行小微与科技企业", 16, OFF, False, 2),
     ("参访与高端小局", 16, OFF, False)])
txt(s, Inches(0.95), Inches(4.85), Inches(7.2), Inches(1.4),
    [("单场费用不超过 5000 元", 16, AMBER_LT, True, 6),
     ("复旦住房政策研究中心 · 杨浦区科企联 · 上海市科企联", 12, MUTED, False, 3),
     ("拟合作方：浙江泰隆商业银行上海地区机构（以最终确认为准）", 12, MUTED, False)])
txt(s, Inches(8.9), Inches(1.7), Inches(4.1), Inches(4.6),
    [("这一版只抓三件事", 16, AMBER_LT, True, 16),
     ("7 场", 28, WHITE, True, 0),
     ("10/22 至 12/17", 14, OFF, False, 14),
     ("≤5000 元", 28, WHITE, True, 0),
     ("每场决算硬顶", 14, OFF, False, 14),
     ("12 月最密", 28, WHITE, True, 0),
     ("4 场参访 + 3 场小局", 14, OFF, False)])

# ---------------------------------------------------------------------------
# 2 三句话
# ---------------------------------------------------------------------------
s = new("这一版怎么执行", "今天是 9 月 22 日。8 月和 9 月上旬不再回补，从国庆后直接排能落地的场次。", "原则")
principles = [
    ("01  钱有硬顶", "每一场决算不超过 5000 元。报价超了就减项，不追加、不挪到下一场冲销。"),
    ("02  年底走出去", "10 月 1 场热身，11 月 3 场，12 月 3 场。参访和高端小局为主，不上大型年会。"),
    ("03  人被接住", "客户经理陪自己的客户。老客户当天被记住；新客户 5 个工作日内有人联系。"),
]
for i, (t, d) in enumerate(principles):
    y = Inches(1.4) + i * Inches(1.75)
    rrect(s, Inches(0.45), y, Inches(12.4), Inches(1.55), TEAL)
    rect(s, Inches(0.45), y, Inches(0.1), Inches(1.55), AMBER)
    txt(s, Inches(0.85), y + Inches(0.28), Inches(11.6), Inches(0.45),
        [(t, 22, AMBER_LT, True)])
    txt(s, Inches(0.85), y + Inches(0.8), Inches(11.6), Inches(0.5),
        [(d, 16, OFF, False)])

# ---------------------------------------------------------------------------
# 3 相对上一版
# ---------------------------------------------------------------------------
s = new("相对上一版改了什么", "保留合规边界和跟进台账，把花钱方式和评价方式改过来。", "调整")
draw_table(
    s, Inches(0.4), Inches(1.35),
    [2.3, 5.0, 5.1], 0.72,
    ["项目", "上一版", "这一版"],
    [
        ["时间", "8—12 月，每月 1 场大场", "10/22—12/17，共 7 场小场"],
        ["费用", "单场可达数万元，全年约 20 万", "单场硬顶 5000 元，全年不超过 3.5 万"],
        ["形态", "旗舰分享会 + 年底年会", "参访 4 场，高端小局 3 场"],
        ["目标", "服务转化数量为主", "老客户满意，新客户有人跟"],
        ["人力", "单场约 17—20 人", "单场 8—12 人，客户经理陪同"],
        ["礼品", "未单列约束", "默认不采购伴手礼"],
    ],
    body_size=14, header_size=14,
)

# ---------------------------------------------------------------------------
# 4 双目标
# ---------------------------------------------------------------------------
s = new("两项目标，都要能核对", "不开成发布会。客户带得走的是一次被好好对待的体验，和新的一次联系。", "目标")
card(s, Inches(0.4), Inches(1.35), Inches(6.2), Inches(5.4), "存量客户：开心、愿意再来", [
    "客户经理全程坐在自己客户旁边",
    "半天结束，13:30—16:30，不占全天",
    "离场前复述一句：今天为您记下的事",
    "次日问一句感受，不满意 2 日内电话回",
    "7 场累计存量客户不少于 60 人次",
    "每场满意度不低于 90%",
], OK, 15)
card(s, Inches(6.8), Inches(1.35), Inches(6.1), Inches(5.4), "新客户：进得来、有人接", [
    "客户经理提名，不公开海投",
    "老客户可带 1 家同行或上下游",
    "新客户离场时知道对接人是谁",
    "5 个工作日内完成第一次联系",
    "7 场累计新客到场不少于 25 家",
    "不承诺额度、利率和放款",
], AMBER, 15)

# ---------------------------------------------------------------------------
# 5 节奏
# ---------------------------------------------------------------------------
s = new("年底把场次做密", "同一位客户默认不连续出席相邻两场，除非客户自己提出。", "节奏")
blocks = [
    ("10 月", "1 场", "热身", "10/22 参访\n先把线路和陪同跑顺"),
    ("11 月", "3 场", "加密", "1 场参访 + 2 场小局\n老客户与新客户分开照顾"),
    ("12 月", "3 场", "做满", "2 场参访 + 1 场答谢\n把感谢和次年预约做完"),
]
for i, (m, n, k, d) in enumerate(blocks):
    x = Inches(0.45) + i * Inches(4.2)
    rrect(s, x, Inches(1.4), Inches(4.0), Inches(3.55), TEAL)
    txt(s, x + Inches(0.25), Inches(1.6), Inches(3.5), Inches(0.4),
        [(m, 14, AMBER_LT, True)])
    txt(s, x + Inches(0.25), Inches(2.1), Inches(3.5), Inches(0.7),
        [(n, 36, WHITE, True)])
    txt(s, x + Inches(0.25), Inches(2.9), Inches(3.5), Inches(0.4),
        [(k, 18, AMBER_LT, True)])
    txt(s, x + Inches(0.25), Inches(3.45), Inches(3.5), Inches(1.2),
        [(d, 14, OFF, False)])
txt(s, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.6),
    [("排不开时的取舍：先保住参访。若客户经理陪同冲突，11 月 26 日新客小局并入 12 月 17 日，12 月三场不减。", 15, OFF, False, 8),
     ("被访企业临时取消时：改到银行贵宾会议室做同主题小局，照常接待已确认的客户，费用仍不超过 5000 元。", 15, AMBER_LT, False)])

# ---------------------------------------------------------------------------
# 6 日历
# ---------------------------------------------------------------------------
s = new("七场日历", "全部安排在周四下午。场地用银行贵宾室或企业现场，不租酒店。", "日历")
draw_table(
    s, Inches(0.35), Inches(1.28),
    [0.7, 1.7, 1.5, 3.5, 1.6, 3.4], 0.62,
    ["序", "日期", "类型", "主题", "客户数", "这场为了什么"],
    [
        ["1", "10/22", "参访", "智能制造 / 具身智能", "16", "让老客户先感到被重视"],
        ["2", "11/5", "高端小局", "重点客户私享问诊", "12", "把经营问题听清楚"],
        ["3", "11/19", "参访", "链主企业或 AI 应用现场", "16", "老客户带新客户"],
        ["4", "11/26", "高端小局", "意向新客闭门局", "12", "新客户建立信任"],
        ["5", "12/3", "参访", "跨境窗口或园区", "16", "把出海变成看得见"],
        ["6", "12/10", "参访", "客户点名的标杆企业", "18", "年底答谢"],
        ["7", "12/17", "高端小局", "年度收官与次年预约", "16", "感谢，并约下一次"],
    ],
    body_size=13, header_size=13,
)

# ---------------------------------------------------------------------------
# 7 名额
# ---------------------------------------------------------------------------
s = new("名额怎么分", "先保老客户的体验，再留出新客户的位子。工作人员不占客户名额。", "名额")
draw_table(
    s, Inches(0.4), Inches(1.35),
    [2.4, 2.2, 2.4, 2.4, 2.8], 0.7,
    ["类型", "客户人数", "存量客户", "新客户", "邀请方式"],
    [
        ["精品参访", "16—18 人", "约 10—12 人", "约 4—6 人", "客户经理一对一"],
        ["高端小局", "12 人", "重点客户", "按主题另留", "提名，不公开报名"],
        ["11/26 新客局", "12 人", "4 位老客户背书", "8 位新客户", "老客户带 1 家"],
        ["12/17 收官", "16 人", "高满意客户", "在跟的新客户", "只邀愿意再来的人"],
    ],
    body_size=14, header_size=13,
)
txt(s, Inches(0.5), Inches(5.15), Inches(12.3), Inches(1.6),
    [("客户经理 9 月 24 日前交出两张名单：重点存量客户 8 家，意向新客户 5 家。", 16, WHITE, False, 8),
     ("相邻两场名单默认不重复。一位老客户整个季度被邀请不超过 3 次，避免邀烦。", 16, AMBER_LT, False)])

# ---------------------------------------------------------------------------
# 8 预算模板
# ---------------------------------------------------------------------------
s = new("5000 元以内怎么花", "数字是建议列支。单项可以调，单场合计不能过 5000。", "预算")
draw_table(
    s, Inches(0.4), Inches(1.3),
    [2.5, 3.2, 3.2, 3.3], 0.5,
    ["费用项", "参访（16人）", "高端小局（12人）", "收官小局（16人）"],
    [
        ["场地", "0，企业现场", "0，银行贵宾室", "0，银行贵宾室"],
        ["交通", "1800", "0", "0"],
        ["茶水 / 茶叙", "600", "1800", "2400"],
        ["物料", "400", "300", "400"],
        ["应急", "400", "400", "400"],
        ["建议合计", "3200", "2500", "3200"],
        ["硬顶", "5000", "5000", "5000"],
    ],
    body_size=13, header_size=13,
)
txt(s, Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.55),
    [("默认不采购伴手礼、不请付费嘉宾、不请摄影团队。余量留下，不主动花完。", 15, AMBER_LT, False)])

# ---------------------------------------------------------------------------
# 9 闸门
# ---------------------------------------------------------------------------
s = new("费用闸门：过线就停", "Excel「决算闸门」里，金额大于 5000 会显示超支。", "闸门")
gates = [
    ("立项", "预算表超过 5000 元，不得发出邀请。"),
    ("采购前", "活动前 5 天复核报价。超了就减茶歇、减交通或减物料。"),
    ("决算", "录入超过 5000 元，下一场改为零采购：只用场地和自带茶水。"),
    ("礼品", "要采购须合规书面同意，且整场仍然不超过 5000 元。"),
]
for i, (t, d) in enumerate(gates):
    col, row = i % 2, i // 2
    x = Inches(0.45) + col * Inches(6.4)
    y = Inches(1.4) + row * Inches(2.5)
    rrect(s, x, y, Inches(6.15), Inches(2.25), TEAL)
    txt(s, x + Inches(0.3), y + Inches(0.35), Inches(5.6), Inches(0.45),
        [(t, 20, AMBER_LT, True)])
    txt(s, x + Inches(0.3), y + Inches(1.05), Inches(5.6), Inches(0.85),
        [(d, 16, OFF, False)])

# ---------------------------------------------------------------------------
# 10 高端感
# ---------------------------------------------------------------------------
s = new("不花钱堆出来的高端", "高端指人被照顾到，不是酒店和礼品。", "体验")
items = [
    ("人少", "参访不超过 18 位客户，小局 12 位。每个人都有位子、有人认识。"),
    ("有人陪", "客户经理陪自己的客户走完全程，不把客户留在会场。"),
    ("只一件事", "一场只一个主题。讲不完的内容记进服务卡，不往议程里加。"),
    ("半天", "13:30 开始，16:30 结束。参访动线控制在 90 分钟内。"),
    ("被记住", "离场前客户经理说出今天记下的那一件事，客户点头才算完成。"),
    ("次日回", "第二天单独问候，不发群公告充数。不满意的，项目经理要知道。"),
]
for i, (t, d) in enumerate(items):
    col, row = i % 3, i // 3
    x = Inches(0.4) + col * Inches(4.25)
    y = Inches(1.35) + row * Inches(2.7)
    rrect(s, x, y, Inches(4.05), Inches(2.5), TEAL)
    txt(s, x + Inches(0.22), y + Inches(0.3), Inches(3.6), Inches(0.45),
        [(t, 20, AMBER_LT, True)])
    txt(s, x + Inches(0.22), y + Inches(0.95), Inches(3.6), Inches(1.2),
        [(d, 14, OFF, False)])

# ---------------------------------------------------------------------------
# 11 参访
# ---------------------------------------------------------------------------
s = new("四场参访：客户走出去", "被访企业不收场地费。同意接待、路线和拍摄边界，三件事书面确认后再邀请客户。", "参访")
draw_table(
    s, Inches(0.35), Inches(1.3),
    [1.5, 3.3, 3.6, 3.8], 0.85,
    ["日期", "去哪里", "客户带走什么", "启动条件"],
    [
        ["10/22", "智能制造或具身智能企业", "看见真实产线，老客户感到被请去", "国庆后 10/8 前锁定企业"],
        ["11/19", "链主企业或 AI 应用现场", "老客户带 1 家同行或上下游", "10/30 前企业书面同意"],
        ["12/3", "跨境窗口或园区", "出海不再只是听一堂课", "11/12 前确认开放范围"],
        ["12/10", "客户点名的标杆企业", "答谢，主题来自客户愿望", "11/20 前从客户回访里选题"],
    ],
    body_size=13, header_size=13,
)
txt(s, Inches(0.45), Inches(6.15), Inches(12.4), Inches(0.7),
    [("现场秩序：无拍摄区、随行引导、材料不外带。客户经理负责自己的客户，不负责拉业务承诺。", 14, OFF, False)])

# ---------------------------------------------------------------------------
# 12 高端小局
# ---------------------------------------------------------------------------
s = new("三场高端小局：客户坐下来", "银行贵宾会议室。条线负责人到场问候即可，不承诺某位行领导出席。", "小局")
card(s, Inches(0.4), Inches(1.3), Inches(4.1), Inches(5.4), "11/5 私享问诊", [
    "12 位重点存量客户",
    "每家 8 分钟",
    "只问清一件事和下一步",
    "当场写下责任人与日期",
    "不在现场判断能不能贷款",
], OK, 15)
card(s, Inches(4.65), Inches(1.3), Inches(4.1), Inches(5.4), "11/26 新客闭门", [
    "4 位老客户 + 8 位新客户",
    "老客户只讲自己的经营体验",
    "不讲“我行保证审批”",
    "新客户离场拿到对接人",
    "5 个工作日内完成首联",
], AMBER, 15)
card(s, Inches(8.9), Inches(1.3), Inches(4.0), Inches(5.4), "12/17 年度收官", [
    "邀请满意、愿意再来的客户",
    "感谢具体到人，不念稿",
    "未结事项当面约定时点",
    "问明年还想看哪一类企业",
    "当场记下次年意向",
], AMBER, 15)

# ---------------------------------------------------------------------------
# 13 SLA
# ---------------------------------------------------------------------------
s = new("客户离开之后怎么接", "服务事项仍要有人跟，但考核先看客户有没有被接住。", "时限")
draw_table(
    s, Inches(0.45), Inches(1.35),
    [2.2, 5.5, 4.5], 0.7,
    ["时点", "必须做完的事", "算完成的标准"],
    [
        ["当天", "服务卡写下一件事", "客户点头，客户经理签名"],
        ["次日", "单独问感受", "满意 / 一般 / 不满意，有一句原因"],
        ["2 日内", "不满意的电话回访", "项目经理知道原因和处理"],
        ["3 个工作日", "台账建档", "客户类型、场次、需求可查"],
        ["5 个工作日", "新客户第一次联系", "联系记录落到客户经理名下"],
        ["10 个工作日", "服务事项第一次反馈", "受理、补件、转介或不适配，说清楚"],
    ],
    body_size=14, header_size=14,
)

# ---------------------------------------------------------------------------
# 14 话术边界
# ---------------------------------------------------------------------------
s = new("邀请和现场可以说什么", "客户高兴，来自被尊重；不来自被承诺。", "话术")
card(s, Inches(0.4), Inches(1.35), Inches(6.2), Inches(5.4), "可以这样说", [
    "这场人少，客户经理会陪着您",
    "去一家值得看的企业，半天结束",
    "活动不收费",
    "您可以带一位同行或合作伙伴",
    "新客户 5 个工作日内会有人联系",
    "贷款能否办理，由银行自己审批",
], OK, 16)
card(s, Inches(6.8), Inches(1.35), Inches(6.1), Inches(5.4), "不要这样说", [
    "保证给额度、给利率、给放款时间",
    "行长一定会见您",
    "来了就有内部名额",
    "送礼、返点或换开户",
    "这场是发布某款联名产品",
    "在微信群里收财务报表",
], AMBER, 16)

# ---------------------------------------------------------------------------
# 15 责任
# ---------------------------------------------------------------------------
s = new("谁对客户负责", "客户开不开心，先问陪同的客户经理，再问项目经理。", "责任")
draw_table(
    s, Inches(0.4), Inches(1.35),
    [3.4, 4.4, 4.4], 0.72,
    ["事项", "谁拍板", "谁去办"],
    [
        ["客户提名与陪同", "客户经理", "客户经理"],
        ["满意度与次日回访", "客户经理", "客户经理"],
        ["参访企业落实", "项目经理", "科企联 / 项目办"],
        ["单场不超过 5000 元", "项目经理", "会务"],
        ["对外说法", "银行合规窗口", "主持人、客户经理"],
        ["新客户 5 日内联系", "客户经理", "客户经理"],
        ["不满意升级", "项目经理", "客户经理先打电话"],
    ],
    body_size=15, header_size=14,
)

# ---------------------------------------------------------------------------
# 16 KPI
# ---------------------------------------------------------------------------
s = new("年底对得上的数字", "不按贷款金额给这场活动打分。", "指标")
draw_table(
    s, Inches(0.4), Inches(1.3),
    [4.2, 2.2, 2.2, 3.6], 0.62,
    ["指标", "10月", "11月", "12月 / 累计"],
    [
        ["场次", "1", "3", "3，累计 7"],
        ["存量客户人次", "12", "26", "38，累计不少于 60"],
        ["新客户到场", "4", "14", "12，累计不少于 25"],
        ["老客户转介到场", "—", "6", "累计不少于 8"],
        ["满意度", "≥90%", "≥90%", "每场都要 ≥90%"],
        ["新客 5 日内联系", "100%", "100%", "100%"],
        ["单场决算", "≤5000", "≤5000", "7 场全部 ≤5000"],
    ],
    body_size=14, header_size=14,
)

# ---------------------------------------------------------------------------
# 17 行动
# ---------------------------------------------------------------------------
s = new("从 9 月 23 日起", "国庆前不办活动。先把客户名单和第一家参访企业定下来。", "开工")
draw_table(
    s, Inches(0.4), Inches(1.28),
    [1.7, 7.6, 2.9], 0.58,
    ["日期", "动作", "谁来做"],
    [
        ["9/23", "确认 7 场框架和 5000 元硬顶", "四方联系人"],
        ["9/24", "交出重点客户 8 家、意向新客 5 家", "客户经理"],
        ["9/28", "储备 3 家愿意免费接待的企业", "项目办 / 科企联"],
        ["9/30", "确认话术：不承诺贷款、默认无伴手礼", "银行合规"],
        ["10/8", "敲定 10/22 参访企业和客户名单", "项目经理"],
        ["10/9", "一对一发邀请", "客户经理"],
        ["10/19", "预走访线路", "会务 + 被访企业"],
        ["10/22", "第一场，当天收回服务卡", "全体陪同"],
    ],
    body_size=14, header_size=13,
)

# ---------------------------------------------------------------------------
# 18 收尾
# ---------------------------------------------------------------------------
s = new()
rect(s, 0, 0, SW, Pt(5), AMBER)
rect(s, Inches(0.7), Inches(1.7), Inches(0.12), Inches(1.2), AMBER)
txt(s, Inches(1.05), Inches(1.6), Inches(11), Inches(0.4),
    [("执行时记住这三句", 16, AMBER_LT, False)])
txt(s, Inches(1.0), Inches(2.2), Inches(11.5), Inches(2.4),
    [("一场不超五千。", 32, WHITE, True, 10),
     ("年底多走出去。", 32, WHITE, True, 10),
     ("老客户有人陪，新客户五日内有回音。", 28, WHITE, True)])
txt(s, Inches(1.0), Inches(5.2), Inches(11.5), Inches(1.4),
    [("配套表格：七场日历、决算闸门、客户出席与回访、新客跟进。金额、名单和日期填进表里再生效。", 15, AMBER_LT, False, 8),
     ("本方案不构成授信、利率、期限或放款承诺。金融业务由泰隆银行独立审批。", 13, MUTED, False)])

out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "deliverables")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "泰隆银行客户参访与高端小局_2026年10-12月.pptx")
prs.save(out_path)
print(f"已生成：{out_path}（共 {PAGE} 页）")
