#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按既有劳动合同版式重排《胡继刚-2》，不改条款文字。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generate_labor_contract as base

SONG = base.SONG
HEI = base.HEI
INK = base.INK
MUTED = base.MUTED
LABEL_FILL = base.LABEL_FILL
HEAD_FILL = base.HEAD_FILL

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source" / "劳动合同书_上海贝迪创建材科技有限公司_胡继刚-2.docx"
OUT = ROOT / "output" / "劳动合同书_上海贝迪创建材科技有限公司_胡继刚-2.docx"

COMPANY = "上海贝迪创建材科技有限公司"
ARTICLE_RE = re.compile(r"^(第[一二三四五六七八九十百零]+条)[　\s]+(.*)$", re.S)
CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十百零]+章")


def norm(text: str) -> str:
    return re.sub(r"\s+", "", text)


def cell_lines(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for tr in table._tbl.findall(qn("w:tr")):
        cells = []
        for tc in tr.findall(qn("w:tc")):
            cells.append("".join(node.text or "" for node in tc.findall(".//" + qn("w:t"))))
        rows.append(cells)
    return rows


def display_label(text: str) -> str:
    return text.replace("\u3000", "").strip()


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    base.set_paragraph_format(first_header.paragraphs[0], before=0, after=0, line=1.0)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    base.set_paragraph_format(hp, before=0, after=2, line=1.0, align="left")
    left = hp.add_run(COMPANY)
    base.set_run_font(left, SONG, 9, color=MUTED)
    right = hp.add_run("　劳动合同书")
    base.set_run_font(right, SONG, 9, color=MUTED)
    base.add_bottom_border(hp, sz="8")

    def fill_footer(paragraph) -> None:
        base.set_paragraph_format(paragraph, before=2, after=0, line=1.0, align="center")
        base.add_top_border(paragraph, sz="6")
        lead = paragraph.add_run("第 ")
        base.set_run_font(lead, SONG, 9, color=MUTED)
        base.add_field(paragraph, "PAGE")
        mid = paragraph.add_run(" 页")
        base.set_run_font(mid, SONG, 9, color=MUTED)

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    fill_footer(first_footer.paragraphs[0])
    footer = section.footer
    footer.is_linked_to_previous = False
    fill_footer(footer.paragraphs[0])


def add_party_block(doc, rows: list[list[str]]) -> None:
    """rows 含甲方、乙方两段，每段首行为跨列标题。"""
    blocks: list[list[list[str]]] = []
    current: list[list[str]] = []
    for row in rows:
        if len(row) == 1 and row[0]:
            if current:
                blocks.append(current)
            current = [row]
        else:
            current.append(row)
    if current:
        blocks.append(current)

    for index, block in enumerate(blocks):
        if index:
            base.add_text(doc, "", size=6, after=2, line=1.0)
        title = block[0][0]
        data = [row for row in block[1:] if any(cell.strip() for cell in row)]
        table = base.add_table(doc, 1 + len(data), [4.2, 12.4])
        table.cell(0, 0).merge(table.cell(0, 1))
        base.write_cell(table.cell(0, 0), title, bold=True, size=11, font=HEI, fill=HEAD_FILL)
        base.set_row_cant_split(table.rows[0], 0.78)
        for row_index, row in enumerate(data, start=1):
            label = display_label(row[0]) if row else ""
            value = row[1] if len(row) > 1 else ""
            base.write_cell(table.cell(row_index, 0), label, bold=True, fill=LABEL_FILL, align="center")
            base.write_cell(table.cell(row_index, 1), value, bold=(label in {"名称", "姓名"}))
            base.set_row_cant_split(table.rows[row_index], 0.7)


def add_article_paragraph(doc, text: str) -> None:
    matched = ARTICLE_RE.match(text)
    paragraph = doc.add_paragraph()
    base.set_paragraph_format(
        paragraph,
        before=8,
        after=1,
        line=1.2,
        indent_cm=0.85,
        align="justify",
    )
    if matched:
        lead = paragraph.add_run(f"{matched.group(1)}　")
        base.set_run_font(lead, SONG, 12, bold=True)
        body = paragraph.add_run(matched.group(2))
        base.set_run_font(body, SONG, 12)
    else:
        run = paragraph.add_run(text)
        base.set_run_font(run, SONG, 12)


def add_body_paragraph(doc, text: str) -> None:
    base.add_text(
        doc,
        text,
        before=1,
        after=1,
        line=1.2,
        indent_cm=0.85,
        align="justify",
    )


def add_signature_table(doc) -> None:
    sign = base.add_table(doc, 6, [8.3, 8.3])
    base.write_cell(sign.cell(0, 0), "甲方（盖章）", bold=True, font=HEI, size=11, fill=HEAD_FILL, align="center")
    base.write_cell(sign.cell(0, 1), "乙方（本人签字）", bold=True, font=HEI, size=11, fill=HEAD_FILL, align="center")
    base.write_cell(sign.cell(1, 0), f"甲方：{COMPANY}", bold=True, align="center")
    base.write_cell(sign.cell(1, 1), "乙方：胡继刚", bold=True, align="center")
    base.write_cell(sign.cell(2, 0), "", align="center")
    base.write_cell(sign.cell(2, 1), "", align="center")
    base.write_cell(sign.cell(3, 0), "法定代表人或授权代表（签字）：", align="left")
    base.write_cell(sign.cell(3, 1), "", align="left")
    base.write_cell(sign.cell(4, 0), "签署日期：＿＿＿＿年＿＿月＿＿日", align="left")
    base.write_cell(sign.cell(4, 1), "签署日期：＿＿＿＿年＿＿月＿＿日", align="left")
    base.write_cell(sign.cell(5, 0), "签署地点：＿＿＿＿＿＿＿＿＿＿", align="left")
    base.write_cell(sign.cell(5, 1), "签署地点：＿＿＿＿＿＿＿＿＿＿", align="left")
    for index, height in ((0, 0.78), (1, 0.9), (2, 3.6), (3, 1.15), (4, 0.82), (5, 0.82)):
        base.set_row_cant_split(sign.rows[index], height)


def add_receipt_table(doc, text: str) -> None:
    left, right = text.split("领取日期：", 1)
    table = base.add_table(doc, 1, [8.3, 8.3])
    base.write_cell(table.cell(0, 0), left.strip(), align="left")
    base.write_cell(table.cell(0, 1), "领取日期：" + right.strip(), align="left")
    base.set_row_cant_split(table.rows[0], 1.3)


def new_document() -> Document:
    doc = Document()
    base.configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.7)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.45)
    configure_header_footer(section)
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        section._sectPr.remove(grid)
    if doc.paragraphs and not doc.paragraphs[0].text:
        doc.element.body.remove(doc.paragraphs[0]._element)
    core = doc.core_properties
    core.title = "劳动合同书"
    core.subject = f"{COMPANY}与胡继刚"
    core.category = "劳动合同"
    core.last_modified_by = COMPANY
    core.comments = ""
    for child in core._element:
        if child.tag.endswith("creator"):
            child.text = COMPANY
    return doc


def build() -> Document:
    source = Document(SOURCE)
    paragraphs = [paragraph.text for paragraph in source.paragraphs]
    doc = new_document()

    number = next(text for text in paragraphs if text.startswith("合同编号"))
    base.add_text(doc, number, size=10.5, align="right", before=0, after=2, line=1.0)
    base.add_text(
        doc,
        "劳 动 合 同 书",
        size=22,
        font=HEI,
        bold=True,
        align="center",
        before=2,
        after=2,
        line=1.0,
    )
    subtitle = base.add_text(
        doc,
        "（固定期限）",
        size=12,
        font=HEI,
        align="center",
        before=0,
        after=6,
        line=1.0,
        color=MUTED,
    )
    base.add_bottom_border(subtitle, sz="16", color="1A1A1A")
    add_party_block(doc, cell_lines(source.tables[0]))

    for text in paragraphs:
        if not text.strip():
            continue
        if text.startswith("合同编号") or text in {"劳动合同书", "（固定期限）"}:
            continue
        if text == "签字盖章页":
            base.add_text(
                doc,
                text,
                size=16,
                font=HEI,
                bold=True,
                align="center",
                before=16,
                after=6,
                line=1.0,
            )
            continue
        if text.startswith("双方已阅读"):
            add_body_paragraph(doc, text)
            base.add_text(doc, "", size=6, before=6, after=2, line=1.0)
            add_signature_table(doc)
            continue
        if text == "合同文本领取确认":
            base.add_text(doc, text, size=12, font=HEI, bold=True, before=14, after=4, align="left")
            continue
        if text.startswith("乙方签字："):
            add_receipt_table(doc, text)
            continue
        if CHAPTER_RE.match(text):
            base.add_text(
                doc,
                text,
                size=14,
                font=HEI,
                bold=True,
                before=12,
                after=2,
                line=1.05,
                align="left",
                keep_next=True,
            )
            continue
        if ARTICLE_RE.match(text):
            add_article_paragraph(doc, text)
            continue
        add_body_paragraph(doc, text)
    return doc


def output_text(doc: Document) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in cell_lines(table):
            parts.extend(row)
    return "\n".join(parts)


def assert_same_content(source: Document, formatted: Document) -> None:
    produced = norm(output_text(formatted))
    missing: list[str] = []
    for paragraph in source.paragraphs:
        text = paragraph.text.strip()
        if not text or text == "劳动合同书":
            continue
        if norm(text) not in produced:
            missing.append(text)
    for table in source.tables:
        for tc in table._tbl.findall(".//" + qn("w:tc")):
            for paragraph in tc.findall(qn("w:p")):
                text = "".join(node.text or "" for node in paragraph.findall(".//" + qn("w:t"))).strip()
                if text and norm(text) not in produced:
                    missing.append(text)
    required_lines = [
        "甲方：上海贝迪创建材科技有限公司",
        "乙方：胡继刚",
        "（盖章）",
        "（本人签字）",
        "法定代表人或授权代表（签字）：",
        "签署日期：＿＿＿＿年＿＿月＿＿日",
        "签署地点：＿＿＿＿＿＿＿＿＿＿",
    ]
    for line in required_lines:
        if norm(line) not in produced:
            missing.append(line)
    if missing:
        raise SystemExit("排版后缺少原文：" + " / ".join(missing[:12]))
    if "上海贝迪创建科技有限公司" in output_text(formatted):
        raise SystemExit("用人单位名称被改掉了")
    if "康桥东路" in produced or "1987年4月25日" in produced:
        raise SystemExit("排版写入了原文没有的内容")


def main() -> None:
    source = Document(SOURCE)
    doc = build()
    assert_same_content(source, doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"已生成：{OUT}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # noqa: BLE001
        print(f"生成失败：{exc}", file=sys.stderr)
        raise
