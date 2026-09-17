#!/usr/bin/env python3
"""核对报名操作指引 Word 是否包含关键信息，且已去掉套话。"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "deliverables" / "重庆市九龙坡区城市更新十五五实施方案_报名及采购文件获取操作指引.docx"

REQUIRED = [
    "重庆市九龙坡区高质量推动城市更新“十五五”时期实施方案",
    "重庆大学设计总院报名及采购文件获取操作指引",
    "JLP26C00100",
    "重庆市九龙坡区住房和城乡建设委员会",
    "竞争性磋商",
    "132万元",
    "2026年9月18日18:00",
    "2026年9月23日14:30",
    "城乡规划编制甲级资质",
    "重庆大学建筑规划设计研究总院有限公司",
    "复旦大学住房政策研究中心",
    "重庆市政府采购网",
    "CA数字证书",
    "报名审核通过",
    "合法获取采购文件",
    "三信建设咨询集团有限公司",
    "023-63024313",
    "023-68036909",
    "赵茜",
    "联合体",
    "我的投标项目",
    "摄像头",
]

FORBIDDEN = [
    "最重要的一句话",
    "建议暂按以下结构推进",
    "这是整个流程中最关键的一步",
    "请不要只回复",
    "需要特别注意",
    "必须注意",
    "不要卡最后时间",
    "建议不要等",
    "作为AI",
    "根据您提供",
    "希望对您有帮助",
    "当然可以",
    "总的来说",
    "综上所述",
]


def extract_text(doc: Document) -> str:
    parts: list[str] = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main() -> int:
    if not DOCX.exists():
        print(f"未找到文件：{DOCX}", file=sys.stderr)
        return 1

    doc = Document(DOCX)
    text = extract_text(doc)
    size_kb = DOCX.stat().st_size / 1024

    missing = [item for item in REQUIRED if item not in text]
    leftover = [item for item in FORBIDDEN if item in text]

    print(f"文件：{DOCX.name}")
    print(f"大小：{size_kb:.1f} KB")
    print(f"段落：{len(doc.paragraphs)}")
    print(f"表格：{len(doc.tables)}")
    print(f"必含条目：{len(REQUIRED) - len(missing)}/{len(REQUIRED)}")

    if missing:
        print("缺少：")
        for item in missing:
            print(f"  - {item}")
    if leftover:
        print("仍含套话：")
        for item in leftover:
            print(f"  - {item}")

    if missing or leftover:
        return 1
    if len(doc.tables) < 6:
        print(f"表格数量偏少：{len(doc.tables)}", file=sys.stderr)
        return 1
    print("核对通过。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
