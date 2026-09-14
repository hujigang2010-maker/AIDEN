# -*- coding: utf-8 -*-
"""校验三件套是否生成完整、内容是否写入。"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
DELIV = ROOT / "deliverables"

PPT = DELIV / "东昇聚变_政府事务与上海产业落地_90天工作设想.pptx"
DOC = DELIV / "东昇聚变_政府事务与上海产业落地_90天工作设想.docx"
XLS = DELIV / "东昇聚变_政府事务与上海产业落地_90天工作台账.xlsx"

EXPECTED_SLIDES = 16
EXPECTED_SHEETS = ["90天总览", "分周计划", "产出物", "政府关系作战图", "选址比选", "现场六指标", "风险与支持"]
KEYWORDS = ["东昇聚变", "90天", "浦东", "杨浦", "补位不抢位", "金桥", "张江", "马桥"]


def ppt_text(path: Path) -> str:
    prs = Presentation(path)
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text)
    return "\n".join(chunks)


def fail(msg: str):
    print(f"FAIL  {msg}")
    sys.exit(1)


def main():
    for p in (PPT, DOC, XLS):
        if not p.exists():
            fail(f"缺少文件 {p}")
        if p.stat().st_size < 8000:
            fail(f"文件过小 {p} {p.stat().st_size}")
        if not zipfile.is_zipfile(p):
            fail(f"不是有效 Office 文件 {p}")

    prs = Presentation(PPT)
    if len(prs.slides) != EXPECTED_SLIDES:
        fail(f"PPT 页数 {len(prs.slides)} != {EXPECTED_SLIDES}")
    if abs(prs.slide_width - 12191640) > 200:
        fail(f"PPT 宽度异常 {prs.slide_width}")

    text = ppt_text(PPT)
    doc = Document(DOC)
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    wb = load_workbook(XLS)
    if wb.sheetnames != EXPECTED_SHEETS:
        fail(f"工作表不符 {wb.sheetnames}")

    blob = text + "\n" + doc_text
    for sh in wb.sheetnames:
        ws = wb[sh]
        for row in ws.iter_rows(max_row=30, max_col=8, values_only=True):
            blob += "\n" + " ".join("" if v is None else str(v) for v in row)

    missing = [k for k in KEYWORDS if k not in blob]
    if missing:
        fail(f"关键词缺失 {missing}")

    if "专职拿地" not in blob and "不是本岗主责" not in blob:
        fail("未写清拿地边界")

    print("PASS  PPT/DOCX/XLSX 校验通过")
    print(f"  PPT  {len(prs.slides)} 页  {PPT.stat().st_size} bytes")
    print(f"  DOCX {len(doc.paragraphs)} 段  {DOC.stat().st_size} bytes")
    print(f"  XLSX {wb.sheetnames}  {XLS.stat().st_size} bytes")


if __name__ == "__main__":
    main()
