#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""核对工作经历表格与原文一致。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from generate_work_history import (  # noqa: E402
    RECORDS,
    find_gaps,
    main_row,
    split_dept_title,
    split_row,
    write_outputs,
)


class WorkHistoryTests(unittest.TestCase):
    def test_record_count_and_order(self) -> None:
        self.assertEqual(len(RECORDS), 18)
        self.assertEqual([record["no"] for record in RECORDS], [str(i) for i in range(1, 19)])
        starts = [record["start"] for record in RECORDS]
        self.assertEqual(starts, sorted(starts))

    def test_verbatim_endpoints(self) -> None:
        first = RECORDS[0]
        self.assertEqual(first["start"], "2015-09")
        self.assertEqual(first["end"], "2015-10")
        self.assertEqual(first["org"], "上海睿泽股权投资管理有限公司")
        self.assertEqual(first["dept_title"], "工程部/资料员、施工员")
        self.assertEqual(first["work"], "施工资料及现场管理")

        nanjing = RECORDS[15]
        self.assertEqual(nanjing["region"], "江苏省南京市")
        self.assertEqual(nanjing["org"], "江苏省建筑装饰设计研究院有限公司")
        self.assertEqual(nanjing["work"], "工程技术及项目管理")

        current = RECORDS[17]
        self.assertEqual(current["end"], "今")
        self.assertEqual(current["org"], "上海贝迪创建科技有限公司")
        self.assertEqual(current["dept_title"], "工程部/高级工程师")

        self.assertEqual(RECORDS[1]["work"], "工程管理级施工协调")
        self.assertEqual(RECORDS[3]["dept_title"], "工程部总监")
        self.assertEqual(RECORDS[11]["work"], "钢结构施工技术管理")
        self.assertTrue(all(record["country"] == "中华人民共和国" for record in RECORDS))

    def test_known_gaps_only(self) -> None:
        self.assertEqual(
            find_gaps(),
            [
                "第1段与第2段之间空缺 2015-11",
                "第13段与第14段之间空缺 2023-09",
            ],
        )

    def test_department_split(self) -> None:
        self.assertEqual(split_dept_title("工程部/资料员、施工员"), ("工程部", "资料员、施工员"))
        self.assertEqual(split_dept_title("工程部总监"), ("工程部", "总监"))
        self.assertEqual(split_dept_title("工程部副总经理"), ("工程部", "副总经理"))
        self.assertEqual(split_dept_title("工程部/高级工程师"), ("工程部", "高级工程师"))

    def test_generated_files_match_source(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            paths = write_outputs(Path(folder))
            workbook = load_workbook(paths["xlsx"])
            main = workbook["工作经历"]
            self.assertEqual([main.cell(3, col).value for col in range(1, 8)], [
                "序号",
                "起止时间",
                "国家地区",
                "国家省市",
                "工作单位",
                "部门或职务",
                "主要工作内容",
            ])
            for index, record in enumerate(RECORDS):
                excel_row = 4 + index
                expected = main_row(record)
                expected[0] = int(expected[0])
                actual = [main.cell(excel_row, col).value for col in range(1, 8)]
                self.assertEqual(actual, expected)

            split = workbook["分列填写"]
            for index, record in enumerate(RECORDS):
                actual = [split.cell(4 + index, col).value for col in range(1, 11)]
                expected = split_row(record)
                expected[0] = int(expected[0])
                self.assertEqual(actual, expected)

            document = Document(paths["docx"])
            table = document.tables[0]
            self.assertEqual(len(table.rows), 19)
            for index, record in enumerate(RECORDS):
                actual = [table.cell(index + 1, col).text for col in range(7)]
                self.assertEqual(actual, main_row(record))

            html = paths["html"].read_text(encoding="utf-8")
            for record in RECORDS:
                self.assertIn(record["org"], html)
                self.assertIn(record["dept_title"], html)
                self.assertIn(record["work"], html)


if __name__ == "__main__":
    unittest.main()
