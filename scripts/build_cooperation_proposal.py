#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《杨浦科技企业成长与产业协同合作计划》内部建议稿。

综合 2026-09-24 洽谈记录研判与后续商务结构设计。
输出为联合会内部决策用 PPT，不是已经向对方发出的报价。
"""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
W, H = 13.333, 7.5
TOTAL = 17

NAVY = "16324F"
INK = "1C2833"
MUTED = "5E6D7A"
TEAL = "1F5C57"
TEAL_SOFT = "E6F1EF"
GOLD = "8C6A2F"
GOLD_SOFT = "F6F0E2"
BG = "F6F3EC"
WHITE = "FFFFFF"
CARD = "FFFCF7"
LINE = "E4DACC"
RED = "8C3D3D"
RED_SOFT = "F8EEEE"
PAPER = "EFEAE1"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def font_run(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rpr, qn(tag))
        node.set("typeface", FONT)


def rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    return shape


def paint(slide, color):
    shape = rect(slide, 0, 0, W, H, color)
    tree = slide.shapes._spTree
    element = shape._element
    tree.remove(element)
    tree.insert(2, element)


def write(slide, paragraphs, x, y, w, h, anchor="t"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    frame._txBody.bodyPr.set("anchor", anchor)
    for index, para in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = para.get("align", PP_ALIGN.LEFT)
        paragraph.space_before = Pt(para.get("before", 0))
        paragraph.space_after = Pt(para.get("after", 0))
        if "spacing" in para:
            paragraph.line_spacing = para["spacing"]
        parts = para["text"].split("\n")
        for part_index, part in enumerate(parts):
            if part_index:
                etree.SubElement(paragraph._p, qn("a:br"))
            run = paragraph.add_run()
            run.text = part
            font_run(run, para.get("size", 14), para.get("color", INK), para.get("bold", False))
    return box


def t(text, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, before=0, after=0, spacing=1.08):
    return {
        "text": text,
        "size": size,
        "color": color,
        "bold": bold,
        "align": align,
        "before": before,
        "after": after,
        "spacing": spacing,
    }


def blank_slide(prs, color=BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint(slide, color)
    return slide


def footer(slide, page):
    rect(slide, 0.5, 7.02, 12.33, 0.01, LINE)
    write(slide, [t("内部建议稿  ·  谈判起点，不是对方已确认的预算", 11, MUTED)], 0.5, 7.1, 9.2, 0.28)
    write(
        slide,
        [t(f"{page:02d}  /  {TOTAL:02d}", 11, MUTED, align=PP_ALIGN.RIGHT)],
        10.4,
        7.1,
        2.43,
        0.28,
    )


def header(slide, kicker, title):
    rect(slide, 0, 0, W, 0.07, TEAL)
    write(slide, [t(kicker, 12, GOLD, True)], 0.52, 0.22, 12.2, 0.26)
    write(slide, [t(title, 26, NAVY, True, spacing=1.0)], 0.5, 0.48, 12.3, 0.48)
    rect(slide, 0.52, 1.06, 0.86, 0.03, GOLD)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def card(slide, x, y, w, h, fill=CARD, line=LINE):
    return rect(slide, x, y, w, h, fill, line)


def slide_cover(prs):
    slide = blank_slide(prs, NAVY)
    rect(slide, 0, 0, 0.14, H, GOLD)
    write(slide, [t("内部建议稿", 14, "E4C98A", True)], 0.72, 0.48, 6, 0.32)
    write(slide, [t("2026年9月24日洽谈  ·  约1小时39分钟", 14, "C9D4DE")], 0.72, 0.84, 8, 0.3)
    write(
        slide,
        [
            t("杨浦科技企业成长", 40, WHITE, True, spacing=1.0),
            t("与产业协同合作计划", 40, WHITE, True, before=4, spacing=1.0),
        ],
        0.7,
        1.7,
        12,
        1.7,
    )
    rect(slide, 0.74, 3.6, 1.4, 0.035, GOLD)
    write(
        slide,
        [
            t("成都新谷投资集团（上海新谷湾 / 新谷国际）", 18, "F4F7F8", spacing=1.05),
            t("×  上海市杨浦区科技企业联合会", 18, "F4F7F8", before=6, spacing=1.05),
        ],
        0.72,
        3.85,
        11,
        0.9,
    )
    write(
        slide,
        [
            t(
                "综合洽谈原文研判与商务结构设计。供联合会内部决策，不作为已经发出的报价。",
                14,
                "C9D4DE",
            )
        ],
        0.72,
        5.05,
        10.5,
        0.4,
    )
    labels = [
        (0.72, "本地合作先落地"),
        (3.55, "双城合作留后手"),
        (6.38, "海外合作等条件成熟"),
    ]
    for x, label in labels:
        rect(slide, x, 6.15, 2.65, 0.52, "1E4468")
        write(slide, [t(label, 13, WHITE, True, align=PP_ALIGN.CENTER)], x, 6.27, 2.65, 0.3)
    notes(
        slide,
        "本稿把两轮判断收成一个项目：以会员建立关系，以联合服务站做载体，"
        "以招商和早期资本对接形成收入，以成都和新加坡作为后续验证，不在第一阶段无偿铺开。",
    )


def slide_judgment(prs):
    slide = blank_slide(prs)
    header(slide, "01   总体判断", "值得推进，但只从一个有偿试点开始")
    footer(slide, 2)
    columns = [
        (
            "01",
            "上海本地",
            "最容易落地",
            "企业服务、精准招商、小型活动、服务站和早期资本对接。直接对应出租率、现金流和品牌节点。",
        ),
        (
            "02",
            "上海 — 成都",
            "长期价值最高",
            "长三角企业有真实需求时，再用对方合计约80万平方米的载体和既有飞地通道。首期不投入无偿组织成本。",
        ),
        (
            "03",
            "海外延伸",
            "先核实再宣传",
            "对方波士顿平台已运行，新加坡正在选址。两周内由业务负责人摸清用途，成熟后再谈联合设点。",
        ),
    ]
    for index, (num, title, tag, body) in enumerate(columns):
        x = 0.5 + index * 4.2
        card(slide, x, 1.32, 4.0, 4.15)
        rect(slide, x, 1.32, 4.0, 0.08, TEAL if index == 0 else GOLD)
        write(slide, [t(num, 18, GOLD, True)], x + 0.24, 1.56, 1.2, 0.36)
        write(slide, [t(title, 22, NAVY, True)], x + 0.24, 2.05, 3.5, 0.42)
        write(slide, [t(tag, 13, TEAL, True)], x + 0.24, 2.55, 3.5, 0.3)
        write(slide, [t(body, 14, INK, spacing=1.15)], x + 0.24, 3.1, 3.52, 2.0)
    card(slide, 0.5, 5.62, 12.33, 1.2, TEAL_SOFT, "C9DDD9")
    write(
        slide,
        [
            t("首期只争取三件成果", 15, TEAL, True),
            t(
                "推动上海公司按章程申请副会长；拿下三个月产业服务试点并写明专项预算；把挂牌做成联合服务站，而不是单独交付一块牌子。",
                14,
                INK,
                before=4,
            ),
        ],
        0.74,
        5.78,
        11.85,
        0.92,
    )


def slide_signal(prs):
    slide = blank_slide(prs)
    header(slide, "02   合作从哪里来", "对方因《杨浦一百》主动上门")
    footer(slide, 3)
    card(slide, 0.5, 1.32, 6.15, 5.45, NAVY, NAVY)
    write(slide, [t("对方原话", 12, "E4C98A", True)], 0.78, 1.55, 5.5, 0.28)
    write(
        slide,
        [
            t(
                "“我就是看到你们那个报告，包括里面有我们好几家企业……想的有没有这个机会点？”",
                20,
                WHITE,
                True,
                spacing=1.2,
            )
        ],
        0.78,
        2.05,
        5.55,
        1.7,
    )
    write(
        slide,
        [
            t("这份白皮书已经做成了三件事", 14, "E4C98A", True, after=8),
            t("让园区自己找到联合会。", 15, WHITE, before=6),
            t("让对方相信联合会掌握可核验的企业信息。", 15, WHITE, before=4),
            t("让合作从对方的问题开始，而不是从一次拜访开始。", 15, WHITE, before=4),
        ],
        0.78,
        4.0,
        5.55,
        2.3,
    )
    signals = [
        ("约 00:12", "承接办公需求，并愿带企业去川渝考察"),
        ("约 00:16", "小型活动空间可以提供给联合会使用"),
        ("约 01:02", "直接问：能不能帮忙在这里挂个牌"),
        ("约 01:07", "活动、选址有结果后，可以谈费用"),
        ("约 01:29", "主动询问入会方式和会费"),
    ]
    write(slide, [t("五处明确表态", 14, NAVY, True)], 6.9, 1.32, 5.8, 0.32)
    for index, (when, what) in enumerate(signals):
        y = 1.78 + index * 0.95
        fill = GOLD_SOFT if index in (2, 3) else CARD
        card(slide, 6.9, y, 5.9, 0.84, fill)
        write(slide, [t(when, 12, GOLD, True)], 7.1, y + 0.1, 1.5, 0.26)
        write(slide, [t(what, 14, INK, True if index in (2, 3) else False)], 7.1, y + 0.38, 5.5, 0.36)
    notes(
        slide,
        "01:02 和 01:07 商业含义最强：前者说明需要产业平台，后者说明对方不是只想免费拿资源。"
        "有付费意愿还不等于预算已批，也不等于接受某个报价。",
    )


def slide_profile(prs):
    slide = blank_slide(prs)
    header(slide, "03   对方经营画像", "上海项目要先解决现金流")
    footer(slide, 4)
    facts = [
        ("成都载体", "高新区、温江、新津\n合计约80万平方米"),
        ("上海口径", "洽谈约6000平方米\n官网约1.67万平方米"),
        ("在园企业", "约37家\n多家已入《杨浦一百》"),
        ("出租率", "约七成多\n今年退租后正在回升"),
        ("招商配置", "团队约8人\n其中招商约3人"),
        ("租金水平", "约2元/平方米/天\n不含物业费"),
        ("固定成本", "物业费每月20元\n停车位月租700元"),
        ("资本现状", "总部3只基金\n上海接触偏中后期"),
    ]
    for index, (label, value) in enumerate(facts):
        col, row = index % 4, index // 4
        x, y = 0.5 + col * 3.2, 1.32 + row * 1.55
        card(slide, x, y, 3.05, 1.42)
        write(slide, [t(label, 13, TEAL, True)], x + 0.16, y + 0.16, 2.7, 0.28)
        write(slide, [t(value, 14, INK, spacing=1.12)], x + 0.16, y + 0.52, 2.72, 0.74)
    card(slide, 0.5, 4.55, 12.33, 2.25, GOLD_SOFT, "E6D7B4")
    write(slide, [t("报价前先要一页房源单", 15, GOLD, True)], 0.74, 4.72, 11.8, 0.3)
    write(
        slide,
        [
            t(
                "同一栋楼售予多个小业主，物业在开发商手里，空置也要付物业费。成交服务费只覆盖对方书面确认、由其掌控的房源。",
                14,
                INK,
                before=2,
            ),
            t(
                "若按1.67万平方米、约三成空置估算，空置物业费约120万元/年。此为测算，不是报价。\n园内另有约1元/平方米/天的抛租，录音称相关合同大约到2027年6月。",
                14,
                INK,
                before=6,
            ),
        ],
        0.74,
        5.1,
        11.85,
        1.5,
    )
    notes(
        slide,
        "面积存在两个口径，签约前必须让对方列出可控套间、空置、对外报价，以及上海团队可以当场决定的租金区间。"
        "单一民营股东决策灵活，但仍需总部书面授权价格带，避免每单回成都打报告。",
    )


def slide_needs(prs):
    slide = blank_slide(prs)
    header(slide, "04   需求判断", "方案要回答园区的经营问题")
    footer(slide, 5)
    card(slide, 0.5, 1.32, 6.15, 5.45)
    rect(slide, 0.5, 1.32, 0.08, 5.45, TEAL)
    write(slide, [t("对方会认真听的四件事", 16, TEAL, True)], 0.82, 1.52, 5.5, 0.34)
    wants = [
        ("出租率与现金流", "空置仍要付物业费。活动如果不能带来企业和租金，很难单独立项。"),
        ("匹配阶段的资本", "中后期机构帮不上入园企业。晶丰新感长大后被免费场地加投资挖走。"),
        ("能持续来人的牌子", "上海项目知名度弱，区位优势主要是离复旦近。对方当场问过挂牌。"),
        ("长三角桥头堡", "总部希望这边的企业拓展西南时先想到新谷，而不是只做一栋楼的出租。"),
    ]
    for index, (title, body) in enumerate(wants):
        y = 2.05 + index * 1.1
        write(slide, [t(title, 15, NAVY, True)], 0.82, y, 5.5, 0.3)
        write(slide, [t(body, 13, MUTED, spacing=1.08)], 0.82, y + 0.32, 5.5, 0.66)
    card(slide, 6.85, 1.32, 5.98, 5.45, RED_SOFT, "E7CFCF")
    write(slide, [t("第一份协议不纳入", 16, RED, True)], 7.12, 1.52, 5.5, 0.34)
    stops = [
        "税收分成或返税平台。双方都认为政策会收回。",
        "会员费覆盖无限次招商、策划和融资服务。",
        "未获学校审批就挂“复旦大学××中心”。",
        "把成都总部企业作为导入这栋楼的主方案。",
        "无产权、规划和预算确认时做改造研究。",
        "合同里承诺一定招满若干平方米。",
    ]
    for index, item in enumerate(stops):
        write(slide, [t(f"{index + 1}    {item}", 14, INK, spacing=1.05)], 7.12, 2.1 + index * 0.7, 5.45, 0.6)


def slide_map(prs):
    slide = blank_slide(prs)
    header(slide, "05   机会取舍", "十个方向收成一个首期项目")
    footer(slide, 6)
    card(slide, 0.5, 1.3, 7.55, 5.48)
    write(slide, [t("首期打包推进", 16, TEAL, True)], 0.74, 1.48, 7, 0.32)
    write(
        slide,
        [t("下面五项是一个项目的五个部件，不分别向对方铺开。", 13, MUTED)],
        0.74,
        1.88,
        7,
        0.28,
    )
    items = [
        ("01", "副会长单位", "建立正式关系，年费对应章程权益"),
        ("02", "三个月招商试点", "精准对接科技企业，结果另计服务费"),
        ("03", "联合服务站", "挂牌、沙龙、接待放在同一节点"),
        ("04", "早期资本对接", "先做在园企业融资需求，再请机构"),
        ("05", "小型产业活动", "园区出场地，定制场次单独报价"),
    ]
    for index, (num, title, desc) in enumerate(items):
        y = 2.32 + index * 0.84
        write(slide, [t(num, 16, GOLD, True)], 0.78, y, 0.7, 0.36)
        write(slide, [t(title, 16, NAVY, True)], 1.55, y, 2.5, 0.32)
        write(slide, [t(desc, 13, MUTED)], 4.05, y + 0.02, 3.7, 0.36)
    layers = [
        (TEAL_SOFT, "C9DDD9", TEAL, "平行核实", "新加坡选址。两周内约对方出海负责人和陆总，先问清是孵化器、办公室还是服务平台。"),
        (GOLD_SOFT, "E6D7B4", GOLD, "第二阶段", "双城考察、核聚变与机器人生态、专题研究、资产盘活。都要有预算和真实需求再立项。"),
        (RED_SOFT, "E7CFCF", RED, "不作为产品", "税收分成、未授权复旦冠名、保证招满、免费完整白皮书、无预算的改公寓方案。"),
    ]
    for index, (fill, line, color, title, body) in enumerate(layers):
        y = 1.3 + index * 1.86
        card(slide, 8.22, y, 4.6, 1.74, fill, line)
        write(slide, [t(title, 15, color, True)], 8.42, y + 0.16, 4.2, 0.3)
        write(slide, [t(body, 13, INK, spacing=1.08)], 8.42, y + 0.54, 4.2, 1.05)


def slide_offer(prs):
    slide = blank_slide(prs)
    header(slide, "06   首期方案", "先报一个对方能审批的方案")
    footer(slide, 7)
    offers = [
        ("1", "副会长单位", "1 万元 / 年", "以上海公司为入会主体，按章程和审批程序办理。权益是交流、展示和资讯，不含专项招商。"),
        ("2", "三个月试点", "2 万元", "招商需求梳理、首批科技企业与早期机构对接、一场小型产业交流。谈判区间可到 2–3 万元。"),
        ("3", "结果服务费", "0.5–1 个月租金", "企业签约入驻后另计，以基础租金为协商起点。大型或特殊项目另行议价。"),
    ]
    for index, (num, title, price, body) in enumerate(offers):
        x = 0.5 + index * 4.2
        card(slide, x, 1.32, 4.02, 3.35)
        write(slide, [t(num, 20, GOLD, True)], x + 0.22, 1.5, 0.6, 0.36)
        write(slide, [t(title, 18, NAVY, True)], x + 0.22, 2.0, 3.55, 0.36)
        write(slide, [t(price, 20, TEAL, True)], x + 0.22, 2.48, 3.55, 0.4)
        write(slide, [t(body, 13, INK, spacing=1.12)], x + 0.22, 3.05, 3.58, 1.4)
    card(slide, 0.5, 4.84, 12.33, 1.95, NAVY, NAVY)
    write(slide, [t("若对方接受建议报价", 14, "E4C98A", True)], 0.76, 5.04, 11.8, 0.28)
    write(
        slide,
        [
            t("首期基础合作约 3 万元，另保留招商结果收益。对方若暂批不下专项预算，至少先锁定入会和结果付费机制。", 16, WHITE, before=6),
            t("同一项工作只进一个收费科目：会员费里的交流，不再收活动费；试点已含的那场活动，不再单收策划费。", 14, "D5DEE6", before=8),
        ],
        0.76,
        5.4,
        11.8,
        1.2,
    )
    notes(
        slide,
        "所有金额都是谈判起点。1万元会员费容易把后续服务做穿，所以专项必须另签。"
        "试点费取推荐值2万元，保留2–3万元的谈判空间。",
    )


def slide_leasing(prs):
    slide = blank_slide(prs)
    header(slide, "07   招商试点", "做企业成长服务，不做普通中介")
    footer(slide, 8)
    write(
        slide,
        [t("对方已有招商团队，市场上也不缺中介。壁垒来自企业库、投资机构和对接过程。", 14, MUTED)],
        0.52,
        1.24,
        12.2,
        0.32,
    )
    blocks = [
        ("科技企业", "引荐确有空间需求、并且付得起租金的人工智能和硬科技企业。", "办公空间与书面授权的租金区间"),
        ("投资机构", "带早期机构及其被投企业到现场。中后期资金不作为主引擎。", "落地条件、注册与园区服务"),
        ("产业活动", "一场20–40人的主题交流，会后只登记到现场的带看。", "场地、接待和在园企业"),
    ]
    for index, (title, ours, theirs) in enumerate(blocks):
        x = 0.5 + index * 4.2
        card(slide, x, 1.68, 4.02, 2.85)
        write(slide, [t(title, 16, NAVY, True)], x + 0.2, 1.84, 3.6, 0.32)
        write(slide, [t("联合会", 12, TEAL, True)], x + 0.2, 2.28, 3.6, 0.24)
        write(slide, [t(ours, 13, INK, spacing=1.08)], x + 0.2, 2.54, 3.62, 0.85)
        write(slide, [t("新谷湾", 12, GOLD, True)], x + 0.2, 3.42, 3.6, 0.24)
        write(slide, [t(theirs, 13, INK, spacing=1.08)], x + 0.2, 3.68, 3.62, 0.68)
    metrics = [
        ("合同承诺", "过程与台账"),
        ("复盘参照", "90天看线索质量"),
        ("带看口径", "到现场才计数"),
        ("租金授权", "先拿书面价格带"),
    ]
    for index, (label, value) in enumerate(metrics):
        x = 0.5 + index * 3.2
        card(slide, x, 4.7, 3.05, 1.15, TEAL_SOFT, "C9DDD9")
        write(slide, [t(label, 12, TEAL, True)], x + 0.16, 4.84, 2.75, 0.24)
        write(slide, [t(value, 14, NAVY, True)], x + 0.16, 5.16, 2.75, 0.48)
    write(
        slide,
        [t("行业带看成交比约110:1。宝龙样本为数月约1000至1600平方米。内部复盘看90天线索质量，合同不保证面积。", 13, MUTED)],
        0.52,
        6.05,
        12.3,
        0.7,
    )
    notes(
        slide,
        "内部可以用800至1200平方米、两单可签约面积做复盘参照，因为区位弱于滨江宝龙。"
        "早期企业可以推荐，但需要投资方协同或租金支付能力；传统零售沿用对方自己的租金梯度筛选。",
    )


def slide_station(prs):
    slide = blank_slide(prs)
    header(slide, "08   联合服务站", "把挂牌做成长期服务节点")
    footer(slide, 9)
    card(slide, 0.5, 1.3, 12.33, 1.45, NAVY, NAVY)
    write(slide, [t("推荐名称", 12, "E4C98A", True)], 0.74, 1.46, 3, 0.24)
    write(
        slide,
        [t("杨浦科技企业成长与产业协同服务站", 22, WHITE, True)],
        0.74,
        1.78,
        11.8,
        0.42,
    )
    write(
        slide,
        [t("设在新谷国际，先试运行一年。名称覆盖人工智能、机器人、先进制造和跨区域协同，不把运营责任做成完整孵化器。", 13, "D5DEE6")],
        0.74,
        2.26,
        11.8,
        0.34,
    )
    events = [
        ("杨浦 AI 产业下午茶", "20–30人", "每月或双月", "发现企业需求"),
        ("科技企业融资闭门会", "10–15家", "每季度", "融资与选址并行"),
        ("未来产业创新参访", "小范围", "每季度", "连接会员与潜在客户"),
    ]
    for index, (name, scale, freq, goal) in enumerate(events):
        x = 0.5 + index * 4.2
        card(slide, x, 2.95, 4.02, 2.15)
        write(slide, [t(name, 15, NAVY, True)], x + 0.18, 3.1, 3.66, 0.4)
        write(slide, [t(f"{scale}   ·   {freq}", 13, TEAL, True)], x + 0.18, 3.58, 3.66, 0.3)
        write(slide, [t(goal, 14, INK)], x + 0.18, 4.05, 3.66, 0.4)
    card(slide, 0.5, 5.28, 12.33, 1.52, GOLD_SOFT, "E6D7B4")
    write(slide, [t("两条约定", 14, GOLD, True)], 0.74, 5.44, 11.8, 0.26)
    write(
        slide,
        [
            t("场地由对方提供，只适合中小型活动。联合会的策划、邀约和组织，基础场次可以资源共建，定制招商和资本专场单独报价。", 14, INK, before=4),
            t("涉及联合会或复旦大学的名称，须取得相应主体授权。引入复旦资源，先走专家交流和合规的产学研合作。", 14, INK, before=4),
        ],
        0.74,
        5.74,
        11.85,
        0.92,
    )


def slide_capital(prs):
    slide = blank_slide(prs)
    header(slide, "09   早期资本", "从在园企业的融资需求做起")
    footer(slide, 10)
    write(
        slide,
        [t("对方有早期企业，缺少与阶段匹配的资本服务；联合会能接触机构和项目。第一阶段不设基金，也不承诺融资成功。", 14, MUTED)],
        0.52,
        1.24,
        12.3,
        0.4,
    )
    steps = [
        ("1", "筛企业", "从约37家在园企业中选出有融资需求的5–8家，形成需求清单。"),
        ("2", "核阶段", "只请投资策略对得上的早期机构。先核对，再发会议邀请。"),
        ("3", "做闭门会", "“杨浦未来产业投融资闭门对接会”，机构、创始人、园区同场。"),
        ("4", "接空间", "机构看中项目时，园区同步给出办公和成长服务。"),
    ]
    for index, (num, title, body) in enumerate(steps):
        x = 0.5 + index * 3.2
        card(slide, x, 1.78, 3.05, 2.55)
        write(slide, [t(num, 18, GOLD, True)], x + 0.16, 1.94, 2.7, 0.32)
        write(slide, [t(title, 16, NAVY, True)], x + 0.16, 2.34, 2.7, 0.32)
        write(slide, [t(body, 13, INK, spacing=1.1)], x + 0.16, 2.78, 2.72, 1.35)
    card(slide, 0.5, 4.52, 12.33, 2.25)
    write(slide, [t("一条可以重复使用的链条", 15, TEAL, True)], 0.74, 4.7, 11.8, 0.3)
    chain = ["企业库识别需求", "机构匹配阶段", "园区提供空间", "落地后持续服务"]
    for index, label in enumerate(chain):
        x = 0.74 + index * 3.05
        rect(slide, x, 5.18, 2.7, 0.55, TEAL if index % 2 == 0 else NAVY)
        write(slide, [t(label, 13, WHITE, True, align=PP_ALIGN.CENTER)], x, 5.32, 2.7, 0.3)
    write(
        slide,
        [t("普通对接和活动可以约定服务费。融资交易、资金募集和投资顾问，由具备相应资质的主体依法开展。", 13, MUTED)],
        0.74,
        5.9,
        11.8,
        0.55,
    )
    notes(
        slide,
        "录音里的企业名称（东升相关主体、华数科技、华太激光、晶丰新感等）是转写，外发材料前核对工商登记。"
        "东升巨变的重资产选址在对接市级部门，不默认纳入这次园区招商。产业方向不等于对方愿意投资该产业。",
    )


def slide_singapore(prs):
    slide = blank_slide(prs)
    header(slide, "10   平行工作线", "新加坡两周内摸清，先不对外挂名")
    footer(slide, 11)
    card(slide, 0.5, 1.3, 6.2, 5.48)
    write(slide, [t("这次会要问清的三件事", 16, TEAL, True)], 0.74, 1.5, 5.7, 0.34)
    questions = [
        ("选址用途", "筹建孵化器、租赁办公室，还是设立产业服务平台。三种用途对应完全不同的合作。"),
        ("决策人", "约对方出海负责人和陆总。上海招商团队不能代替这条线拍板。"),
        ("企业名单", "哪些会员企业需要新加坡主体或跨境业务，并可以在杨浦保留联络点。"),
    ]
    for index, (title, body) in enumerate(questions):
        y = 2.1 + index * 1.45
        write(slide, [t(title, 16, NAVY, True)], 0.74, y, 5.6, 0.32)
        write(slide, [t(body, 14, INK, spacing=1.1)], 0.74, y + 0.4, 5.65, 0.9)
    card(slide, 6.9, 1.3, 5.93, 5.48, GOLD_SOFT, "E6D7B4")
    write(slide, [t("这条线现在做到这里", 16, GOLD, True)], 7.14, 1.5, 5.45, 0.34)
    limits = [
        "对方原话是“刚布的新加坡，正处于选址阶段”。",
        "联合会有约两年的新加坡跨境业务经验，2026年另有澳大利亚安排。经验可以交流，安排本身不自动成为联合项目。",
        "用途未明之前，对外材料不使用“联合海外中心”。",
        "波士顿平台只做办公时间的资源对接，不在首期单独立项。",
        "收费等项目范围清楚后按项目报价。",
    ]
    for index, item in enumerate(limits):
        write(slide, [t(item, 14, INK, spacing=1.08)], 7.14, 2.1 + index * 0.85, 5.45, 0.8)


def slide_phase2(prs):
    slide = blank_slide(prs)
    header(slide, "11   第二阶段", "有需求、有预算之后再展开")
    footer(slide, 12)
    cards = [
        ("双城考察", "“长三角—成渝科技产业协同”先做小团。8–12家企业，行前写明市场、厂房、伙伴或当地基金需求。对方已有台州、海宁、绍兴等通道，且不承接外地楼宇代运营。"),
        ("未来产业", "核聚变、具身智能和机器人可以做走访与产业链梳理，并与联合会已有研究衔接。园区里有相关企业，不等于集团会投资该方向。"),
        ("研究产品", "《杨浦一百》继续作为获客工具。对方园区可先成为一个更新案例。独立课题，例如跨区域发展观察，有经费再做，不免费交付白皮书。"),
        ("资产盘活", "租金、小面积产品和改公寓都谈过，对方也提到改造预算有限。产权、规划和预算未确认前，不做完整改造研究。政策口径以主管部门书面意见为准。"),
    ]
    for index, (title, body) in enumerate(cards):
        col, row = index % 2, index // 2
        x, y = 0.5 + col * 6.4, 1.32 + row * 2.7
        card(slide, x, y, 6.2, 2.52)
        rect(slide, x, y, 0.08, 2.52, TEAL if row == 0 else GOLD)
        write(slide, [t(title, 18, NAVY, True)], x + 0.32, y + 0.22, 5.6, 0.38)
        write(slide, [t(body, 14, INK, spacing=1.12)], x + 0.32, y + 0.72, 5.6, 1.6)


def slide_pricing(prs):
    slide = blank_slide(prs)
    header(slide, "12   收费边界", "同一项工作只收一次费")
    footer(slide, 13)
    columns = [2.3, 2.7, 5.5, 1.85]
    header_row = ["合作产品", "谈判起点", "服务边界", "时点"]
    data = [
        ["副会长单位", "1万元 / 年", "章程权益；不含专项招商和融资", "立刻"],
        ["三个月试点", "2万–3万元", "需求梳理、精准对接、一场小型活动", "立刻"],
        ["招商成果", "0.5–1个月基础租金", "只计书面确认的可控房源", "签约后"],
        ["定制活动", "0.8万–2万元 / 场", "试点未包含的闭门会和专场", "按场"],
        ["服务站", "按人员和场次", "挂牌不绑定无限运营", "试运行"],
        ["成都考察", "按项目", "只组织有书面需求的小团", "第二阶段"],
        ["研究报告", "按课题", "不免费做完整白皮书", "第二阶段"],
        ["海外对接", "按项目", "选址用途明确后再报价", "条件成熟"],
    ]
    x0, y0 = 0.48, 1.28
    height = 0.5
    cursor_x = x0
    for width, label in zip(columns, header_row):
        rect(slide, cursor_x, y0, width, height, NAVY)
        write(slide, [t(label, 12, WHITE, True)], cursor_x + 0.1, y0 + 0.12, width - 0.16, 0.28)
        cursor_x += width
    for r, row in enumerate(data):
        y = y0 + height + r * 0.48
        fill = WHITE if r % 2 == 0 else PAPER
        cursor_x = x0
        for width, label in zip(columns, row):
            rect(slide, cursor_x, y, width, 0.48, fill, LINE)
            write(slide, [t(label, 12, INK, bold=(cursor_x == x0))], cursor_x + 0.1, y + 0.1, width - 0.16, 0.3)
            cursor_x += width
    write(
        slide,
        [t("金额均为内部谈判起点。首期推荐组合是会员1万元加试点2万元；试点若已含一场活动，该场不再另收基础策划费。", 13, MUTED)],
        0.5,
        6.55,
        12.3,
        0.4,
    )


def slide_timeline(prs):
    slide = blank_slide(prs)
    header(slide, "13   九十天", "从建立关系到明确项目")
    footer(slide, 14)
    rows = [
        ("第 1–10 天", "确定负责人，取得房源单、决策权限和预算口径", "合作需求清单"),
        ("第 11–20 天", "确定入会主体、试点范围和收费科目", "意向书或专项协议"),
        ("第 21–40 天", "对接首批科技企业与早期机构，安排到场参观", "需求与对接台账"),
        ("第 41–60 天", "举办第一场产业或投融资闭门活动", "活动记录与具体线索"),
        ("第 61–75 天", "跟进选址和融资，只推进有下一步的对象", "可核实的洽谈进展"),
        ("第 76–90 天", "复盘过程成果，决定是否进入年度合作", "阶段报告与下期方案"),
    ]
    for index, (when, task, output) in enumerate(rows):
        y = 1.28 + index * 0.9
        card(slide, 0.5, y, 12.33, 0.8)
        rect(slide, 0.5, y, 0.08, 0.8, TEAL if index < 2 else GOLD)
        write(slide, [t(when, 14, NAVY, True)], 0.78, y + 0.22, 2.1, 0.36)
        write(slide, [t(task, 14, INK)], 3.0, y + 0.22, 6.3, 0.4)
        write(slide, [t(output, 13, TEAL, True, align=PP_ALIGN.RIGHT)], 9.2, y + 0.22, 3.35, 0.36)
    notes(
        slide,
        "新加坡沟通放在第1–10天并行启动，不占用招商试点的主节奏。"
        "考核放在负责人、房源、预算、到场台账和一场闭门活动，不考核承诺的招租面积。",
    )


def slide_message(prs):
    slide = blank_slide(prs)
    header(slide, "14   对外口径", "会后可以直接发出的微信")
    footer(slide, 15)
    write(slide, [t("发送前核对称呼。录音中出现赵总、夏总，以对方确认的姓名为准。", 13, MUTED)], 0.52, 1.22, 12, 0.28)
    card(slide, 0.5, 1.58, 12.33, 5.22)
    paragraphs = [
        t("赵总、夏总，你们好。", 14, INK, before=0, after=8),
        t("上次交流后，我们把双方资源又对了一遍。建议先落一个可控的试点，不把多条线同时铺开。", 14, INK, after=8),
        t("一是产业招商与资本对接。围绕贵司上海园区的人工智能和科技企业定位，由联合会引荐确有空间或融资需求的企业、以及阶段匹配的早期投资机构，到现场交流。活动和选址形成结果后，按专项约定结算。", 14, INK, after=8),
        t("二是联合服务站。上次提到的挂牌，建议做成“科技企业成长与产业协同服务站”，放在贵司园区，先做小型沙龙和企业对接。", 14, INK, after=8),
        t("三是双城与海外留作下一步。成都载体适合有明确考察需求的企业小团走访；新加坡业务建议先由双方具体负责人做一次专题沟通，摸清选址用途后再谈合作方式。", 14, INK, after=8),
        t("副会长单位可以同步办理。试点前希望先确认贵司可控房源、上海团队能当场决定的租金区间，以及试点经费的审批口径。这样双方投入可控，也更容易在这个季度形成成果。", 14, INK),
    ]
    write(slide, paragraphs, 0.74, 1.76, 11.85, 4.9)


def slide_decision(prs):
    slide = blank_slide(prs)
    header(slide, "15   请内部拍板", "这周确定三件事即可启动")
    footer(slide, 16)
    decisions = [
        ("入会主体", "是否以上海新谷湾公司申请副会长，并按章程走审批。"),
        ("试点报价", "是否以1万元会费加2万元试点费作为正式谈判起点。"),
        ("服务站名称", "是否采用“科技企业成长与产业协同服务站”，一年试运行。"),
    ]
    for index, (title, body) in enumerate(decisions):
        x = 0.5 + index * 4.2
        card(slide, x, 1.32, 4.02, 2.15)
        write(slide, [t(f"0{index + 1}", 14, GOLD, True)], x + 0.2, 1.48, 3.6, 0.26)
        write(slide, [t(title, 18, NAVY, True)], x + 0.2, 1.82, 3.6, 0.36)
        write(slide, [t(body, 14, INK, spacing=1.1)], x + 0.2, 2.3, 3.62, 0.95)
    card(slide, 0.5, 3.66, 12.33, 3.1, NAVY, NAVY)
    write(slide, [t("更长的用法", 14, "E4C98A", True)], 0.76, 3.86, 11.8, 0.28)
    write(
        slide,
        [
            t("宝龙、创智汇之后，又有园区因为《杨浦一百》主动上门。零散的园区合作可以收成一条可重复的服务链。", 16, WHITE, before=8, spacing=1.12),
            t("科技企业数据库  →  成长需求  →  投资机构匹配  →  产业空间匹配  →  跨区域落地  →  持续服务", 15, "E4C98A", before=10),
            t("新谷湾是这条链上优先级较高的一环。成都载体、总部基金、新加坡和波士顿，留到试点跑通后再逐项验证。", 14, "D5DEE6", before=10),
        ],
        0.76,
        4.2,
        11.85,
        2.3,
    )


def slide_appendix(prs):
    slide = blank_slide(prs)
    header(slide, "附录   使用说明", "这页留给汇报人，不对外转发")
    footer(slide, 17)
    points = [
        ("依据", "2026-09-24 录音及双方当场表态。公开载体信息仅用于核对上海新谷湾、新谷国际与湾谷科技园的主体关系。"),
        ("口径", "出租率、租金、人员、基金和融资情节均来自洽谈，尚未独立尽调。企业名称经过语音转写，外发前核对工商名。"),
        ("执行会长", "录音记为“志刚”。公开活动报道中，复旦大学住房政策研究中心秘书长、联合会执行会长为胡继刚。发送材料前与本人确认。"),
        ("金额", "全部为内部谈判起点，不是对方预算，也不是市场统一定价。"),
        ("承诺", "联合会对接过程、活动组织和台账负责，不对最终租赁或融资成功作无条件保证。"),
    ]
    for index, (title, body) in enumerate(points):
        y = 1.28 + index * 1.08
        write(slide, [t(title, 14, TEAL, True)], 0.55, y, 1.6, 0.3)
        write(slide, [t(body, 14, INK, spacing=1.08)], 2.25, y, 10.4, 0.9)


def build():
    presentation = Presentation()
    presentation.slide_width = Inches(W)
    presentation.slide_height = Inches(H)
    presentation.core_properties.title = "杨浦科技企业成长与产业协同合作计划"
    presentation.core_properties.subject = "成都新谷投资集团（上海新谷湾）与杨浦区科技企业联合会内部合作建议稿"
    presentation.core_properties.author = "上海市杨浦区科技企业联合会"
    presentation.core_properties.category = "内部建议稿"
    for builder in (
        slide_cover,
        slide_judgment,
        slide_signal,
        slide_profile,
        slide_needs,
        slide_map,
        slide_offer,
        slide_leasing,
        slide_station,
        slide_capital,
        slide_singapore,
        slide_phase2,
        slide_pricing,
        slide_timeline,
        slide_message,
        slide_decision,
        slide_appendix,
    ):
        builder(presentation)
    if len(presentation.slides) != TOTAL:
        raise SystemExit(f"页数不符：{len(presentation.slides)} != {TOTAL}")
    output = "/workspace/output/杨浦新谷湾合作建议稿-2026.pptx"
    presentation.save(output)
    print(output)


if __name__ == "__main__":
    build()
