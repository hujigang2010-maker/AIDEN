#!/usr/bin/env python3
"""生成九龙坡城市更新项目联合投标工作方案 Word 文件。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


SONG = "宋体"
HEI = "黑体"
KAI = "楷体"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NAVY_HEX = "1F4E79"
LABEL_HEX = "E7EEF5"
ALT_HEX = "F7F9FC"
WARN_HEX = "FDECEC"
NOTE_HEX = "FFF8E7"
LINE_HEX = "1F4E79"

PAGE_WIDTH_CM = 16.0


def _set_run_font(
    run,
    *,
    name: str = SONG,
    size: float = 12,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def set_keep_next(paragraph: Paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag == qn("w:keepNext"):
            pPr.remove(child)
    pPr.append(OxmlElement("w:keepNext"))


def set_keep_lines(paragraph: Paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag == qn("w:keepLines"):
            pPr.remove(child)
    pPr.append(OxmlElement("w:keepLines"))


def _set_paragraph_format(
    paragraph: Paragraph,
    *,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before: float = 0,
    space_after: float = 4,
    line_spacing: float = 1.32,
    first_line_indent: float | None = None,
    left_indent: float | None = None,
    keep_with_next: bool = False,
) -> None:
    pf = paragraph.paragraph_format
    paragraph.alignment = alignment
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)
    set_keep_lines(paragraph)
    if keep_with_next:
        set_keep_next(paragraph)


def add_run_text(
    paragraph: Paragraph,
    text: str,
    *,
    name: str = SONG,
    size: float = 12,
    bold: bool = False,
    color: RGBColor | None = None,
):
    run = paragraph.add_run(text)
    _set_run_font(run, name=name, size=size, bold=bold, color=color)
    return run


def add_para(
    doc: Document,
    text: str = "",
    *,
    name: str = SONG,
    size: float = 12,
    bold: bool = False,
    color: RGBColor | None = None,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before: float = 0,
    space_after: float = 4,
    line_spacing: float = 1.32,
    first_line_indent: float | None = None,
    left_indent: float | None = None,
    keep_with_next: bool = False,
) -> Paragraph:
    p = doc.add_paragraph()
    _set_paragraph_format(
        p,
        alignment=alignment,
        space_before=space_before,
        space_after=space_after,
        line_spacing=line_spacing,
        first_line_indent=first_line_indent,
        left_indent=left_indent,
        keep_with_next=keep_with_next,
    )
    if text:
        add_run_text(p, text, name=name, size=size, bold=bold, color=color)
    return p


def add_section_heading(doc: Document, text: str) -> Paragraph:
    p = add_para(
        doc,
        text,
        name=HEI,
        size=14,
        bold=True,
        color=NAVY,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=10,
        space_after=6,
        line_spacing=1.25,
        keep_with_next=True,
    )
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), NAVY_HEX)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_subheading(doc: Document, text: str) -> Paragraph:
    return add_para(
        doc,
        text,
        name=HEI,
        size=12,
        bold=True,
        color=NAVY,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=8,
        space_after=3,
        line_spacing=1.25,
        keep_with_next=True,
    )


def set_cell_shading(cell: _Cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell: _Cell, top: int = 50, bottom: int = 50, left: int = 70, right: int = 70) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcMar"):
            tcPr.remove(child)
    tcMar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table: Table, color: str = "8FA4BC", sz: str = "8") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_table_width(table: Table, width_cm: float = PAGE_WIDTH_CM) -> None:
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    twips = int(width_cm * 567)
    tblW.set(qn("w:w"), str(twips))
    tblW.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    for child in list(trPr):
        if child.tag == qn("w:tblHeader"):
            trPr.remove(child)
    trPr.append(OxmlElement("w:tblHeader"))


def keep_table_together(table: Table) -> None:
    rows = list(table.rows)
    for row in rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        keep_next = row != rows[-1]
        for cell in row.cells:
            for p in cell.paragraphs:
                if keep_next:
                    set_keep_next(p)


def set_col_widths(table: Table, widths_cm: list[float]) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is None:
        tblGrid = OxmlElement("w:tblGrid")
        tblPr.addnext(tblGrid)
    else:
        for child in list(tblGrid):
            tblGrid.remove(child)
    for width in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 567)))
        tblGrid.append(grid_col)
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
            tc = row.cells[idx]._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 567)))
            tcW.set(qn("w:type"), "dxa")


def fill_cell(
    cell: _Cell,
    text: str,
    *,
    name: str = SONG,
    size: float = 10.5,
    bold: bool = False,
    color: RGBColor | None = None,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    fill: str | None = None,
    compact: bool = False,
) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if compact:
        set_cell_margins(cell, top=36, bottom=36, left=50, right=50)
    else:
        set_cell_margins(cell)
    if fill:
        set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    _set_paragraph_format(
        p,
        alignment=alignment,
        space_before=0,
        space_after=0,
        line_spacing=1.2,
    )
    add_run_text(p, text, name=name, size=size, bold=bold, color=color or INK)


def add_kv_table(doc: Document, rows: list[tuple[str, str, dict]]) -> Table:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_table_borders(table, color="8FA4BC", sz="4")
    set_col_widths(table, [4.2, 11.8])
    for i, (label, value, opts) in enumerate(rows):
        label_fill = NAVY_HEX if opts.get("alert") else LABEL_HEX
        value_fill = WARN_HEX if opts.get("alert") else (ALT_HEX if i % 2 else "FFFFFF")
        label_color = WHITE if opts.get("alert") else NAVY
        value_color = RED if opts.get("alert") or opts.get("emphasis") else INK
        fill_cell(
            table.rows[i].cells[0],
            label,
            name=HEI,
            size=10.5,
            bold=True,
            color=label_color,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            fill=label_fill,
        )
        fill_cell(
            table.rows[i].cells[1],
            value,
            name=SONG,
            size=11,
            bold=bool(opts.get("alert") or opts.get("emphasis")),
            color=value_color,
            fill=value_fill,
        )
    keep_table_together(table)
    return table


def add_grid_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[float],
    *,
    header_fill: str = NAVY_HEX,
    font_size: float = 10.5,
    together: bool = True,
    compact: bool = False,
) -> Table:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_table_borders(table, color="8FA4BC", sz="4")
    set_col_widths(table, col_widths)
    for i, header in enumerate(headers):
        fill_cell(
            table.rows[0].cells[i],
            header,
            name=HEI,
            size=font_size,
            bold=True,
            color=WHITE,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            fill=header_fill,
            compact=compact,
        )
    set_repeat_header(table.rows[0])
    center_headers = {"序号", "完成情况", "电话", "联系人", "时间", "牵头", "主责", "重大总院", "复旦中心"}
    for r_i, row in enumerate(rows):
        fill = ALT_HEX if r_i % 2 else "FFFFFF"
        for c_i, value in enumerate(row):
            header = headers[c_i]
            if header in center_headers or "★" in value:
                align = WD_ALIGN_PARAGRAPH.CENTER
            else:
                align = WD_ALIGN_PARAGRAPH.LEFT
            is_alert = "18:00" in value or "14:30" in value
            lead_fudan = header == "牵头" and value.startswith("复旦")
            fill_cell(
                table.rows[r_i + 1].cells[c_i],
                value,
                name=SONG,
                size=font_size,
                bold=is_alert or lead_fudan,
                color=RED if is_alert else (NAVY if lead_fudan else INK),
                alignment=align,
                fill=WARN_HEX if is_alert else fill,
                compact=compact,
            )
    if together:
        keep_table_together(table)
    else:
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return table


def add_callout(doc: Document, title: str, body: str, *, fill: str = WARN_HEX) -> Table:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_table_borders(table, color=LINE_HEX, sz="12")
    cell = table.rows[0].cells[0]
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    p1 = cell.paragraphs[0]
    _set_paragraph_format(p1, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2, line_spacing=1.2)
    add_run_text(p1, title, name=HEI, size=11, bold=True, color=RED)

    p2 = cell.add_paragraph()
    _set_paragraph_format(p2, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=0, line_spacing=1.3)
    add_run_text(p2, body, name=SONG, size=11, color=INK)
    keep_table_together(table)
    return table


def add_quote(doc: Document, text: str) -> Table:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_table_borders(table, color="C9A227", sz="8")
    cell = table.rows[0].cells[0]
    cell.text = ""
    set_cell_shading(cell, NOTE_HEX)
    set_cell_margins(cell, top=70, bottom=70, left=110, right=110)
    p = cell.paragraphs[0]
    _set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=0, line_spacing=1.3)
    add_run_text(p, text, name=KAI, size=11, color=INK)
    keep_table_together(table)
    return table


def add_bullets(doc: Document, items: list[str], *, indent: float = 0.4) -> None:
    for idx, item in enumerate(items):
        p = add_para(
            doc,
            "",
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            space_after=2,
            line_spacing=1.28,
            left_indent=indent,
            keep_with_next=idx < len(items) - 1,
        )
        add_run_text(p, "•  ", name=HEI, size=12, color=NAVY)
        add_run_text(p, item, name=SONG, size=12, color=INK)


def add_numbered(doc: Document, items: list[str]) -> None:
    for idx, item in enumerate(items, start=1):
        p = add_para(
            doc,
            "",
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            space_after=2,
            line_spacing=1.28,
            left_indent=0.4,
            keep_with_next=idx < len(items),
        )
        add_run_text(p, f"{idx}.  ", name=HEI, size=12, bold=True, color=NAVY)
        add_run_text(p, item, name=SONG, size=12, color=INK)


def set_run_east_asia_on_style(style, font_name: str, size_pt: float) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def add_page_number(paragraph: Paragraph) -> None:
    run = paragraph.add_run()
    _set_run_font(run, name=SONG, size=9, color=MUTED)

    def fld(kind: str, instr: str | None = None):
        if kind in {"begin", "end", "separate"}:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr or ""
            run._r.append(el)

    fld("begin")
    fld("instr", " PAGE ")
    fld("separate")
    page_run = paragraph.add_run("1")
    _set_run_font(page_run, name=SONG, size=9, color=MUTED)
    end_run = paragraph.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "end")
    end_run._r.append(el)


def setup_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    _set_paragraph_format(hp, alignment=WD_ALIGN_PARAGRAPH.RIGHT, space_before=0, space_after=2, line_spacing=1.0)
    add_run_text(hp, "内部工作文件  ·  联合投标工作方案", name=SONG, size=9, color=MUTED)
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), NAVY_HEX)
    pBdr.append(bottom)
    pPr.append(pBdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    _set_paragraph_format(fp, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=0, line_spacing=1.0)
    add_run_text(fp, "第 ", name=SONG, size=9, color=MUTED)
    add_page_number(fp)
    add_run_text(fp, " 页    项目编号：JLP26C00100", name=SONG, size=9, color=MUTED)


def build_document(output_path: Path) -> Path:
    doc = Document()
    set_run_east_asia_on_style(doc.styles["Normal"], SONG, 12)
    setup_header_footer(doc)

    props = doc.core_properties
    props.title = "九龙坡城市更新项目联合投标工作方案"
    props.subject = "重庆市九龙坡区高质量推动城市更新“十五五”时期实施方案"
    props.author = "复旦大学住房政策研究中心"
    props.category = "内部工作文件"
    props.comments = "竞争性磋商项目 JLP26C00100 联合投标工作方案"

    add_para(
        doc,
        "重庆市九龙坡区高质量推动城市更新",
        name=HEI,
        size=18,
        bold=True,
        color=NAVY,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=2,
        space_after=2,
        line_spacing=1.2,
        keep_with_next=True,
    )
    add_para(
        doc,
        "“十五五”时期实施方案",
        name=HEI,
        size=18,
        bold=True,
        color=NAVY,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=4,
        line_spacing=1.2,
        keep_with_next=True,
    )
    add_para(
        doc,
        "联合投标工作方案",
        name=HEI,
        size=16,
        bold=True,
        color=INK,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=2,
        line_spacing=1.2,
        keep_with_next=True,
    )
    add_para(
        doc,
        "报名、分工、技术方案框架及时间安排",
        name=SONG,
        size=12,
        color=MUTED,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=4,
        line_spacing=1.2,
        keep_with_next=True,
    )
    add_para(
        doc,
        "竞争性磋商　·　项目编号 JLP26C00100　·　2026年9月",
        name=SONG,
        size=10.5,
        color=MUTED,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=8,
        line_spacing=1.15,
        keep_with_next=True,
    )

    add_callout(
        doc,
        "办理时限",
        "采购文件获取截止：2026年9月18日18:00。响应文件递交截止及磋商时间：2026年9月23日14:30。"
        "须在重庆政府采购网完成网上报名且审核通过，方视为合法获取采购文件；公告网页所附文件仅供阅览，不能替代正式报名。",
    )

    add_section_heading(doc, "一、项目基本信息")
    add_kv_table(
        doc,
        [
            ("项目名称", "重庆市九龙坡区高质量推动城市更新“十五五”时期实施方案", {}),
            ("项目编号", "JLP26C00100", {"emphasis": True}),
            ("采购单位", "重庆市九龙坡区住房和城乡建设委员会", {}),
            ("采购方式", "竞争性磋商", {}),
            ("项目预算及最高限价", "132万元", {}),
            ("采购文件获取截止", "2026年9月18日18:00", {"alert": True}),
            ("递交截止 / 磋商时间", "2026年9月23日14:30", {"alert": True}),
            ("是否接受联合体", "接受。联合体资格细则以磋商文件附件为准。", {}),
            (
                "核心专业资质",
                "具备自然资源部颁发的有效城乡规划编制甲级资质，或城乡规划（国土空间规划）编制甲级资质。",
                {},
            ),
        ],
    )

    add_section_heading(doc, "二、合作基础与报名主体")
    add_para(
        doc,
        "重庆大学建筑规划设计研究总院有限公司官方资料显示，其具备城乡规划编制甲级资质，并同时具有建筑、市政、工程咨询等资质，"
        "与本项目核心资质门槛及重庆本地执行要求相匹配。现阶段由该院完成项目报名及正式采购文件获取。",
        first_line_indent=0.74,
    )
    add_kv_table(
        doc,
        [
            ("拟牵头单位", "重庆大学建筑规划设计研究总院有限公司", {"emphasis": True}),
            (
                "拟联合研究单位",
                "复旦大学相关主体 / 复旦大学住房政策研究中心研究团队",
                {},
            ),
        ],
    )
    add_para(
        doc,
        "时间较紧：9月18日18:00前须完成合法获取采购文件，9月23日14:30前递交响应文件。本项目接受联合体，"
        "但联合体具体资格要求以磋商文件附件为准。",
        first_line_indent=0.74,
        space_before=6,
    )

    add_section_heading(doc, "三、联合投标结构")
    add_para(
        doc,
        "拟按下列结构推进联合投标：",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_kv_table(
        doc,
        [
            ("牵头方", "重庆大学建筑规划设计研究总院有限公司", {"emphasis": True}),
            ("联合体研究方 / 智库方", "复旦大学一侧（主体名称待核）", {}),
        ],
    )
    add_para(
        doc,
        "重庆大学设计总院负责甲级规划资质、重庆本地基础、规划技术、空间方案、电子投标和商务组织；"
        "复旦一侧负责城市更新战略、住房政策、政策机制、全国案例、“十五五”实施逻辑、评估机制及智库成果。",
        first_line_indent=0.74,
        space_before=6,
    )
    add_subheading(doc, "须当日核清的主体资格问题")
    add_para(
        doc,
        "“复旦大学住房政策研究中心”能否作为联合体成员签署采购合同，须对照磋商文件确认。"
        "若研究中心不是独立法人、不具备独立政府采购供应商资格，正式联合体成员应写“复旦大学”，"
        "不宜单独写“复旦大学住房政策研究中心”。研究中心可在项目团队和成果署名中体现，例如：",
        first_line_indent=0.74,
    )
    add_kv_table(
        doc,
        [
            ("联合体成员", "复旦大学", {"emphasis": True}),
            ("项目实施团队", "复旦大学住房政策研究中心", {}),
        ],
    )
    add_para(
        doc,
        "请重庆大学设计总院商务人员对照磋商文件确认，必要时当日电话询问采购代理机构（023-63024313）。"
        "联合体资格细则以采购文件附件为准。最终主体名称以磋商文件要求及代理机构确认口径为准。",
        first_line_indent=0.74,
        space_before=6,
    )

    add_section_heading(doc, "四、双方工作分工")
    add_para(
        doc,
        "复旦一侧不承担规划制图主责，定位为战略、政策与智库研究，与重庆大学设计总院的规划技术能力形成互补。",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_grid_table(
        doc,
        ["工作模块", "重大总院", "复旦中心", "牵头"],
        [
            ["政府采购报名", "★★★★★", "配合", "重大总院"],
            ["重庆采购网注册 / CA", "★★★★★", "核主体资格", "重大总院"],
            ["联合体协议", "起草", "审核、盖章", "重大总院"],
            ["甲级规划资质", "提供", "—", "重大总院"],
            ["商务标总体组织", "★★★★★", "提供本方材料", "重大总院"],
            ["重庆 / 九龙坡现状研究", "★★★★★", "★★★", "重大总院"],
            ["城市更新空间规划", "★★★★★", "★★", "重大总院"],
            ["更新项目梳理 / 空间布局", "★★★★★", "★★★", "重大总院"],
            ["“十五五”战略判断", "★★★★", "★★★★★", "双方"],
            ["城市更新政策机制", "★★★", "★★★★★", "复旦"],
            ["住房与房地产专题", "★★★", "★★★★★", "复旦"],
            ["全国案例比较", "★★★", "★★★★★", "复旦"],
            ["城市更新实施机制", "★★★★", "★★★★★ 主笔", "复旦"],
            ["评估 / 监测指标体系", "★★★", "★★★★★ 主笔", "复旦"],
            ["投融资 / 政策工具研究", "★★★★", "★★★★", "双方"],
            ["成果汇报 / PPT", "★★★★", "★★★★★", "双方"],
            ["最终投标文件汇编上传", "★★★★★", "配合", "重大总院"],
        ],
        [5.4, 3.2, 4.0, 3.4],
        font_size=9.5,
        together=False,
        compact=True,
    )

    add_section_heading(doc, "五、报名及采购文件获取")
    add_subheading(doc, "（一）确认重庆市政府采购供应商账号")
    add_para(
        doc,
        "登录重庆市政府采购网，确认重庆大学建筑规划设计研究总院有限公司已注册为重庆市政府采购供应商。"
        "已有账号的，使用企业账号登录；尚未注册的，立即完成供应商注册及企业信息认证。"
        "参加本项目须先注册成为重庆市政府采购供应商。",
        first_line_indent=0.74,
    )

    add_subheading(doc, "（二）确认CA数字证书可用")
    add_para(
        doc,
        "本项目为全流程电子招投标项目。须确认重庆大学设计总院已持有可在重庆市政府采购网使用的正式CA数字证书，"
        "并完成CA证书安装、电子签章测试、登录测试、投标文件加密及签章环境测试。",
        first_line_indent=0.74,
    )
    add_para(
        doc,
        "首次参加重庆政府采购电子招投标项目的，进入：重庆市政府采购网 → 在线开评标 → 电子招投标中心，"
        "下载并按下列手册办理：",
        first_line_indent=0.74,
        space_before=3,
        keep_with_next=True,
    )
    add_numbered(
        doc,
        [
            "《正式CA签章流程手册》",
            "《政府采购全程电子化采购系统供应商操作手册》",
            "《政府采购CA版供应商投标前软件安装手册》",
        ],
    )
    add_para(
        doc,
        "CA及投标环境须在报名阶段完成测试，勿延至2026年9月23日递交当日。",
        first_line_indent=0.74,
        space_before=3,
        color=RED,
        bold=True,
    )

    add_subheading(doc, "（三）检索项目")
    add_para(
        doc,
        "登录重庆市政府采购网后，按下列路径进入：个人中心 → 在线开评标 → 电子标书在线获取。",
        first_line_indent=0.74,
    )
    add_para(
        doc,
        "按项目编号 JLP26C00100，或项目名称“重庆市九龙坡区高质量推动城市更新‘十五五’时期实施方案”检索。"
        "核对采购人为重庆市九龙坡区住房和城乡建设委员会、预算为132万元后进入项目。",
        first_line_indent=0.74,
    )

    add_subheading(doc, "（四）网上报名并获取采购文件")
    add_para(
        doc,
        "按系统提示依次办理：报名申请 → 确认供应商信息 → 提交报名 → 在线获取竞争性磋商文件 → 确认报名审核状态。",
        first_line_indent=0.74,
    )
    add_para(
        doc,
        "公告网页所附《磋商文件（定稿）》仅供阅览，不能替代正式报名。须登录重庆市政府采购网完成网上报名申请，"
        "且报名审核通过，方视为合法获取采购文件；否则响应文件可能被拒绝。请确认报名状态为“审核通过”或“已成功获取采购文件”。",
        first_line_indent=0.74,
    )

    add_subheading(doc, "（五）下载并保存全套正式资料")
    add_para(
        doc,
        "报名成功后，下载并保存项目全部资料，至少包括：",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_numbered(
        doc,
        [
            "《重庆市九龙坡区高质量推动城市更新“十五五”时期实施方案竞争性磋商文件》",
            "项目公告附件",
            "《重庆市全流程电子招投标项目投标人办理正式CA签章流程手册》",
            "《政府采购全程电子化采购系统供应商操作手册》",
            "《政府采购CA版供应商投标前软件安装手册》",
            "联合体相关格式文件",
            "响应文件格式",
            "最新补遗、澄清、答疑和更正文件",
            "平台后续新增的项目资料",
        ],
    )
    add_para(
        doc,
        "正式采购文件获取后，请立即整体打包发送复旦大学住房政策研究中心项目组。",
        first_line_indent=0.74,
        space_before=3,
    )

    add_subheading(doc, "（六）报名完成核对")
    add_para(
        doc,
        "重庆大学设计总院商务人员完成操作后，请按下列五项核对，并保存报名成功页面、电子标书获取成功页面及项目状态截图，便于双方确认。",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_grid_table(
        doc,
        ["序号", "核对事项", "完成情况"],
        [
            ["①", "重庆市政府采购网供应商账号正常", "□ 已确认"],
            ["②", "CA数字证书正常", "□ 已确认"],
            ["③", "JLP26C00100已进入“我的项目 / 我的投标项目”", "□ 已确认"],
            ["④", "网上报名申请已提交，并显示审核通过或合法获取采购文件", "□ 已确认"],
            ["⑤", "正式磋商文件及全部附件已下载", "□ 已确认"],
        ],
        [1.5, 11.7, 2.8],
    )

    add_section_heading(doc, "六、联合体报名处理")
    add_para(
        doc,
        "本项目接受联合体。报名阶段如系统要求填写联合体信息，请先对照正式磋商文件确认下列事项，再行填写：",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_bullets(
        doc,
        [
            "城乡规划甲级资质由哪一方持有",
            "联合体牵头方是否须为甲级规划单位",
            "联合体成员数量有无限制",
            "双方是否均须注册为重庆政府采购供应商",
            "联合体成员是否均须完成CA办理",
            "联合体协议在报名阶段还是响应文件阶段提交",
            "联合体各方业绩能否共同计入商务评分",
        ],
    )
    add_para(
        doc,
        "文件或系统规定不明确的，由重庆大学设计总院商务人员电话向采购代理机构确认后再操作。",
        first_line_indent=0.74,
        space_before=3,
    )
    add_subheading(doc, "联系方式")
    add_grid_table(
        doc,
        ["单位", "联系人", "电话"],
        [
            ["采购代理机构：三信建设咨询集团有限公司", "赵老师 / 赵茜", "023-63024313"],
            ["采购人：重庆市九龙坡区住房和城乡建设委员会", "吕老师", "023-68036909"],
        ],
        [9.0, 3.6, 3.4],
    )
    add_para(
        doc,
        "向代理机构确认时，可采用下列口径：",
        first_line_indent=0.74,
        space_before=6,
        keep_with_next=True,
    )
    add_quote(
        doc,
        "本项目拟由具有城乡规划甲级资质的重庆大学建筑规划设计研究总院作为牵头方，与复旦大学相关研究主体组成联合体。"
        "请确认电子标书获取阶段，是仅由联合体牵头方报名获取，还是联合体各方均需完成供应商注册及报名？"
        "另请确认：复旦大学住房政策研究中心能否作为独立联合体成员，或须以复旦大学作为联合体成员、研究中心作为实施团队？",
    )
    add_para(
        doc,
        "以采购代理机构正式答复为准。",
        first_line_indent=0.74,
        space_before=4,
    )

    add_section_heading(doc, "七、今日须完成的四项工作")
    add_para(
        doc,
        "现阶段先把投标资格和评分规则核清，再展开技术方案撰写。",
        first_line_indent=0.74,
    )

    add_subheading(doc, "（一）重庆大学设计总院完成正式报名")
    add_para(
        doc,
        "采购公告要求：只有在重庆市政府采购网完成网上报名并审核通过，才视为合法获取采购文件；公告网页附件仅供阅览。"
        "获取文件截止为9月18日18:00。请当日确认：重庆政府采购网账号 → CA → 报名 → 下载正式磋商文件 → 下载最新澄清文件。此项优先办理。",
        first_line_indent=0.74,
    )

    add_subheading(doc, "（二）双方拆解正式磋商文件")
    add_para(
        doc,
        "拿到完整磋商文件后，先拆五张表：资格表、评分表、业绩表、人员评分表、技术评分表。重点核对：联合体双方业绩是否均可计分。"
        "若规定联合体任一成员业绩均可计分，则复旦大学相关城市更新、住房政策、规划研究业绩可计入商务评分；"
        "若只计算牵头方业绩，则商务分须主要依靠重庆大学设计总院，联合体结构亦须相应调整。",
        first_line_indent=0.74,
    )

    add_subheading(doc, "（三）交换投标资产包")
    add_para(
        doc,
        "材料以评分能用上的为限，形成可直接调用的投标素材，不必堆砌长篇介绍。",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_grid_table(
        doc,
        ["提供方", "材料内容"],
        [
            [
                "重庆大学设计总院",
                "营业执照、甲级资质、类似业绩、项目负责人简历、规划师团队、重庆/九龙坡业绩、获奖和荣誉。",
            ],
            [
                "复旦一侧",
                "复旦大学主体材料、住房政策研究中心介绍、专家团队、城市更新/住房/政策研究业绩、白皮书与研究报告、政府研究课题、相关案例及专家职称学历材料。建议汇总为《复旦可用于九龙坡项目的投标素材包》。",
            ],
        ],
        [3.6, 12.4],
        together=False,
    )

    add_subheading(doc, "（四）预先明确联合体利益分配")
    add_para(
        doc,
        "中标前先书面确认下列事项：合同签署主体、收款主体、复旦承担的成果范围、工作量比例、费用比例、知识产权、成果署名、"
        "专家差旅、税费、修改次数、政府汇报出席安排。本项目预算132万元，双方宜按实际工作量确定费用比例，避免按各自完整成本核算后合作效率偏低。",
        first_line_indent=0.74,
    )

    add_section_heading(doc, "八、技术方案框架（6+1）")
    add_para(
        doc,
        "技术标不宜写成传统规划编制流程，而应体现“重庆大学规划能力 + 复旦政策智库能力”的组合特征。建议采用下列6+1结构：",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_grid_table(
        doc,
        ["模块", "核心内容", "主责"],
        [
            ["01 战略研判", "“十四五”复盘、“十五五”趋势、国家及重庆城市更新政策", "复旦"],
            ["02 城市更新体检", "人口、产业、住房、公共空间、存量资产、设施等", "重大+复旦"],
            ["03 更新空间体系", "重点片区、更新单元、空间结构、分类策略", "重大"],
            ["04 “十五五”重点任务", "居住、产业、商业、公共服务、历史文化等", "双方"],
            ["05 项目库与行动计划", "重点项目清单、年度时序、实施主体", "重大"],
            ["06 政策与实施机制", "土地、住房、资金、运营、审批、社会参与机制", "复旦"],
            ["+1 动态实施体系", "指标体系、年度评估、项目库更新、常年咨询机制", "复旦+重大"],
        ],
        [4.0, 8.8, 3.2],
        together=False,
        compact=True,
    )
    add_para(
        doc,
        "其中“+1 动态实施体系”用于把成果从一本“十五五实施方案”延伸为："
        "规划 → 项目库 → 年度行动计划 → 年度评估 → 动态维护。该机制亦可作为后续持续服务的工作接口。",
        first_line_indent=0.74,
        space_before=6,
    )

    add_section_heading(doc, "九、复旦主责章节")
    add_para(
        doc,
        "复旦一侧不承担过多绘图。下列五章由复旦做实后，即便空间规划由重庆大学设计总院完成，成果仍能体现住房政策研究中心的专业贡献：",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_numbered(
        doc,
        [
            "城市更新政策与发展趋势研判",
            "九龙坡住房与人口发展专题",
            "城市更新实施机制与政策工具",
            "国内典型城市更新案例及重庆启示",
            "“十五五”实施评估与动态监测机制",
        ],
    )

    add_section_heading(doc, "十、项目组织架构")
    add_para(
        doc,
        "对外不以“两家单位拼标”表述，而以联合项目组呈现：规划技术 + 政策研究 + 高校智库。",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_grid_table(
        doc,
        ["角色", "人选方向"],
        [
            ["项目总负责人", "重庆大学设计总院高级规划专家"],
            ["首席政策专家 / 智库负责人", "复旦大学住房政策研究中心专家"],
            ["城市更新规划组", "重庆大学设计总院"],
            ["住房与政策研究组", "复旦大学住房政策研究中心"],
            ["产业与存量空间组", "双方协同"],
            ["实施机制与评估组", "复旦主责、重大配合"],
        ],
        [5.6, 10.4],
        together=False,
    )

    add_section_heading(doc, "十一、时间安排")
    add_grid_table(
        doc,
        ["时间", "工作安排"],
        [
            ["9月17日", "确定牵头结构；重庆大学设计总院完成/确认重庆采购网报名；取得正式标书；双方互换材料。"],
            ["9月18日18:00前", "确认报名审核通过并合法获取正式采购文件；完成资格/评分/业绩矩阵；联合体协议基本定稿。"],
            ["9月19日", "双方完成技术方案框架；锁定项目负责人和核心团队。"],
            ["9月20日", "各模块初稿完成；完成商务证明材料。"],
            ["9月21日", "合并第一版完整响应文件；开始电子投标系统测试。"],
            ["9月22日", "完整审标：资格、签章、报价、人员、业绩逐项核查。"],
            ["9月23日上午", "完成最终文件上传；准备CA和远程磋商。请于上午完成，避免临近截止时刻操作。"],
            ["9月23日14:30", "响应文件递交截止并开始磋商。"],
        ],
        [4.2, 11.8],
        together=False,
    )
    add_para(
        doc,
        "本项目采用全流程电子投标。响应文件须通过平台投标客户端编制，并在重庆市政府采购网"
        "“我的投标项目 → 在线投标”中提交。开标时须使用CA，并配备带摄像头、麦克风的电脑，在规定时间内完成远程解密和磋商。",
        first_line_indent=0.74,
        space_before=6,
    )

    add_section_heading(doc, "十二、投标启动会")
    add_para(
        doc,
        "重庆大学设计总院取得正式磋商文件后，双方召开一次约30分钟的投标启动会，只确认下列五项，其余事项会后推进：",
        first_line_indent=0.74,
        keep_with_next=True,
    )
    add_numbered(
        doc,
        [
            "谁是正式联合体成员",
            "谁牵头",
            "商务评分是否足够",
            "谁负责哪几章",
            "谁负责最终投标上传",
        ],
    )

    add_section_heading(doc, "十三、当前工作重点")
    add_callout(
        doc,
        "首要任务",
        "重庆大学设计总院当前首要任务，是在2026年9月18日18:00前完成重庆政府采购网正式报名，"
        "并确认系统显示“合法获取采购文件”或“报名审核通过”。完成后，将系统正式下载的全部磋商文件和附件发送复旦项目组，"
        "双方按评分标准确定联合体结构并开展联合投标。",
    )

    add_para(
        doc,
        "— 完 —",
        name=SONG,
        size=10.5,
        color=MUTED,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=14,
        space_after=0,
        line_spacing=1.15,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    output = root / "deliverables" / "重庆市九龙坡区城市更新十五五实施方案_报名及采购文件获取操作指引.docx"
    build_document(output)
    print(f"已生成：{output}")


if __name__ == "__main__":
    main()
