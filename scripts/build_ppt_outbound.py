# -*- coding: utf-8 -*-
"""单独生成《出海行业方向与当地支持诉求》领事对接专报 PPT（原模板）。"""
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import content as C
import ppt_theme as T

OUT = Path(__file__).resolve().parent.parent / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)
OUT_FILE = OUT / "复兴岛_出海行业方向与当地支持诉求_领事对接专报.pptx"

TOTAL = 12
FOOTER_L = "领事对接专报｜行业方向 × 当地支持｜V3.3"


def main():
    prs = T.new_prs()
    SW, SH = prs.slide_width, prs.slide_height

    def slide():
        return T.add_slide(prs)

    def hdr(s, kicker, title, idx):
        T.header(s, kicker, title, idx, TOTAL, SW)

    def ftr(s):
        T.footer(s, FOOTER_L)

    # 01 封面
    s = slide()
    T.rect(s, 0, 0, Inches(0.18), SH, T.TEAL)
    T.rect(s, 0, Inches(5.55), SW, Pt(3), T.TEAL)
    T.rect(s, 0, Inches(5.63), Inches(4.8), Pt(3), T.GOLD)
    T.text(s, Inches(0.9), Inches(0.9), Inches(11.5), Inches(0.4),
           [[("杨浦区 · 复兴岛｜给东亚/东南亚总领事的对接专报", 13, T.TEAL2, True)]])
    T.text(s, Inches(0.9), Inches(1.45), Inches(11.7), Inches(2.0),
           [[(C.OUTBOUND_BRIEF_TITLE, 32, T.WHITE, True)],
            [(C.OUTBOUND_BRIEF_SUB, 16, T.GOLD, True)],
            [("配套主会：" + C.PROJECT_FULL, 13, T.LIGHT, False)]],
           space_after=8)
    T.text(s, Inches(0.9), Inches(3.85), Inches(11.5), Inches(1.4),
           [[("核心问题", 14, T.TEAL2, True)],
            [("企业出海聚焦哪些行业？落地后最需要当地什么支持？", 15, T.WHITE, True)],
            [("总领事如何有的放矢回答，甚至协助解决？", 15, T.WHITE, True)]],
           space_after=4)
    T.text(s, Inches(0.9), Inches(5.9), Inches(11.5), Inches(0.7),
           [[(C.EVENT_DATE + "　｜　分论坛 C / 国际酒会使用　｜　" + C.VERSION, 13, T.WHITE, True)],
            [("视觉模板与主方案一致 · 可单独呈送外事礼宾组与各总领事", 12, T.GREY, False)]],
           space_after=3)

    # 02 用途
    s = slide()
    hdr(s, "WHY THIS BRIEF", "专报用途：让总领事有的放矢", 2)
    ftr(s)
    for i, line in enumerate(C.OUTBOUND_BRIEF_PURPOSE):
        y = Inches(1.45) + i * Inches(1.15)
        T.rect(s, Inches(0.55), y, Inches(12.2), Inches(1.0), T.CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.rect(s, Inches(0.55), y, Inches(0.12), Inches(1.0), T.GOLD if i == 0 else T.TEAL)
        T.text(s, Inches(0.9), y + Inches(0.28), Inches(11.5), Inches(0.5),
               [[(line, 16, T.WHITE, False)]])

    # 03 六大方向总览
    s = slide()
    hdr(s, "SIX DIRECTIONS", "复兴岛侧重点出海 · 六大行业方向", 3)
    ftr(s)
    accents = [T.GOLD, T.TEAL, T.TEAL2, T.GOLD, T.TEAL, T.TEAL2]
    for i, ind in enumerate(C.OUTBOUND_INDUSTRIES):
        col, row = i % 3, i // 3
        x = Inches(0.4) + col * Inches(4.25)
        y = Inches(1.35) + row * Inches(2.7)
        T.card(
            s, x, y, Inches(4.05), Inches(2.5),
            f"{ind['code']} {ind['name']}",
            [f"代表企业：{ind['enterprises']}", f"优先市场：{ind['markets']}", ind["plan"][:36] + "…"],
            accent=accents[i], tsize=13, bsize=11,
        )

    # 04–06 行业详表（每页 2 个）
    for page_i, start in enumerate((0, 2, 4)):
        s = slide()
        pair = C.OUTBOUND_INDUSTRIES[start:start + 2]
        hdr(s, "INDUSTRY DETAIL", f"行业详解 {pair[0]['code']}–{pair[1]['code']}：方向 · 支持 · 领事可做", 4 + page_i)
        ftr(s)
        for j, ind in enumerate(pair):
            x = Inches(0.35) + j * Inches(6.45)
            T.rect(s, x, Inches(1.3), Inches(6.25), Inches(5.4), T.CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            T.rect(s, x, Inches(1.3), Inches(6.25), Pt(4), T.GOLD if j == 0 else T.TEAL)
            T.text(
                s, x + Inches(0.2), Inches(1.45), Inches(5.85), Inches(1.35),
                [
                    [(f"{ind['code']}  {ind['name']}", 16, T.WHITE, True)],
                    [(f"企业：{ind['enterprises']}", 12, T.GOLD, True)],
                    [(f"计划：{ind['plan']}", 11, T.LIGHT, False)],
                    [(f"市场：{ind['markets']}", 11, T.TEAL2, False)],
                ],
                space_after=3,
            )
            T.text(
                s, x + Inches(0.2), Inches(2.95), Inches(5.85), Inches(0.35),
                [[("出海后需要的当地支持", 12, T.GOLD, True)]],
            )
            runs = []
            for sp in ind["supports"]:
                runs.append([("· " + sp, 12, T.LIGHT, False)])
            T.text(s, x + Inches(0.2), Inches(3.3), Inches(5.85), Inches(2.0), runs, space_after=2)
            T.text(
                s, x + Inches(0.2), Inches(5.45), Inches(5.85), Inches(1.0),
                [
                    [("总领事可如何帮助", 12, T.TEAL2, True)],
                    [(ind["consul_can"], 12, T.WHITE, False)],
                ],
                space_after=2,
            )

    # 07 企业对照
    s = slide()
    hdr(s, "ENTERPRISE MAP", "杨浦重点企业 × 出海行业对照", 7)
    ftr(s)
    mapping = {
        "优刻得": "01 智算云与数据基础设施",
        "智谱": "02 大模型与行业智能应用",
        "苏度": "03 具身智能与智能制造装备",
        "卓益得": "03 具身智能与智能制造装备",
        "傲鲨": "03 具身智能与智能制造装备",
        "清宝": "03 具身智能与智能制造装备",
        "声网": "02 大模型与行业智能应用 / 06 文旅协同",
        "道客": "01 智算云与数据基础设施",
        "商米": "04 智能终端与物理 AI 商业场景",
        "美团": "05 低空经济与城市无人系统",
    }
    rows = []
    for a, b, c, _ in C.YANGPU_ENTERPRISES:
        key = next((k for k in mapping if k in a), "—")
        rows.append([a, b, mapping.get(key, "—"), c[:28] + ("…" if len(c) > 28 else "")])
    T.add_table(
        s, Inches(0.3), Inches(1.28), Inches(12.7), Inches(5.5),
        ["企业", "赛道", "对应出海方向", "杨浦关联"],
        rows,
        col_widths=[Inches(2.4), Inches(2.0), Inches(4.3), Inches(4.0)],
        font_size=10,
    )

    # 08 国别匹配
    s = slide()
    hdr(s, "COUNTRY MATCH", "国别 × 优先行业：圆桌分桌依据", 8)
    ftr(s)
    rows = [[a, b, c, d] for a, b, c, d in C.OUTBOUND_COUNTRY_MATCH]
    T.add_table(
        s, Inches(0.25), Inches(1.2), Inches(12.8), Inches(5.5),
        ["国家", "总领事", "优先对接行业", "会谈切入点"],
        rows,
        col_widths=[Inches(1.3), Inches(2.4), Inches(3.8), Inches(5.3)],
        font_size=9,
    )

    # 09 支持包
    s = slide()
    hdr(s, "SUPPORT PACKS", "当地支持五包：领事可直接认领回应", 9)
    ftr(s)
    for i, (name, desc) in enumerate(C.OUTBOUND_SUPPORT_PACKS):
        y = Inches(1.4) + i * Inches(1.0)
        T.rect(s, Inches(0.5), y, Inches(12.3), Inches(0.88), T.CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.rect(s, Inches(0.5), y, Inches(2.4), Inches(0.88), T.TEAL if i % 2 == 0 else T.GOLD)
        T.text(s, Inches(0.5), y, Inches(2.4), Inches(0.88),
               [[(name, 14, T.WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, Inches(3.15), y + Inches(0.25), Inches(9.3), Inches(0.45),
               [[(desc, 15, T.LIGHT, False)]])

    # 10 提问脚本
    s = slide()
    hdr(s, "ROUNDTABLE SCRIPT", "分论坛 C / 酒会建议提问（企业×领事共用）", 10)
    ftr(s)
    for i, q in enumerate(C.OUTBOUND_ROUNDTABLE_QUESTIONS):
        y = Inches(1.4) + i * Inches(1.0)
        T.rect(s, Inches(0.55), y, Inches(12.2), Inches(0.88), T.CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.rect(s, Inches(0.55), y, Inches(0.9), Inches(0.88), T.GOLD)
        T.text(s, Inches(0.55), y, Inches(0.9), Inches(0.88),
               [[(f"Q{i+1}", 18, T.WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, Inches(1.7), y + Inches(0.25), Inches(10.7), Inches(0.45),
               [[(q, 14, T.WHITE, False)]])

    # 11 会后闭环
    s = slide()
    hdr(s, "CLOSED LOOP", "会前装订 · 会上分桌 · 会后跟踪", 11)
    ftr(s)
    for i, step in enumerate(C.OUTBOUND_NEXT):
        y = Inches(1.8) + i * Inches(1.4)
        T.rect(s, Inches(1.0), y, Inches(11.3), Inches(1.15), T.CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        T.rect(s, Inches(1.0), y, Inches(1.1), Inches(1.15), T.TEAL if i < 2 else T.GOLD)
        T.text(s, Inches(1.0), y, Inches(1.1), Inches(1.15),
               [[(str(i + 1), 28, T.WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.text(s, Inches(2.4), y + Inches(0.35), Inches(9.6), Inches(0.5),
               [[(step, 16, T.WHITE, True)]])

    # 12 封底
    s = slide()
    T.rect(s, 0, 0, Inches(0.18), SH, T.TEAL)
    T.rect(s, 0, Inches(3.2), SW, Pt(3), T.TEAL)
    T.text(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(1.2),
           [[(C.OUTBOUND_BRIEF_TITLE, 26, T.WHITE, True)],
            [("行业说清楚 · 支持列明白 · 领事能接住", 16, T.GOLD, True)]],
           align=PP_ALIGN.CENTER, space_after=10)
    T.text(s, Inches(0.9), Inches(3.5), Inches(11.5), Inches(1.8),
           [[("配套主方案：" + C.PROJECT_NAME, 14, T.LIGHT, True)],
            [("使用场景：外事邀约附页 · 分论坛 C 桌牌资料 · 总领事会前预读", 13, T.TEAL2, True)],
            [("正式对外口径与出席确认，仍以区外办 / 市外办指导为准", 12, T.GREY, False)],
            [(C.VERSION, 12, T.GREY, False)]],
           align=PP_ALIGN.CENTER, space_after=5)

    prs.save(OUT_FILE)
    print(f"已生成：{OUT_FILE}（{TOTAL} 页）")


if __name__ == "__main__":
    main()
