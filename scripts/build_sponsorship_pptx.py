#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成王德峰教授《传习录》与阳明心学专题活动的赞助方案 PPT。

对外版可发送给企业；内部版含招商目标、收支测算与三方机制，请勿外发。
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"

# 画面为 16:9。颜色取宣纸、朱砂、墨色，避免做成通用商务蓝模板。
W = 13.333333
H = 7.5

INK = RGBColor(0x14, 0x11, 0x0E)
INK_2 = RGBColor(0x2C, 0x26, 0x22)
PAPER = RGBColor(0xF6, 0xF1, 0xE8)
PAPER_2 = RGBColor(0xEF, 0xE8, 0xDC)
CARD = RGBColor(0xFF, 0xFC, 0xF7)
CINNABAR = RGBColor(0x8E, 0x2A, 0x28)
CINNABAR_DEEP = RGBColor(0x6E, 0x21, 0x1F)
GOLD = RGBColor(0xA6, 0x84, 0x45)
GOLD_LINE = RGBColor(0xC4, 0xA2, 0x65)
GOLD_TEXT = RGBColor(0x7A, 0x5A, 0x22)
CREAM = RGBColor(0xF7, 0xF1, 0xE6)
CREAM_DIM = RGBColor(0xD9, 0xCB, 0xB6)
TEXT = RGBColor(0x1C, 0x19, 0x16)
MUTED = RGBColor(0x5C, 0x56, 0x4E)
LINE = RGBColor(0xE3, 0xD9, 0xC8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PINE = RGBColor(0x1E, 0x3A, 0x32)
PINE_SOFT = RGBColor(0xE7, 0xF0, 0xEC)
WASH = RGBColor(0xF8, 0xF0, 0xEA)
GOLD_WASH = RGBColor(0xF8, 0xF3, 0xE6)
HEADER_GOLD = RGBColor(0x6E, 0x54, 0x24)

FONT = "微软雅黑"
BOXES: list[dict] = []


def rgb_of(color: RGBColor) -> RGBColor:
    return color


def set_run_font(run, size: float, color: RGBColor, bold: bool = False, font: str = FONT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    run.font.italic = False
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
        node = r_pr.find(qn(tag))
        if node is not None:
            r_pr.remove(node)
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = r_pr.makeelement(qn(tag), {"typeface": font})
        r_pr.append(el)


def _solid(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def add_rect(slide, x, y, w, h, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    _solid(shape, color)
    return shape


def add_round(slide, x, y, w, h, fill: RGBColor, line: RGBColor | None = None, radius: float = 0.08):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    _solid(shape, fill)
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    paragraphs: list[dict],
    *,
    anchor=MSO_ANCHOR.TOP,
    page: int = 0,
    pad: float = 0.02,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(pad)
    tf.margin_right = Inches(pad)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    expanded: list[dict] = []
    for item in paragraphs:
        parts = str(item.get("text", "")).split("\n")
        for j, part in enumerate(parts):
            piece = dict(item)
            piece["text"] = part
            if j > 0:
                piece["before"] = 0
                piece["after"] = item.get("after", 0) if j == len(parts) - 1 else 0
            expanded.append(piece)
    for i, item in enumerate(expanded):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = item.get("align", PP_ALIGN.LEFT)
        p.level = 0
        p.space_before = Pt(item.get("before", 0))
        p.space_after = Pt(item.get("after", 0))
        if item.get("spc") is not None:
            p.line_spacing = item["spc"]
        run = p.add_run()
        run.text = item.get("text", "")
        set_run_font(
            run,
            item.get("size", 14),
            item.get("color", TEXT),
            item.get("bold", False),
            item.get("font", FONT),
        )
    plain = "\n".join(item.get("text", "") for item in expanded)
    size = max(item.get("size", 14) for item in paragraphs)
    BOXES.append({"page": page, "text": plain, "size": size, "w": w, "h": h, "pad": pad})
    return box


def t(text, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT, before=0, after=0, spc=1.05):
    return {
        "text": text,
        "size": size,
        "color": color,
        "bold": bold,
        "align": align,
        "before": before,
        "after": after,
        "spc": spc,
    }


def new_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def paint(slide, color=PAPER):
    add_rect(slide, 0, 0, W, H, color)


def footer(slide, page: int, total: int, dark: bool = False):
    color = RGBColor(0xB7, 0xAA, 0x9C) if dark else MUTED
    rule = RGBColor(0x3A, 0x33, 0x2C) if dark else LINE
    add_rect(slide, 0.58, 7.08, W - 1.16, 0.01, rule)
    add_text(
        slide,
        0.58,
        7.12,
        9.2,
        0.30,
        [t("《传习录》与阳明心学  ·  企业品牌合作赞助方案", 11, color)],
        anchor=MSO_ANCHOR.MIDDLE,
        page=page,
    )
    add_text(
        slide,
        10.4,
        7.12,
        2.35,
        0.30,
        [t(f"{page:02d}  /  {total:02d}", 11, color, align=PP_ALIGN.RIGHT)],
        anchor=MSO_ANCHOR.MIDDLE,
        page=page,
    )


def header(slide, kicker: str, title: str, page: int, total: int):
    paint(slide, PAPER)
    add_rect(slide, 0, 0, W, 0.08, CINNABAR)
    add_text(slide, 0.58, 0.28, 10, 0.28, [t(kicker, 12, CINNABAR, True)], page=page)
    add_text(slide, 0.58, 0.54, 12.1, 0.48, [t(title, 26, TEXT, True)], page=page)
    add_rect(slide, 0.58, 1.12, 0.72, 0.035, GOLD_LINE)
    footer(slide, page, total)


def bullets(items: list[str], size=13, color=TEXT, after=5, spc=1.08):
    rows = []
    for i, item in enumerate(items):
        rows.append(t(item, size, color, after=after if i < len(items) - 1 else 0, spc=spc))
    return rows


# ---------------------------------------------------------------------------
# 对外版
# ---------------------------------------------------------------------------

def slide_cover(prs, page, total):
    s = new_slide(prs)
    paint(s, INK)
    add_rect(s, 0, 0, 0.12, H, CINNABAR)
    add_round(s, 0.62, 0.42, 0.72, 0.72, CINNABAR, radius=0.12)
    add_text(
        s, 0.62, 0.50, 0.72, 0.56,
        [t("心", 26, CREAM, True, PP_ALIGN.CENTER)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )
    add_text(
        s, 1.52, 0.48, 7.5, 0.58,
        [t("企业品牌合作赞助方案", 16, GOLD_LINE, True)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )
    add_text(
        s, 8.6, 0.48, 4.1, 0.58,
        [t("2026.10.31   上海", 14, CREAM_DIM, align=PP_ALIGN.RIGHT)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )
    add_text(s, 0.62, 1.85, 12, 0.32, [t("特邀主讲", 14, GOLD_LINE, True)], page=page)
    add_text(s, 0.62, 2.18, 12, 0.72, [t("王德峰教授", 44, CREAM, True)], page=page)
    add_text(
        s, 0.62, 2.95, 12, 0.38,
        [t("复旦大学哲学学院原教授、博士生导师", 16, CREAM_DIM)],
        page=page,
    )
    add_rect(s, 0.62, 3.52, 1.7, 0.025, GOLD_LINE)
    add_text(s, 0.62, 3.75, 12, 0.78, [t("《传习录》与阳明心学", 40, CREAM, True)], page=page)
    add_text(s, 0.62, 4.58, 12, 0.46, [t("专题讲座暨企业品牌合作", 22, GOLD_LINE)], page=page)
    add_text(
        s, 0.62, 5.25, 12, 0.40,
        [t("2026年10月31日（星期六）    全天专题研修    上海市", 16, CREAM)],
        page=page,
    )
    add_rect(s, 0.62, 5.85, 12.05, 0.012, RGBColor(0x3A, 0x33, 0x2C))
    add_text(s, 0.62, 6.02, 12, 0.28, [t("拟联合主办", 12, GOLD_LINE, True)], page=page)
    add_text(
        s, 0.62, 6.32, 12.05, 0.36,
        [t("复旦大学住房政策研究中心    ·    上海市杨浦区科技企业联合会    ·    首乾书院", 14, CREAM)],
        page=page,
    )
    add_text(
        s, 0.62, 6.85, 12.05, 0.40,
        [t("主办单位名称及相关标识的对外使用，以各机构正式授权为准。活动场地确认后另行通知。", 12, RGBColor(0xB7, 0xAA, 0x9C))],
        page=page,
    )


def slide_quote(prs, page, total):
    s = new_slide(prs)
    paint(s, PAPER)
    add_rect(s, 0, 0, W, 0.08, CINNABAR)
    add_text(s, 0.7, 0.85, 8, 0.32, [t("活动理念", 13, CINNABAR, True)], page=page)
    add_text(
        s, 0.7, 1.45, 11.8, 1.9,
        [
            t("在不确定的时代，寻找确定的内心；", 30, TEXT, True, after=8, spc=1.15),
            t("在纷繁复杂的世界中，重拾知行合一的智慧。", 30, TEXT, True, spc=1.15),
        ],
        page=page,
    )
    add_rect(s, 0.7, 3.6, 1.35, 0.03, GOLD_LINE)
    add_text(
        s, 0.7, 3.78, 10, 0.36,
        [t("主题    《传习录》与阳明心学", 15, MUTED)],
        page=page,
    )
    concepts = [
        ("心即理", "理解人生选择"),
        ("知行合一", "落实现实行动"),
        ("致良知", "安顿个人修养"),
    ]
    x = 0.7
    for name, desc in concepts:
        add_round(s, x, 4.45, 3.7, 1.45, CARD, LINE, 0.08)
        add_rect(s, x + 0.18, 4.62, 0.46, 0.045, CINNABAR)
        add_text(s, x + 0.28, 4.78, 3.2, 0.48, [t(name, 20, TEXT, True)], page=page)
        add_text(s, x + 0.28, 5.3, 3.2, 0.4, [t(desc, 14, MUTED)], page=page)
        x += 3.95
    footer(s, page, total)


def slide_toc(prs, page, total):
    s = new_slide(prs)
    header(s, "目录", "方案结构", page, total)
    items = [
        ("01", "活动缘起", "哲学智慧如何回应今天的人生命题", "04"),
        ("02", "活动安排", "时间、形式、人群与规模", "05"),
        ("03", "联合主办", "三家机构的内容、渠道与组织能力", "06"),
        ("04", "赞助体系", "战略、特别、品牌与实物支持", "07"),
        ("05", "合作价值", "品牌形象、客户触达与内部文化", "13"),
        ("06", "合作方式", "签约主体、权益边界与时间", "15"),
    ]
    for i, (num, title, desc, dest) in enumerate(items):
        col = i % 2
        row = i // 2
        x = 0.58 + col * 6.25
        y = 1.5 + row * 1.72
        add_round(s, x, y, 5.95, 1.52, CARD, LINE, 0.08)
        add_text(s, x + 0.28, y + 0.28, 1.3, 0.9, [t(num, 26, CINNABAR, True)], anchor=MSO_ANCHOR.MIDDLE, page=page)
        add_text(s, x + 1.6, y + 0.26, 3.9, 0.48, [t(title, 20, TEXT, True)], page=page)
        add_text(s, x + 1.6, y + 0.78, 3.9, 0.48, [t(desc, 13, MUTED)], page=page)


def slide_background(prs, page, total):
    s = new_slide(prs)
    header(s, "01    活动缘起", "以哲学智慧，回应时代命题", page, total)
    add_text(
        s, 0.58, 1.42, 7.15, 2.55,
        [
            t("在经济环境不断变化、人工智能快速发展、社会竞争日益激烈的今天，人们对人生价值、内心秩序与精神成长的关注不断增加。", 15, TEXT, after=10, spc=1.2),
            t("王阳明的心学，尤其是“心即理”“知行合一”“致良知”，为理解人生选择、个人修养与现实行动提供了重要的哲学视角。", 15, TEXT, after=10, spc=1.2),
            t("活动特邀王德峰教授，以《传习录》与阳明心学为主题展开。它既是一次哲学学习，也是面向企业经营者、专业人士及传统文化爱好者的思想交流。", 15, TEXT, spc=1.2),
        ],
        page=page,
    )
    points = [
        ("01", "心即理", "为理解人生选择提供哲学视角。"),
        ("02", "知行合一", "把认识带到现实行动之中。"),
        ("03", "致良知", "指向个人修养与内心秩序。"),
    ]
    y = 1.42
    for num, name, desc in points:
        add_round(s, 8.0, y, 4.75, 1.28, CARD, LINE, 0.08)
        add_text(s, 8.2, y + 0.16, 0.7, 0.36, [t(num, 12, CINNABAR, True)], page=page)
        add_text(s, 8.9, y + 0.12, 3.5, 0.40, [t(name, 18, TEXT, True)], page=page)
        add_text(s, 8.2, y + 0.62, 4.3, 0.5, [t(desc, 13, MUTED)], page=page)
        y += 1.42
    add_round(s, 0.58, 5.85, 12.17, 1.0, CARD, LINE, 0.08)
    add_rect(s, 0.58, 5.85, 0.08, 1.0, CINNABAR)
    add_text(
        s, 0.9, 5.95, 11.6, 0.8,
        [t("活动将通过三家机构各自的合法宣传渠道推广，吸引关注人文思想与个人成长的社会群体。", 15, TEXT)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )


def slide_facts(prs, page, total):
    s = new_slide(prs)
    header(s, "02    活动安排", "活动基本情况", page, total)
    rows = [
        ("主题", "《传习录》与阳明心学"),
        ("主讲嘉宾", "王德峰教授"),
        ("日期", "2026年10月31日（星期六）"),
        ("活动形式", "全天专题研修"),
        ("活动地点", "上海市（场地确认后通知）"),
        ("上午", "《传习录》与阳明心学（上）"),
        ("下午", "《传习录》与阳明心学（下）"),
        ("参与人群", "企业家、管理者、专业人士\n哲学及传统文化爱好者"),
        ("活动规模", "计划 200 人，以场地和实际报名为准"),
        ("合作形式", "企业赞助、品牌支持、实物支持"),
    ]
    for i, (label, value) in enumerate(rows):
        col = 0 if i < 5 else 1
        row = i if i < 5 else i - 5
        x = 0.58 + col * 6.25
        y = 1.4 + row * 1.02
        add_round(s, x, y, 6.05, 0.9, CARD, LINE, 0.1)
        add_text(s, x + 0.22, y + 0.1, 1.55, 0.7, [t(label, 13, CINNABAR, True)], anchor=MSO_ANCHOR.MIDDLE, page=page)
        add_text(s, x + 1.8, y + 0.1, 4.05, 0.7, [t(value, 14, TEXT, spc=1.05)], anchor=MSO_ANCHOR.MIDDLE, page=page)


def slide_hosts(prs, page, total):
    s = new_slide(prs)
    header(s, "03    联合主办", "三家机构，一种互补", page, total)
    add_text(
        s, 0.58, 1.38, 12.1, 0.7,
        [t("本次合作的基础，不只是主讲人的学术影响力，而是三家机构结合起来的内容公信、企业渠道和活动组织能力。下列特点描述各自更贴近的资源，具体分工以三方约定为准。", 15, TEXT, spc=1.15)],
        page=page,
    )
    cards = [
        (CINNABAR, "内容公信", "复旦大学住房政策研究中心", "以研究机构的公共信誉\n支撑活动的学术品质\n连接思想资源与社会对话"),
        (PINE, "企业渠道", "上海市杨浦区科技企业联合会", "连接科技企业与经营者\n服务企业家和管理者\n进入真实的交流场景"),
        (GOLD_TEXT, "活动组织", "首乾书院", "以人文学习与现场组织\n服务关注传统文化\n与个人成长的听众"),
    ]
    for i, (accent, role, name, desc) in enumerate(cards):
        x = 0.58 + i * 4.15
        add_round(s, x, 2.3, 3.98, 3.55, CARD, LINE, 0.08)
        add_rect(s, x, 2.3, 3.98, 0.08, accent)
        add_text(s, x + 0.28, 2.58, 3.4, 0.36, [t(role, 13, accent, True)], page=page)
        add_text(s, x + 0.24, 3.08, 3.52, 0.7, [t(name, 16, TEXT, True, spc=1.05)], page=page)
        add_text(s, x + 0.24, 3.85, 3.52, 1.6, [t(desc, 15, MUTED, spc=1.2)], page=page)
    add_text(
        s, 0.58, 6.05, 12.1, 0.85,
        [t("拟联合主办。主办单位名称及相关标识的对外使用，以各机构正式授权为准。活动推广通过三家机构各自的合法宣传渠道进行。", 13, MUTED, spc=1.15)],
        page=page,
    )


def slide_system(prs, page, total):
    s = new_slide(prs)
    header(s, "04    赞助体系", "一级战略，两级支持，另开实物", page, total)
    add_text(
        s, 0.58, 1.36, 12.1, 0.55,
        [t("采用“1 家战略支持 + 2—3 家特别支持 + 若干品牌支持”的结构，并引入茶饮、礼品等实物支持。门槛按单场文化活动的实际决策尺度设置。", 15, TEXT)],
        page=page,
    )
    tiers = [
        (CINNABAR, "战略支持单位", "3万元", "限 1 家", "本次最高级别商业支持"),
        (INK, "特别支持单位", "1万元", "2—3 家", "重要企业合作伙伴"),
        (GOLD_TEXT, "品牌支持单位", "5,000元", "3—5 家", "参与高品质文化活动"),
        (PINE, "实物及服务支持", "按贡献", "另行确认", "茶饮、礼品、场地与服务"),
    ]
    for i, (accent, name, price, quota, desc) in enumerate(tiers):
        x = 0.58 + i * 3.15
        add_round(s, x, 2.15, 3.02, 3.55, CARD, LINE, 0.08)
        add_rect(s, x, 2.15, 3.02, 0.08, accent)
        add_text(s, x + 0.2, 2.45, 2.62, 0.7, [t(name, 16, TEXT, True, spc=1.1)], page=page)
        add_text(s, x + 0.2, 3.3, 2.62, 0.7, [t(price, 28, accent, True)], page=page)
        add_text(s, x + 0.2, 4.15, 2.62, 0.4, [t(quota, 16, TEXT, True)], page=page)
        add_text(s, x + 0.2, 4.7, 2.62, 0.7, [t(desc, 13, MUTED, spc=1.15)], page=page)
    add_text(
        s, 0.58, 5.9, 12.1, 0.95,
        [t("上述级别与权益为招商建议。具体展示位置、形式及交付数量，须在合作协议中写明。赞助席位计入活动入场容量，与正式学员享有相同的课程权益。", 14, MUTED, spc=1.15)],
        page=page,
    )


def _tier_head(s, kicker, title, price, quota, positioning, page, total):
    header(s, kicker, title, page, total)
    add_round(s, 9.15, 0.32, 3.6, 0.7, CINNABAR, radius=0.1)
    add_text(
        s, 9.25, 0.36, 3.4, 0.62,
        [t(f"{price}    {quota}", 15, WHITE, True, PP_ALIGN.CENTER)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )
    add_text(s, 0.58, 1.32, 12.1, 0.4, [t(positioning, 15, MUTED)], page=page)


def slide_strategic(prs, page, total):
    s = new_slide(prs)
    _tier_head(
        s, "04    赞助体系", "战略支持单位", "3万元", "限 1 家",
        "定位：本次活动最高级别商业支持单位。", page, total,
    )
    columns = [
        ("品牌宣传", [
            "享有“活动战略支持单位”称谓",
            "名称或标识列入经审核的主视觉海报",
            "现场主背景板展示企业标识",
            "签到处或指定区域设置品牌展示",
            "主办方实际发布的活动介绍中列名",
            "主持人在开场或结束时鸣谢",
        ]),
        ("企业参与", [
            "提供 4 个全天研修席位",
            "可邀请负责人、重要客户或伙伴",
            "经现场安排许可，可设 1 处展示区",
            "资料放置于自愿领取的资料区",
            "参与正常交流，与现场嘉宾建立联系",
        ]),
        ("活动传播", [
            "在活动回顾内容中鸣谢",
            "获得可合法使用的现场照片及回顾资料",
            "可在授权范围内宣传本次支持",
        ]),
    ]
    for i, (name, items) in enumerate(columns):
        x = 0.58 + i * 4.15
        add_round(s, x, 1.85, 3.98, 4.15, CARD, LINE, 0.08)
        add_text(s, x + 0.24, 2.02, 3.5, 0.42, [t(name, 16, CINNABAR, True)], page=page)
        add_text(
            s, x + 0.24, 2.52, 3.52, 3.25,
            bullets([f"·  {item}" for item in items], 13, TEXT, after=6, spc=1.05),
            page=page,
        )
    add_round(s, 0.58, 6.15, 12.15, 0.78, WASH, radius=0.08)
    add_text(
        s, 0.8, 6.22, 11.75, 0.64,
        [t("特别说明：战略支持不包含王德峰教授个人商业代言、私下会见、专属授课或与教授合影。", 14, CINNABAR_DEEP)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )


def slide_special(prs, page, total):
    s = new_slide(prs)
    _tier_head(
        s, "04    赞助体系", "特别支持单位", "1万元", "建议 2—3 家",
        "定位：本次活动的重要企业合作伙伴。", page, total,
    )
    add_round(s, 0.58, 1.85, 8.15, 4.95, CARD, LINE, 0.08)
    items = [
        "享有“活动特别支持单位”称谓",
        "活动海报及现场背景板展示名称或标识",
        "活动主持人口头鸣谢",
        "提供 2 个全天研修席位",
        "可在指定区域展示企业宣传资料",
        "活动回顾中展示企业名称",
        "获得经授权的活动现场照片和传播素材",
    ]
    add_text(
        s, 0.9, 2.08, 7.55, 4.5,
        bullets([f"·  {item}" for item in items], 18, TEXT, after=22, spc=1.0),
        page=page,
    )
    add_round(s, 8.95, 1.85, 3.8, 4.95, INK, radius=0.08)
    add_text(s, 9.2, 2.08, 3.35, 0.4, [t("适合对象", 14, GOLD_LINE, True)], page=page)
    add_text(
        s, 9.2, 2.52, 3.3, 1.15,
        [t("关注企业家群体、人文教育及品质生活的企业。", 15, CREAM, spc=1.15)],
        page=page,
    )
    for i, name in enumerate(["企业服务", "科技创新", "文化教育", "健康管理"]):
        y = 3.9 + i * 0.65
        add_round(s, 9.25, y, 3.2, 0.52, RGBColor(0x2C, 0x26, 0x22), radius=0.12)
        add_text(s, 9.25, y, 3.2, 0.52, [t(name, 15, CREAM, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE, page=page)


def slide_brand(prs, page, total):
    s = new_slide(prs)
    _tier_head(
        s, "04    赞助体系", "品牌支持单位", "5,000元", "建议 3—5 家",
        "定位：以清晰的成本，参与一场高品质文化活动。", page, total,
    )
    add_round(s, 0.58, 1.85, 8.15, 4.95, CARD, LINE, 0.08)
    items = [
        "享有“活动品牌支持单位”称谓",
        "在赞助单位展示区域列示名称或标识",
        "提供 1 个全天研修席位",
        "席位与正式购票学员享有相同课程权益",
        "可在指定位置陈列品牌介绍资料",
        "在活动回顾中统一鸣谢",
        "获得活动现场的合规传播素材",
    ]
    add_text(
        s, 0.9, 2.08, 7.55, 4.5,
        bullets([f"·  {item}" for item in items], 18, TEXT, after=20, spc=1.0),
        page=page,
    )
    add_round(s, 8.95, 1.85, 3.8, 4.95, GOLD_WASH, radius=0.08)
    add_text(
        s, 9.2, 2.2, 3.3, 4.2,
        [
            t("更适合", 14, GOLD_TEXT, True, after=14),
            t("希望参与品牌文化建设，\n扩大本地知名度的\n中小企业。", 18, TEXT, True, spc=1.15),
        ],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )


def slide_inkind(prs, page, total):
    s = new_slide(prs)
    header(s, "04    赞助体系", "实物及服务支持", page, total)
    add_text(
        s, 0.58, 1.34, 12.1, 0.42,
        [t("不设置统一金额。根据实际提供的产品或服务，确认称谓、现场展示和鸣谢方式。", 15, TEXT)],
        page=page,
    )
    rows = [
        ("茶饮支持", "茶叶、茶饮品牌", "饮用茶及茶歇"),
        ("礼品支持", "文创、生活方式品牌", "参会伴手礼"),
        ("图书支持", "出版社、书店", "哲学及传统文化书籍"),
        ("餐饮支持", "餐饮企业", "工作餐、茶歇"),
        ("场地支持", "酒店、会议中心", "场地及会议服务"),
        ("技术支持", "科技、会议服务企业", "签到、摄影、音视频"),
    ]
    for i, (kind, who, what) in enumerate(rows):
        col = i % 3
        row = i // 3
        x = 0.58 + col * 4.15
        y = 1.9 + row * 1.85
        add_round(s, x, y, 3.98, 1.7, CARD, LINE, 0.08)
        add_text(s, x + 0.24, y + 0.16, 3.5, 0.38, [t(kind, 16, CINNABAR, True)], page=page)
        add_text(s, x + 0.24, y + 0.58, 3.5, 0.36, [t(who, 13, MUTED)], page=page)
        add_text(s, x + 0.24, y + 1.02, 3.5, 0.42, [t(what, 15, TEXT, True)], page=page)
    add_text(
        s, 0.58, 5.7, 12.15, 1.2,
        [
            t("实物支持单位可获得“指定服务支持单位”称谓，并按实际贡献获得相应的现场展示和鸣谢。", 14, TEXT, after=4),
            t("未经王德峰教授本人或权利方授权，不得将签名书、私人交流、合影等作为赞助回报。", 14, CINNABAR_DEEP),
        ],
        page=page,
    )


def slide_compare(prs, page, total):
    s = new_slide(prs)
    header(s, "04    赞助体系", "赞助权益对照", page, total)
    headers = [
        ("权益", INK),
        ("战略支持", CINNABAR),
        ("特别支持", INK_2),
        ("品牌支持", HEADER_GOLD),
    ]
    sub = ["", "3万元 · 限1家", "1万元 · 2—3家", "5,000元 · 3—5家"]
    data = [
        ["赞助单位称谓", "战略支持单位", "特别支持单位", "品牌支持单位"],
        ["主视觉海报", "核心位置", "常规位置", "支持单位专区"],
        ["现场背景板", "核心位置", "常规位置", "支持单位专区"],
        ["主持人鸣谢", "单独鸣谢", "统一鸣谢", "统一鸣谢"],
        ["全天研修名额", "4 人", "2 人", "1 人"],
        ["企业资料展示", "专属区域", "公共区域", "公共区域"],
        ["活动回顾露出", "重点鸣谢", "统一鸣谢", "统一鸣谢"],
        ["活动照片素材", "提供", "提供", "提供"],
        ["品牌使用授权", "按约定", "按约定", "按约定"],
        ["数量限制", "1 家", "2—3 家", "3—5 家"],
    ]
    xs = [0.5, 3.55, 6.55, 9.55]
    ws = [3.05, 3.0, 3.0, 3.28]
    y = 1.38
    hh = 0.62
    for i, (label, fill) in enumerate(headers):
        add_rect(s, xs[i], y, ws[i] - 0.06, hh, fill)
        add_text(
            s, xs[i], y + 0.04, ws[i] - 0.06, 0.30,
            [t(label, 13, CREAM, True, PP_ALIGN.CENTER)],
            anchor=MSO_ANCHOR.MIDDLE, page=page,
        )
        add_text(
            s, xs[i], y + 0.30, ws[i] - 0.06, 0.26,
            [t(sub[i], 11, CREAM, align=PP_ALIGN.CENTER)],
            anchor=MSO_ANCHOR.MIDDLE, page=page,
        )
    rh = 0.42
    for r, row in enumerate(data):
        yy = y + hh + r * rh
        bg = CARD if r % 2 == 0 else PAPER_2
        for c, value in enumerate(row):
            add_rect(s, xs[c], yy, ws[c] - 0.06, rh - 0.04, bg)
            weight = c == 0
            add_text(
                s, xs[c] + 0.06, yy, ws[c] - 0.18, rh - 0.04,
                [t(value, 13, TEXT, weight, PP_ALIGN.CENTER if c else PP_ALIGN.LEFT)],
                anchor=MSO_ANCHOR.MIDDLE, page=page,
            )
    add_text(
        s, 0.5, 6.45, 12.3, 0.55,
        [t("上述权益为招商建议，须在合作协议中明确具体展示位置、形式及交付数量。实物及服务支持另行确认，不列入本表。", 12, MUTED)],
        page=page,
    )


def slide_value(prs, page, total):
    s = new_slide(prs)
    header(s, "05    合作价值", "企业获得的，不只是标识上墙", page, total)
    cards = [
        ("01", "品牌文化形象", "支持传统哲学和学术交流，体现企业对人文精神、知识传播和社会文化活动的关注。特别适合希望塑造长期品牌形象的企业。"),
        ("02", "目标客户触达", "活动面向企业经营者、管理人员、专业人士和传统文化爱好者。企业可通过现场展示、合规资料发放及正常交流接触相关人群。"),
        ("03", "企业内部文化", "课程席位可用于高管学习、员工奖励或重要客户邀请。一次合作，同时完成文化学习、客户关系维护与品牌支持。"),
    ]
    for i, (num, title, desc) in enumerate(cards):
        x = 0.58 + i * 4.15
        add_round(s, x, 1.48, 3.98, 3.85, CARD, LINE, 0.08)
        add_text(s, x + 0.26, 1.7, 3.4, 0.4, [t(num, 14, CINNABAR, True)], page=page)
        add_text(s, x + 0.26, 2.15, 3.45, 0.7, [t(title, 22, TEXT, True)], page=page)
        add_text(s, x + 0.26, 3.0, 3.45, 2.0, [t(desc, 14, MUTED, spc=1.2)], page=page)
    add_round(s, 0.58, 5.52, 12.15, 1.35, CARD, LINE, 0.08)
    add_rect(s, 0.58, 5.52, 0.08, 1.35, CINNABAR)
    add_text(
        s, 0.9, 5.68, 11.55, 1.05,
        [
            t("不向赞助企业交付未经授权的报名者手机号、微信或其他个人信息。", 14, TEXT, after=4),
            t("赞助席位与正式购票学员享有相同的课程权益。", 14, TEXT),
        ],
        page=page,
    )


def slide_fit(prs, page, total):
    s = new_slide(prs)
    header(s, "05    合作价值", "适宜合作的企业", page, total)
    rows = [
        ("科技企业、人工智能企业", "1万—3万元", "企业家成长，科技与人文结合"),
        ("企业服务、管理咨询公司", "1万—3万元", "受众与企业客户有所重合"),
        ("高端茶叶、文化消费品牌", "5,000元—1万元", "与传统文化主题相契合"),
        ("文化出版、书店、文创企业", "实物或 5,000元", "可提供图书及文化用品"),
        ("酒店、餐饮、会务企业", "实物或服务支持", "场地、餐饮与会务协同"),
    ]
    add_rect(s, 0.58, 1.42, 12.15, 0.5, INK)
    for label, x, w in [("企业类型", 0.75, 4.3), ("建议级别", 5.1, 2.8), ("合作理由", 8.1, 4.3)]:
        add_text(s, x, 1.46, w, 0.42, [t(label, 13, CREAM, True)], anchor=MSO_ANCHOR.MIDDLE, page=page)
    for i, (kind, level, why) in enumerate(rows):
        y = 1.98 + i * 0.72
        bg = CARD if i % 2 == 0 else PAPER_2
        add_rect(s, 0.58, y, 12.15, 0.66, bg)
        add_text(s, 0.75, y, 4.3, 0.66, [t(kind, 14, TEXT, True)], anchor=MSO_ANCHOR.MIDDLE, page=page)
        add_text(s, 5.1, y, 2.8, 0.66, [t(level, 14, CINNABAR, True)], anchor=MSO_ANCHOR.MIDDLE, page=page)
        add_text(s, 8.1, y, 4.4, 0.66, [t(why, 14, TEXT)], anchor=MSO_ANCHOR.MIDDLE, page=page)
    add_text(
        s, 0.58, 5.7, 12.15, 1.15,
        [t("为维护学术活动的公信力，本次不接纳高风险金融产品营销、夸大健康功效的产品，以及其他可能损害活动公共形象的商业推广。", 14, MUTED, spc=1.15)],
        page=page,
    )


def slide_rules(prs, page, total):
    s = new_slide(prs)
    header(s, "06    合作方式", "签约、授权与权益边界", page, total)
    add_round(s, 0.58, 1.42, 6.2, 5.35, CARD, LINE, 0.08)
    add_text(s, 0.84, 1.6, 5.7, 0.4, [t("合作如何落地", 16, CINNABAR, True)], page=page)
    steps = [
        ("01", "由具备签约、收款和开票能力的主体与企业签订协议。"),
        ("02", "赞助款通过对公账户收取，不经由个人账户代收。"),
        ("03", "展示位置、形式、数量与交付时间在协议中列明。"),
        ("04", "活动延期、主讲人无法按约授课或取消时，可协商替代方案，或按合同约定退款。"),
        ("05", "企业宣传本次支持、主办方使用企业标识，均限于授权范围。"),
    ]
    y = 2.15
    for num, text in steps:
        add_text(s, 0.84, y, 0.6, 0.55, [t(num, 13, CINNABAR, True)], page=page)
        add_text(s, 1.5, y, 4.95, 0.78, [t(text, 13, TEXT, spc=1.05)], page=page)
        y += 0.85
    add_round(s, 7.0, 1.42, 5.75, 5.35, INK, radius=0.08)
    add_text(s, 7.28, 1.64, 5.2, 0.4, [t("明确不包含", 16, GOLD_LINE, True)], page=page)
    excludes = [
        "主讲人个人商业代言",
        "私下会见或专属授课",
        "未经授权的合影、签名书",
        "报名者手机号、微信等个人信息",
        "协议约定之外的品牌露出",
    ]
    add_text(
        s, 7.28, 2.25, 5.2, 2.7,
        bullets([f"·  {item}" for item in excludes], 16, CREAM, after=10),
        page=page,
    )
    add_text(
        s, 7.28, 5.15, 5.2, 1.3,
        [t("未经批准的品牌露出如何处理、违约责任及退款安排，写入赞助合同。", 14, CREAM_DIM, spc=1.2)],
        page=page,
    )


def slide_time(prs, page, total):
    s = new_slide(prs)
    header(s, "06    合作方式", "从意向到现场", page, total)
    add_text(
        s, 0.58, 1.4, 12.15, 0.7,
        [t("活动于 10 月 31 日举办。建议在 10 月 15 日前确认主要赞助意向，以便审核标识并制作正式海报和现场物料。", 16, TEXT, spc=1.15)],
        page=page,
    )
    steps = [
        ("10月15日前", "确认级别", "谈定级别与席位\n形成书面意向"),
        ("意向确认后", "签约付款", "签订协议并对公\n完成赞助付款"),
        ("物料制作期", "审核露出", "提交标识与文案\n审核后制作物料"),
        ("10月31日", "现场兑现", "展示、鸣谢和席位\n在现场一并落地"),
        ("活动结束后", "回顾传播", "提供经授权的\n照片与回顾资料"),
    ]
    for i, (when, title, desc) in enumerate(steps):
        x = 0.5 + i * 2.54
        add_round(s, x, 2.4, 2.4, 3.55, CARD, LINE, 0.1)
        add_rect(s, x, 2.4, 2.4, 0.08, CINNABAR if i == 0 else GOLD_LINE)
        add_text(s, x + 0.14, 2.62, 2.12, 0.7, [t(when, 13, CINNABAR, True, spc=1.05)], page=page)
        add_text(s, x + 0.14, 3.35, 2.12, 0.7, [t(title, 20, TEXT, True)], page=page)
        add_text(s, x + 0.14, 4.15, 2.12, 1.35, [t(desc, 14, MUTED, spc=1.2)], page=page)


def slide_close(prs, page, total):
    s = new_slide(prs)
    paint(s, INK)
    add_rect(s, 0, 0, 0.12, H, CINNABAR)
    add_text(s, 0.7, 0.48, 10, 0.36, [t("诚邀合作", 14, GOLD_LINE, True)], page=page)
    add_text(
        s, 0.7, 1.15, 12, 1.7,
        [
            t("五百年前的东方智慧，", 32, CREAM, True, after=6),
            t("如何回应今天的人生困境？", 32, CREAM, True),
        ],
        page=page,
    )
    add_text(
        s, 0.7, 3.15, 12, 0.7,
        [t("王德峰教授  ·  《传习录》与阳明心学\n2026年10月31日  ·  上海  ·  全天专题研修", 16, CREAM_DIM, spc=1.25)],
        page=page,
    )
    tiers = [("战略支持", "3万元"), ("特别支持", "1万元"), ("品牌支持", "5,000元")]
    card_w = 3.85
    for i, (name, price) in enumerate(tiers):
        x = 0.7 + i * (card_w + 0.2)
        add_round(s, x, 4.15, card_w, 1.15, RGBColor(0x23, 0x1C, 0x18), radius=0.1)
        add_text(s, x + 0.2, 4.22, card_w - 0.4, 0.38, [t(name, 13, GOLD_LINE)], anchor=MSO_ANCHOR.MIDDLE, page=page)
        add_text(s, x + 0.2, 4.58, card_w - 0.4, 0.55, [t(price, 22, CREAM, True)], anchor=MSO_ANCHOR.MIDDLE, page=page)
    add_text(
        s, 0.7, 5.5, 12, 0.7,
        [t("同时开放茶饮、图书、礼品、场地及相关服务支持。\n期待与优秀企业携手，共赴一场思想与智慧的盛宴。", 15, CREAM, spc=1.2)],
        page=page,
    )
    add_text(
        s, 0.7, 6.55, 12, 0.6,
        [t("欢迎与活动筹备组洽谈合作意向。主办单位名称、标识及主讲人肖像的对外使用，以正式授权为准；具体权益以合作协议为准。", 12, RGBColor(0xB7, 0xAA, 0x9C), spc=1.1)],
        page=page,
    )


EXTERNAL = [
    slide_cover,
    slide_quote,
    slide_toc,
    slide_background,
    slide_facts,
    slide_hosts,
    slide_system,
    slide_strategic,
    slide_special,
    slide_brand,
    slide_inkind,
    slide_compare,
    slide_value,
    slide_fit,
    slide_rules,
    slide_time,
    slide_close,
]


# ---------------------------------------------------------------------------
# 内部版
# ---------------------------------------------------------------------------

def in_footer(slide, page, total):
    add_rect(slide, 0, 7.18, W, 0.32, CINNABAR_DEEP)
    add_text(
        slide, 0.4, 7.18, 9.5, 0.32,
        [t("内部资料  ·  请勿对外发送  ·  赞助招商研讨", 12, CREAM, True)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )
    add_text(
        slide, 10.3, 7.18, 2.6, 0.32,
        [t(f"{page:02d}  /  {total:02d}", 12, CREAM, align=PP_ALIGN.RIGHT)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )


def in_header(slide, kicker, title, page, total):
    paint(slide, PAPER)
    add_rect(slide, 0, 0, W, 0.36, CINNABAR_DEEP)
    add_text(
        slide, 0.5, 0, 12.3, 0.36,
        [t("内部研讨    ·    不进入对外招商稿", 13, CREAM, True)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )
    add_text(slide, 0.55, 0.52, 12, 0.28, [t(kicker, 12, CINNABAR, True)], page=page)
    add_text(slide, 0.55, 0.8, 12.2, 0.46, [t(title, 26, TEXT, True)], page=page)
    in_footer(slide, page, total)


def slide_in_cover(prs, page, total):
    s = new_slide(prs)
    paint(s, INK)
    add_rect(s, 0, 0, W, 0.42, CINNABAR)
    add_text(
        s, 0.6, 0, 12, 0.42,
        [t("内部资料    ·    请勿对外发送", 14, CREAM, True)],
        anchor=MSO_ANCHOR.MIDDLE, page=page,
    )
    add_text(s, 0.7, 1.5, 12, 0.4, [t("赞助招商研讨", 16, GOLD_LINE, True)], page=page)
    add_text(s, 0.7, 2.05, 12, 1.5, [t("《传习录》与阳明心学\n招商口径、收支测算与三方机制", 32, CREAM, True, spc=1.15)], page=page)
    add_text(
        s, 0.7, 3.85, 11.5, 0.9,
        [t("供复旦大学住房政策研究中心、上海市杨浦区科技企业联合会、首乾书院统筹使用。\n本文件中的目标金额不是已经落实的收入。", 16, CREAM_DIM, spc=1.2)],
        page=page,
    )
    chips = [("6万—8万元", "现金招商目标"), ("不设总冠名", "战略支持 3 万元"), ("10月15日前", "确认主要意向")]
    for i, (head, sub) in enumerate(chips):
        x = 0.7 + i * 3.9
        add_round(s, x, 5.15, 3.7, 1.2, RGBColor(0x23, 0x1C, 0x18), radius=0.1)
        add_text(s, x + 0.2, 5.28, 3.3, 0.45, [t(head, 16, GOLD_LINE, True)], page=page)
        add_text(s, x + 0.2, 5.72, 3.3, 0.4, [t(sub, 14, CREAM)], page=page)


def slide_in_frame(prs, page, total):
    s = new_slide(prs)
    in_header(s, "招商定位", "三级现金赞助，外加实物支持", page, total)
    points = [
        ("结构", "1 家战略支持单位 + 2—3 家特别支持单位 + 若干品牌支持单位。"),
        ("门槛", "本次不宜设置过高的赞助门槛，不建议另设 5万—10万元总冠名。"),
        ("现金目标", "以 6万—8万元现金赞助作为主要招商目标，同时引入茶饮、礼品等实物支持。"),
        ("收入属性", "赞助收入用于分担活动成本，不能将其视为已经落实的收入。"),
    ]
    for i, (title, desc) in enumerate(points):
        y = 1.5 + i * 1.25
        add_round(s, 0.55, y, 12.2, 1.12, CARD, LINE, 0.08)
        add_text(s, 0.8, y + 0.14, 2.2, 0.84, [t(title, 16, CINNABAR, True)], anchor=MSO_ANCHOR.MIDDLE, page=page)
        add_text(s, 3.1, y + 0.16, 9.3, 0.82, [t(desc, 16, TEXT, spc=1.15)], anchor=MSO_ANCHOR.MIDDLE, page=page)


def slide_in_no_title(prs, page, total):
    s = new_slide(prs)
    in_header(s, "级别取舍", "不设置 5万—10万元总冠名", page, total)
    add_round(s, 0.55, 1.5, 12.2, 2.15, CARD, LINE, 0.08)
    add_text(s, 0.85, 1.7, 11.6, 0.4, [t("原因", 14, CINNABAR, True)], page=page)
    add_text(
        s, 0.85, 2.15, 11.6, 1.25,
        [
            t("当前已知付费人数较少，最终到场规模尚未确定。单场活动报出高额总冠名，成交难度较大。", 16, TEXT, after=8, spc=1.15),
            t("王德峰教授的个人学术形象，以及主办机构的品牌，不宜被包装成某企业的商业代言。", 16, TEXT, spc=1.15),
        ],
        page=page,
    )
    add_round(s, 0.55, 3.85, 12.2, 2.9, CARD, LINE, 0.08)
    add_text(s, 0.85, 4.05, 11.6, 0.4, [t("因此采用的报价", 14, CINNABAR, True)], page=page)
    add_text(
        s, 0.85, 4.55, 11.6, 1.9,
        [
            t("3 万元战略支持更适合当前阶段，且限 1 家。", 18, TEXT, True, after=8),
            t("对外表达为一次完整合作：正式课程席位、现场品牌展示、活动传播，而不是单独出售标识位置。", 16, TEXT, after=8, spc=1.15),
            t("战略支持明确不包含个人代言、私下会见、专属授课或与教授合影。", 16, TEXT, spc=1.15),
        ],
        page=page,
    )


def _money_table(slide, x, y, headers, rows, col_w, page, row_h=0.46):
    h = 0.48
    xx = x
    for i, label in enumerate(headers):
        add_rect(slide, xx, y, col_w[i] - 0.06, h, INK)
        add_text(
            slide, xx, y, col_w[i] - 0.06, h,
            [t(label, 13, CREAM, True, PP_ALIGN.CENTER)],
            anchor=MSO_ANCHOR.MIDDLE, page=page,
        )
        xx += col_w[i]
    for r, row in enumerate(rows):
        yy = y + h + r * row_h
        bg = CINNABAR if row[-1] == "合计" or (row[0].startswith("合计")) else (CARD if r % 2 == 0 else PAPER_2)
        color = CREAM if bg == CINNABAR else TEXT
        xx = x
        for c, value in enumerate(row):
            add_rect(slide, xx, yy, col_w[c] - 0.06, row_h - 0.05, bg)
            add_text(
                slide, xx + 0.06, yy, col_w[c] - 0.18, row_h - 0.05,
                [t(value, 14, color, True, PP_ALIGN.CENTER if c else PP_ALIGN.LEFT)],
                anchor=MSO_ANCHOR.MIDDLE, page=page,
            )
            xx += col_w[c]


def slide_in_targets(prs, page, total):
    s = new_slide(prs)
    in_header(s, "招商目标", "6 万元为底，8 万元为争取", page, total)
    add_text(
        s, 0.55, 1.4, 6, 0.4,
        [t("方案一    基础目标", 16, TEXT, True)],
        page=page,
    )
    add_text(
        s, 6.9, 1.4, 6, 0.4,
        [t("方案二    争取目标", 16, TEXT, True)],
        page=page,
    )
    headers = ["类型", "数量", "金额"]
    left = [
        ["战略支持单位", "1 家", "30,000元"],
        ["特别支持单位", "2 家", "20,000元"],
        ["品牌支持单位", "2 家", "10,000元"],
        ["合计", "5 家", "60,000元"],
    ]
    right = [
        ["战略支持单位", "1 家", "30,000元"],
        ["特别支持单位", "3 家", "30,000元"],
        ["品牌支持单位", "4 家", "20,000元"],
        ["合计", "8 家", "80,000元"],
    ]
    _money_table(s, 0.55, 1.9, headers, left, [2.5, 1.5, 2.15], page, 0.52)
    _money_table(s, 6.9, 1.9, headers, right, [2.5, 1.5, 2.15], page, 0.52)
    add_round(s, 0.55, 5.15, 12.2, 1.75, WASH, radius=0.08)
    add_text(
        s, 0.8, 5.35, 11.75, 1.4,
        [t("两套方案是招商目标，不是预计一定能完成的实际收入。\n优先争取 1 家战略支持，再通过现有企业关系落实 2—3 家特别支持。\n实物支持另计，不计入上述现金目标。", 16, TEXT, spc=1.15)],
        page=page,
    )


def slide_in_math(prs, page, total):
    s = new_slide(prs)
    in_header(s, "收支测算", "8 万元能减负，不能当成盈利", page, total)
    add_text(
        s, 0.55, 1.42, 12.2, 0.55,
        [t("以下按此前讨论的口径测算，只用于内部判断，不是对外报价，也不是已到账数字。", 15, MUTED)],
        page=page,
    )
    facts = [
        ("10万元", "主讲人费用口径"),
        ("28,788元", "12 人 × 2,399 元学费"),
        ("108,788元", "8 万赞助 + 上述学费"),
    ]
    for i, (num, label) in enumerate(facts):
        x = 0.55 + i * 4.15
        add_round(s, x, 2.1, 3.98, 1.55, CARD, LINE, 0.08)
        add_text(s, x + 0.2, 2.25, 3.55, 0.6, [t(num, 26, CINNABAR, True)], page=page)
        add_text(s, x + 0.2, 2.95, 3.55, 0.45, [t(label, 14, MUTED)], page=page)
    notes = [
        "即使完成 8 万元现金赞助，加上现有学费，合计也只是约 10.88 万元。",
        "此数尚未扣除场地、运营、票务服务等费用。",
        "赞助方案包含的课程席位占用入场容量，不能与同一席位的门票收入重复计算。",
        "结论：8 万元赞助可以显著减轻成本压力，但不能保证活动最终盈利。",
    ]
    add_text(
        s, 0.7, 3.9, 12, 2.9,
        bullets([f"·  {item}" for item in notes], 16, TEXT, after=10, spc=1.08),
        page=page,
    )


def slide_in_governance(prs, page, total):
    s = new_slide(prs)
    in_header(s, "三方机制", "先定收款主体，再谈怎么分", page, total)
    items = [
        ("统一合同主体", "三方确定有签约、收款和开票能力的主体，与企业签订协议。\n不由个人账户代收赞助款。"),
        ("统一活动预算", "写明师资、场地、设备、宣传、人员、茶歇等支出。\n签约前已发生的费用由原承担方负责，除非书面确认分担。"),
        ("收入分类记账", "赞助、门票、其他合作收入分别记账。\n赞助席位单独登记，避免与门票销售重复统计。"),
        ("净收益另行确定", "联合主办不自动按三分之一分配。\n按师资、运营、渠道、场地与财务风险协商规则。"),
        ("写明变更责任", "延期、无法按约授课或取消时，可协商替代或按约退款。\n合同写明未批准露出、违约责任与退款安排。"),
    ]
    for i, (title, desc) in enumerate(items):
        y = 1.42 + i * 1.08
        add_text(s, 0.55, y, 0.5, 0.9, [t(f"{i+1:02d}", 16, CINNABAR, True)], anchor=MSO_ANCHOR.TOP, page=page)
        add_text(s, 1.2, y, 3.3, 0.9, [t(title, 15, TEXT, True)], anchor=MSO_ANCHOR.TOP, page=page)
        add_text(s, 4.5, y, 8.2, 1.0, [t(desc, 13, MUTED, spc=1.05)], page=page)


def slide_in_priority(prs, page, total):
    s = new_slide(prs)
    in_header(s, "联系顺序", "先找已经有信任的企业", page, total)
    rows_h = ["优先级", "企业类型", "推荐级别", "合作理由"]
    rows = [
        ["第一", "科技企业、人工智能企业", "1万—3万元", "企业家成长，科技与人文结合"],
        ["第二", "企业服务、管理咨询公司", "1万—3万元", "受众与企业客户有所重合"],
        ["第三", "高端茶叶、文化消费品牌", "5,000元—1万元", "与传统文化主题契合"],
        ["第四", "文化出版、书店、文创企业", "实物或 5,000元", "可提供图书及文化用品"],
        ["第五", "酒店、餐饮、会务企业", "实物或服务支持", "可以降低活动执行成本"],
    ]
    _money_table(s, 0.45, 1.42, rows_h, rows, [1.45, 4.15, 2.7, 4.5], page, 0.58)
    add_text(
        s, 0.55, 5.15, 12.2, 1.8,
        [
            t("优先联系已经与主办方有合作基础的企业，而不是临时大规模寻找陌生赞助商。招商周期有限，已有信任的企业更容易完成决策。", 15, TEXT, after=6, spc=1.12),
            t("同时避开高风险金融产品营销、夸大健康功效的产品，以及其他可能损害学术活动公信力的商业推广。", 15, TEXT, spc=1.12),
        ],
        page=page,
    )


def slide_in_combo(prs, page, total):
    s = new_slide(prs)
    in_header(s, "推进方式", "用席位和品牌一起促成决定", page, total)
    add_round(s, 0.55, 1.42, 12.2, 1.95, CARD, LINE, 0.08)
    add_text(s, 0.8, 1.52, 11.7, 0.32, [t("建议的组合", 14, CINNABAR, True)], page=page)
    add_text(
        s, 0.8, 1.88, 11.7, 1.35,
        [t("一家企业支付 3 万元，其中包含 4 个正式课程席位、现场品牌展示及活动传播权益。企业既可以派管理层参加学习，也完成一次文化活动支持。这比单纯购买标识展示更容易推动决策。", 15, TEXT, spc=1.15)],
        page=page,
    )
    notes = [
        ("权益对齐", "与 2,399 元购票学员\n享有相同课程权益。\n不降低此前 12 位学员权益。"),
        ("确认节点", "建议 10 月 15 日前确认主要意向。\n确认后即可制作海报和现场物料。"),
        ("长期可能", "若本次跑通，以后可以沉淀为“企业家思想文化讲坛”的赞助体系，而不必每场从零招商。"),
    ]
    for i, (title, desc) in enumerate(notes):
        x = 0.55 + i * 4.15
        add_round(s, x, 3.58, 3.98, 3.25, WASH, radius=0.08)
        add_text(s, x + 0.2, 3.78, 3.55, 0.45, [t(title, 16, CINNABAR, True)], page=page)
        add_text(s, x + 0.2, 4.35, 3.55, 2.2, [t(desc, 15, TEXT, spc=1.2)], page=page)


def slide_in_copy(prs, page, total):
    s = new_slide(prs)
    in_header(s, "转发文案", "可直接发给企业的短文", page, total)
    add_round(s, 0.55, 1.42, 12.2, 5.45, CARD, LINE, 0.06)
    lines = [
        "王德峰教授《传习录》与阳明心学专题讲座",
        "2026年10月31日 · 上海",
        "",
        "五百年前的东方智慧，如何回应今天的人生困境？",
        "本次活动特邀复旦大学哲学学院原教授、博士生导师王德峰教授，围绕《传习录》与阳明心学展开全天专题研修。",
        "拟由复旦大学住房政策研究中心、上海市杨浦区科技企业联合会、首乾书院联合主办。",
        "现诚邀关注传统文化、企业成长与人文精神建设的优秀企业，共同支持本次高品质文化交流活动。",
        "",
        "战略支持单位 30,000元    特别支持单位 10,000元    品牌支持单位 5,000元",
        "同时开放茶饮、图书、礼品、场地及相关服务支持。",
        "合作企业可获得活动品牌展示、企业代表参会席位、现场鸣谢及活动传播等权益。",
        "",
        "期待与优秀企业携手，共赴一场思想与智慧的盛宴。",
    ]
    paras = []
    for i, line in enumerate(lines):
        bold = i in (0, 3, 8)
        size = 16 if i == 0 else 14
        color = CINNABAR if i in (0, 8) else TEXT
        paras.append(t(line if line else " ", size, color, bold, after=2, spc=1.02))
    add_text(s, 0.8, 1.55, 11.75, 5.2, paras, page=page)


INTERNAL = [
    slide_in_cover,
    slide_in_frame,
    slide_in_no_title,
    slide_in_targets,
    slide_in_math,
    slide_in_governance,
    slide_in_priority,
    slide_in_combo,
    slide_in_copy,
]


def build(builders, title: str, subject: str) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = title
    prs.core_properties.subject = subject
    prs.core_properties.author = "活动筹备组"
    prs.core_properties.category = "赞助方案"
    prs.core_properties.language = "zh-CN"
    total = len(builders)
    for i, fn in enumerate(builders, start=1):
        fn(prs, i, total)
    return prs


def overflow_warnings() -> list[str]:
    warnings = []
    for box in BOXES:
        text = box["text"]
        if not text.strip():
            continue
        size = box["size"]
        pad = box["pad"]
        usable_w = box["w"] - pad * 2
        usable_h = box["h"] - 0.06
        if usable_w <= 0.2 or usable_h <= 0.05:
            continue
        char_w = size / 72 * 1.02
        line_h = size / 72 * 1.15
        needed = 0
        for para in text.split("\n"):
            length = len(para.strip())
            if length == 0:
                needed += 1
                continue
            needed += max(1, math.ceil(length / max(1, int(usable_w / char_w))))
        if needed == 1 and (size / 72) <= usable_h:
            continue
        if needed * line_h > usable_h + 0.12:
            snippet = text.replace("\n", " ")[:28]
            warnings.append(
                f"第{box['page']:02d}页可能溢出 ({box['w']:.2f}×{box['h']:.2f}in, {size}pt): {snippet}"
            )
    return warnings


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    external_path = OUT / "王德峰教授专题讲座_企业赞助方案.pptx"
    internal_path = OUT / "王德峰教授专题讲座_赞助招商研讨_内部版.pptx"

    global BOXES
    BOXES = []
    external = build(
        EXTERNAL,
        "《传习录》与阳明心学 · 企业品牌合作赞助方案",
        "王德峰教授专题讲座企业赞助方案（对外）",
    )
    external_warnings = overflow_warnings()
    external.save(external_path)

    BOXES = []
    internal = build(
        INTERNAL,
        "《传习录》与阳明心学 · 赞助招商研讨（内部）",
        "内部资料，请勿对外发送",
    )
    internal_warnings = [f"[内部] {item}" for item in overflow_warnings()]
    internal.save(internal_path)

    print(f"对外版 {len(external.slides)} 页  {external_path}")
    print(f"内部版 {len(internal.slides)} 页  {internal_path}")
    warnings = external_warnings + internal_warnings
    if warnings:
        print("版面预警：")
        for item in warnings:
            print(" -", item)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
