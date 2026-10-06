"""把 book/src 下的 Markdown 分篇合并为完整书稿，并生成 Word 版本。

输出：
- deliverables/以新为鉴_存量时代的城市生存法则.md
- deliverables/以新为鉴_存量时代的城市生存法则.docx
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "book" / "src"
OUT_DIR = ROOT / "deliverables"
TITLE = "以新为鉴_存量时代的城市生存法则"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
RED = RGBColor(0xC8, 0x10, 0x2E)
GREY = RGBColor(0x55, 0x5B, 0x66)
FONT_CN = "Microsoft YaHei"
FONT_BODY = "SimSun"


def merge_markdown() -> str:
    parts = sorted(SRC.glob("*.md"))
    text = "\n\n".join(p.read_text(encoding="utf-8").strip() for p in parts) + "\n"
    OUT_DIR.mkdir(exist_ok=True)
    md_out = text.replace("](../assets/", "](../book/assets/")
    (OUT_DIR / f"{TITLE}.md").write_text(md_out, encoding="utf-8")
    return text


def _set_run_font(run, name=FONT_BODY, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_inline(paragraph, text, size=10.5, color=None, base_bold=False):
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        run = paragraph.add_run(seg)
        bold = base_bold or (i % 2 == 1)
        _set_run_font(run, FONT_CN if bold else FONT_BODY, size, bold,
                      color if color else (NAVY if (i % 2 == 1 and not base_bold) else None))


def _shade(cell, hex_fill):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _paragraph_border(paragraph, color="1F3A5F"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F3F6FA")
    p_pr.append(shd)


def _add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)
    _set_run_font(run, FONT_CN, 9, color=GREY)


def _setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)
    for level, size, color in ((1, 20, NAVY), (2, 15, NAVY), (3, 12.5, RED)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = FONT_CN
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.keep_with_next = True


def _collect_toc(text):
    toc = []
    seen_h1 = False
    for line in text.splitlines():
        if line.startswith("# "):
            if seen_h1:
                toc.append((1, line[2:].strip()))
            seen_h1 = True
        elif line.startswith("## ") and line[3:].strip() != "目录":
            toc.append((2, line[3:].strip()))
    return toc


def _cover(doc, toc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run("以新为鉴"), FONT_CN, 40, True, NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run("存量时代的城市生存法则"), FONT_CN, 20, True, RED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run("一个城市国家穿越五次衰退的民生账本，以及它对中国的启示"),
                  FONT_CN, 12, False, GREY)
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run("书稿第一版初稿｜2026 年 10 月\n数据截至 2026 年 9 月公开发布数据"),
                  FONT_CN, 10.5, False, GREY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(h.add_run("目　录"), FONT_CN, 18, True, NAVY)
    for level, title in toc:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1 if level == 2 else 4)
        p.paragraph_format.line_spacing = 1.2
        if level == 1:
            p.paragraph_format.space_before = Pt(8)
            _set_run_font(p.add_run(title), FONT_CN, 11.5, True, NAVY)
        else:
            p.paragraph_format.left_indent = Cm(0.9)
            _set_run_font(p.add_run(title), FONT_BODY, 10, False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _flush_table(doc, rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, txt in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        _add_inline(cell.paragraphs[0], txt, 9, RGBColor(0xFF, 0xFF, 0xFF), True)
        _shade(cell, "1F3A5F")
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            cell = table.rows[i].cells[j]
            cell.text = ""
            _add_inline(cell.paragraphs[0], row[j] if j < len(row) else "", 9)
            if i % 2 == 0:
                _shade(cell, "F3F6FA")
    doc.add_paragraph()


def build_docx(text: str):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    _setup_styles(doc)
    _cover(doc, _collect_toc(text))
    _add_page_number(sec)

    lines = text.splitlines()
    skip_toc = False
    table_buf = []
    first_h1 = True
    for raw in lines:
        line = raw.rstrip()
        if table_buf and not line.startswith("|"):
            _flush_table(doc, table_buf)
            table_buf = []
        if line.startswith("# "):
            title = line[2:].strip()
            if first_h1:
                first_h1 = False
                continue
            new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
            new_sec.footer.is_linked_to_previous = True
            doc.add_heading(title, level=1)
            skip_toc = False
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            skip_toc = title == "目录"
            if skip_toc:
                continue
            doc.add_heading(title, level=2)
            continue
        if skip_toc:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("|"):
            table_buf.append(line)
            continue
        img = re.match(r"!\[(.*?)\]\((.+?)\)", line)
        if img:
            path = (SRC / img.group(2)).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Cm(16))
            continue
        if line.strip() == "---":
            continue
        if line.startswith(">"):
            content = line.lstrip(">").strip()
            if not content:
                continue
            p = doc.add_paragraph()
            _paragraph_border(p)
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_after = Pt(2)
            bullet = re.match(r"^(-|\d+\.)\s+(.*)", content)
            if bullet:
                prefix = "• " if bullet.group(1) == "-" else bullet.group(1) + " "
                _add_inline(p, prefix + bullet.group(2), 9.5)
            else:
                _add_inline(p, content, 9.5)
            continue
        m = re.match(r"^(\s*)-\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            if m.group(1):
                p.paragraph_format.left_indent = Cm(1.4)
            _add_inline(p, m.group(2))
            continue
        m = re.match(r"^\d+\.\s+(.*)", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            _add_inline(p, line)
            continue
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(21)
        _add_inline(p, line.strip())
    if table_buf:
        _flush_table(doc, table_buf)

    out = OUT_DIR / f"{TITLE}.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    merged = merge_markdown()
    docx_path = build_docx(merged)
    chars = len(re.sub(r"\s", "", merged))
    print(f"md   -> {OUT_DIR / (TITLE + '.md')}")
    print(f"docx -> {docx_path}")
    print(f"非空白字符数：{chars}")
