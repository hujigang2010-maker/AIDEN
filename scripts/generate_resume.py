#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成胡继刚优化版简历（Word + 可预览 HTML）。

主线：产业招商 · 城市更新 · 智能制造。
第八版：对照「优化版」与「优化版-6」优胜劣汰。
版式与 2026 四宗跟第六版；峰会主持、锦天城出海、土木/MBA、求职意向、IELTS 跟第一版。
现任只写复旦大学住房政策研究中心秘书长。万科留在第一页。
一线投资只计 2011–2021。印刷面不出现待补/待核/引进华为。不编造入驻名单。
无甲方盖章原件勿投出本版。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"

NAVY = RGBColor(0x14, 0x2B, 0x4A)
NAVY_HEX = "142B4A"
GOLD = RGBColor(0xC4, 0x8A, 0x2B)
GOLD_HEX = "C48A2B"
GRAY = RGBColor(0x2C, 0x33, 0x3A)
MUTED = RGBColor(0x5C, 0x66, 0x70)
LINE = "D5DDE8"
SOFT = "F3F6FA"


def set_run_font(run, *, name="微软雅黑", size=10.5, bold=False, color=None, east_asia=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    ea = east_asia or name
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), ea)
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="FFFFFF", sz="0"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        el.set(qn("w:space"), "0")
        borders.append(el)
    tcPr.append(borders)


def set_cell_margin(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def v_align(cell, val="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), val)
    tcPr.append(va)


def clear_cell(cell):
    cell.text = ""
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def para_in(cell, text, *, size=10, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text and not cell.paragraphs[0].runs else cell.add_paragraph()
    if cell.paragraphs[0].text == "" and not cell.paragraphs[0].runs:
        p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, name="微软雅黑", size=size, bold=bold, color=color, east_asia="微软雅黑")
    return p


def add_p(doc, text, *, size=10.5, bold=False, color=None, space_before=0, space_after=4, align=None, line=1.12):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, east_asia="微软雅黑")
    return p


def add_hr(doc, color=GOLD_HEX, sz="12"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def section_title(doc, text):
    add_p(doc, text, size=12, bold=True, color=NAVY, space_before=8, space_after=1)
    add_hr(doc, GOLD_HEX, "14")


def bullet(doc, text, *, size=9.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.32)
    pf.first_line_indent = Cm(-0.32)
    pf.space_before = Pt(0)
    pf.space_after = Pt(1.5)
    pf.line_spacing = 1.12
    run = p.add_run("•  " + text)
    set_run_font(run, size=size, color=GRAY, east_asia="微软雅黑")
    return p


def set_col_widths(table, widths_cm):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def job_header(doc, company, title, dates):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [12.6, 5.0])
    left, right = table.rows[0].cells
    clear_cell(left)
    clear_cell(right)
    set_cell_borders(left)
    set_cell_borders(right)
    set_cell_margin(left, 20, 20, 0, 40)
    set_cell_margin(right, 20, 20, 40, 0)
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(company)
    set_run_font(r1, size=10.5, bold=True, color=NAVY, east_asia="微软雅黑")
    r2 = p.add_run("  |  " + title)
    set_run_font(r2, size=9.5, bold=False, color=MUTED, east_asia="微软雅黑")
    p2 = right.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(0)
    r3 = p2.add_run(dates)
    set_run_font(r3, size=9.5, bold=True, color=GOLD, east_asia="微软雅黑")
    return table


def build_docx() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    head = doc.add_table(rows=1, cols=1)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(head, [18.2])
    cell = head.rows[0].cells[0]
    clear_cell(cell)
    shade_cell(cell, NAVY_HEX)
    set_cell_borders(cell, NAVY_HEX, "0")
    set_cell_margin(cell, 90, 90, 140, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("胡继刚")
    set_run_font(r, size=22, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), east_asia="微软雅黑")
    r2 = p.add_run("    HU Jigang")
    set_run_font(r2, size=11, color=RGBColor(0xE8, 0xD5, 0xA3), east_asia="微软雅黑")
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)
    r = p2.add_run("产业招商 · 城市更新 · 智能制造")
    set_run_font(r, size=11, bold=True, color=RGBColor(0xE8, 0xD5, 0xA3), east_asia="微软雅黑")
    p2b = cell.add_paragraph()
    p2b.paragraph_format.space_after = Pt(2)
    r = p2b.add_run("副教授级高级工程师 ｜ 复旦大学硕士研究生 ｜ 原万科上海区域投资副总经理")
    set_run_font(r, size=10, color=RGBColor(0xD5, 0xDE, 0xE8), east_asia="微软雅黑")
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(0)
    r = p3.add_run("现任复旦大学住房政策研究中心秘书长  ｜  杨浦区科技企业联合会执行会长")
    set_run_font(r, size=9, color=RGBColor(0xC5, 0xD0, 0xDC), east_asia="微软雅黑")
    p4 = cell.add_paragraph()
    p4.paragraph_format.space_after = Pt(0)
    r = p4.add_run("1987 年生  ｜  九三学社  ｜  上海  ｜  IELTS 6.0")
    set_run_font(r, size=9, color=RGBColor(0xC5, 0xD0, 0xDC), east_asia="微软雅黑")
    p5 = cell.add_paragraph()
    p5.paragraph_format.space_after = Pt(0)
    r = p5.add_run("18678408669  ｜  262782809@qq.com  ｜  微信 hu262782809")
    set_run_font(r, size=9, color=RGBColor(0xC5, 0xD0, 0xDC), east_asia="微软雅黑")

    add_p(
        doc,
        "求职意向　产业招商 / 城市更新 / 智能制造 / 产业空间运营 / 出海对接　　上海　薪资面议",
        size=9,
        color=MUTED,
        space_before=6,
        space_after=2,
    )

    stats = doc.add_table(rows=1, cols=4)
    stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(stats, [4.55, 4.55, 4.55, 4.55])
    for i, (k, v) in enumerate(
        [("一线投拓", "10 年"), ("2026 已签约", "四宗产业合作"), ("靖江 + 金坛", "43 亿"), ("公开活动", "主办 / 主持")]
    ):
        c = stats.rows[0].cells[i]
        clear_cell(c)
        shade_cell(c, SOFT)
        set_cell_borders(c, LINE, "4")
        set_cell_margin(c, 60, 60, 70, 70)
        para_in(c, k, size=8.5, color=MUTED, space_after=0)
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(v)
        set_run_font(run, size=12, bold=True, color=NAVY, east_asia="微软雅黑")

    add_p(
        doc,
        "土木本科、复旦 MBA。万科上海区域投资副总经理。2011–2021 一线投拓，做成国资合作、政府代建、不良盘活、轻资产委管与物流园，证明能把结构做成。"
        "2021 年起任复旦大学住房政策研究中心秘书长，以杨浦科企联执行会长对接企业、招商与出海。"
        "2026 年已签约产业合作按负责 / 推动 / 辅助 / 方案拟定落笔：写委托与协同，不写入驻名单。"
        "一线投资计至 2021 年。",
        size=9.5,
        color=GRAY,
        space_before=8,
        space_after=2,
        line=1.18,
    )

    section_title(doc, "2026 已签约产业合作")
    add_p(doc, "已签约的是定位与招商委托，不是入驻清单。", size=9, color=MUTED, space_before=1, space_after=3)
    deals = [
        ("项目", "角色", "做什么", "范围"),
        ("创智汇", "负责", "AI 与 IP 定位招商；辅助向智造、出海、人工智能转型", "约 6600㎡ · 杨浦五角场"),
        ("东方枢纽", "推动", "与产业方对接，促成联动招商", "约 143 万㎡ 产业对接范围"),
        ("华为汽车两园区", "辅助", "拿地—规划—定位—招商协同", "静安约 1.5 万㎡、浦东约 2 万㎡"),
        ("森马产业园", "方案拟定", "工业空间定位与招商方案", "约 2 万㎡ 签约切片"),
    ]
    deal_tbl = doc.add_table(rows=len(deals), cols=4)
    deal_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(deal_tbl, [3.4, 2.2, 7.4, 5.2])
    for r, row in enumerate(deals):
        for c, val in enumerate(row):
            cell = deal_tbl.rows[r].cells[c]
            clear_cell(cell)
            set_cell_borders(cell, LINE, "6")
            set_cell_margin(cell, 35, 35, 50, 50)
            if r == 0:
                shade_cell(cell, NAVY_HEX)
                para_in(cell, val, size=8.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if r % 2 == 0:
                    shade_cell(cell, SOFT)
                para_in(cell, val, size=8, bold=(c == 0), color=NAVY if c == 0 else GRAY)

    add_p(
        doc,
        "能力切片　AI 产业落地（创智汇定位招商 · 2026.05 主持商业化圆桌）｜出海对接（北欧会客厅 · 锦天城海外不动产）｜招商落地（2026 四宗委托 · 万科勾地代建收并购）",
        size=9,
        color=GRAY,
        space_before=4,
        space_after=2,
    )

    section_title(doc, "现任")
    job_header(doc, "复旦大学住房政策研究中心", "秘书长", "2021.06 – 至今")
    add_p(doc, "上海 · 中心做研究；科企联做企业、招商、出海与项目", size=9, color=MUTED, space_before=2, space_after=2)
    for x in [
        "城市更新、存量资产与新质产业空间：关注产业如何用空间、政策、资本落地城市。产业方向，不是技术研发岗。",
        "出海：2026.03 北欧创新国际会客厅共同发起，落户杨浦宝龙旭辉广场环创中心（宝龙商办不是雇主）；2024.06 主持锦天城「他山之石——中国投资者的海外不动产战略布局」。",
        "产业活动：2025.05 全球新经济增长引擎峰会主办方代表发言，与竺劲、汪毅、王维军圆桌（人民网上海、中新社上海点名）；2026.05 人工智能商业化落地与硬核投资破局峰会主办方代表并主持圆桌（不是付费落地项目）。",
    ]:
        bullet(doc, x)

    section_title(doc, "社会职务")
    grid = doc.add_table(rows=1, cols=2)
    set_col_widths(grid, [9.1, 9.1])
    for i, (title, body) in enumerate(
        [("杨浦区科技企业联合会", "执行会长 · 企业 / 招商 / 出海 / 项目"), ("复旦不动产资管协会", "创始理事长 · 2018/2019")]
    ):
        cell = grid.rows[0].cells[i]
        clear_cell(cell)
        shade_cell(cell, SOFT)
        set_cell_borders(cell, LINE, "4")
        set_cell_margin(cell, 60, 60, 70, 70)
        v_align(cell, "top")
        para_in(cell, title, size=9.5, bold=True, color=NAVY, space_after=1)
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(body)
        set_run_font(run, size=8, color=MUTED, east_asia="微软雅黑")

    section_title(doc, "专业经历")
    job_header(doc, "万科企业股份有限公司 · 上海区域", "江浙事业部投资副总经理", "2017.03 – 2021.06")
    add_p(doc, "上海。汇报投资合作开发部总经理。团队约 8 人。", size=9, color=MUTED, space_before=2, space_after=2)
    for x in [
        "镇江大港金域蓝湾：收并购，与大港集团各半合作（约 199 亩 / 32 万方）。",
        "嘉兴海宁商住：政府代建落地（18.53 亿元 / 约 23 万㎡，配建安置房）。",
        "上海龙湖滟澜山：法拍不良资产盘活（测算 IRR 约 16.16%）。",
        "靖江印象城：勾地 + 印力委管（约 15 亿元 / 556 亩 / 购物中心约 7 万方，苏中县市首家印象城）；金坛理想城收并购 + 债务重组，底价获取（约 28.05 亿元 / 561 亩）。",
        "主导落地万纬嘉兴平湖（嘉兴首个高标准自动化冷链园）、上海奉贤南桥冷链园；推动进入盐城、台州、常州金坛。",
    ]:
        bullet(doc, x)

    doc.add_page_break()
    section_title(doc, "专业经历（续）")
    job_header(doc, "新城控股集团股份有限公司 · 总部战略投资中心", "投资拓展资深专业经理", "2015.11 – 2017.02")
    for x in [
        "负责山东、江苏等区域战略布局与土地获取。",
        "主导临沂新城吾悦广场勾地，对接市级重点招商政策，系集团首进山东的重资产项目。",
        "以轻资产模式获取青岛新城吾悦广场（品牌输出 + 委管，约 13.68 万㎡）。",
    ]:
        bullet(doc, x)

    job_header(doc, "中南建设集团 / 中南控股集团", "投资拓展经理 / 战略企划管理主管", "2011.06 – 2015.10")
    add_p(doc, "上海 / 南通。2011.06–2012.10 总裁办战略企划；其后新项目发展中心。", size=9, color=MUTED, space_before=2, space_after=2)
    for x in [
        "从集团战略企划到项目投资拓展的完整历练；获取如皋中南世纪城、苏州中南锦苑。",
        "江阴白鹭湾与碧桂园同股同权合作。2014 年底派驻四川任西南拓展负责人。",
        "推进美国养老、日本永旺等国际产业合作对接（推进中，非已签约落地）。",
    ]:
        bullet(doc, x)

    section_title(doc, "代表项目")
    rows = [
        ("项目", "类型", "做法", "规模"),
        ("镇江大港金域蓝湾", "国资合作", "收并购，与大港集团各半", "约 199 亩 / 32 万方"),
        ("嘉兴海宁商住", "政府代建", "代建落地，配建安置房", "18.53 亿 / 约 23 万㎡"),
        ("上海龙湖滟澜山", "不良盘活", "法拍收并购", "测算 IRR 约 16.16%"),
        ("靖江印象城", "轻资产委管", "勾地 + 印力委管 · 苏中县市首家", "约 15 亿 / 556 亩 / 购物中心约 7 万方"),
        ("南桥 / 平湖冷链园", "产业园", "勾地 / 收购改造 · 嘉兴首个高标准冷链园", "物流园"),
        ("金坛理想城", "收并购", "债务重组，底价获取", "约 28.05 亿 / 561 亩"),
        ("临沂新城吾悦广场", "政府招商", "对接市级重点招商政策", "集团首进山东"),
        ("青岛新城吾悦广场", "轻资产", "品牌输出 + 委管", "约 13.68 万㎡"),
        ("如皋世纪城 / 苏州中南锦苑", "早期拓展", "勾地获取；江阴白鹭湾同股同权", "中南阶段代表项目"),
    ]
    proj = doc.add_table(rows=len(rows), cols=4)
    proj.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(proj, [4.6, 2.8, 5.4, 5.4])
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = proj.rows[r].cells[c]
            clear_cell(cell)
            set_cell_borders(cell, LINE, "6")
            set_cell_margin(cell, 35, 35, 50, 50)
            if r == 0:
                shade_cell(cell, NAVY_HEX)
                para_in(cell, val, size=8.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if r % 2 == 0:
                    shade_cell(cell, SOFT)
                para_in(cell, val, size=8, bold=(c == 0), color=NAVY if c == 0 else GRAY)

    section_title(doc, "教育背景")
    job_header(doc, "复旦大学", "硕士研究生 · 工商管理（财务金融方向）", "2018.09 – 2021.06")
    add_p(doc, "在职攻读，与万科上海区域任职并行。", size=9, color=MUTED, space_before=1, space_after=3)
    job_header(doc, "中国海洋大学", "土木工程 · 工学学士", "2007.09 – 2011.07")

    section_title(doc, "资质")
    add_p(
        doc,
        "副教授级高级工程师；会计师中级；PMP；住建委项目管理工程师；IELTS 6.0；普通话一级甲等；BIM；上海市人才引进认证资质。",
        size=9.5,
        color=GRAY,
        space_after=2,
    )
    add_p(
        doc,
        "论文《万科物流地产平台业务发展战略研究》（知网）；《房屋建筑工程人工智能技术的应用》（《住宅与房地产》）。实用新型专利 2 项（建筑工程管理用功能脚架、防护围栏）。",
        size=9.5,
        color=GRAY,
        space_after=2,
    )
    add_p(
        doc,
        "研修：上海交通大学城市治理数字化转型高级研修班；复旦–花旗银行实践课程；复旦–怡安人力与风险管理。",
        size=9.5,
        color=GRAY,
        space_after=2,
    )
    add_p(
        doc,
        "投递方向：产业招商 · 城市更新 · 智能制造 · 产业空间运营 · 出海对接。"
        "能力：投资拓展 → 政企招商 → 产业组织 → 智库与资源配置。研究方向：AI 产业如何落地城市。",
        size=9,
        color=MUTED,
        space_before=4,
        space_after=0,
    )

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "胡继刚-简历-优化版.docx"
    doc.save(path)
    return path


def write_html() -> tuple[Path, Path]:
    OUT.mkdir(parents=True, exist_ok=True)
    tpl = Path(__file__).resolve().parent / "templates"
    full = OUT / "胡继刚-简历-优化版.html"
    brief = OUT / "胡继刚-简历-一页精华.html"
    full.write_text((tpl / "resume-full.html").read_text(encoding="utf-8"), encoding="utf-8")
    brief.write_text((tpl / "resume-brief.html").read_text(encoding="utf-8"), encoding="utf-8")
    return full, brief


def write_notes() -> Path:
    text = """# 胡继刚简历优化说明（第八版 · 两版优胜劣汰）

对照「优化版」（AI × 出海 × 招商）和「优化版-6」（2026 已签约四宗）合并。页眉仍是 **产业招商 · 城市更新 · 智能制造**。现任只写复旦大学住房政策研究中心秘书长。万科留在第一页。2026 四宗按负责 / 推动 / 辅助 / 方案拟定落笔，不写「引进华为」。

**闸门：有甲方盖章原件，国资更新和产业导入岗可以开始投；没有文件，先别对外发这版。** 面试逐条口径见 `output/2026已签约产业合作-面试口径.md`。

## 优胜劣汰

- 采用第六版：版式、四宗主表、诚实动词、万科项目顺序、点名证据。
- 采用第一版：2026.05 AI 峰会主持、2024.06 锦天城出海、土木本科/复旦 MBA、求职意向、IELTS 6.0、能力切片。
- 淘汰：工商联现任秘书长、山东商会、联考面试官、评审专家、「超级孵化器」、把高工写进任职标题。

## 印刷面

- 创智汇约 6600㎡：负责 AI 与 IP 定位招商；载体在杨浦五角场，不是创智天地。
- 东方枢纽约 143 万㎡：推动产业对接、联动招商；公开东站约 133 万㎡，口径见面试稿。
- 华为汽车两园区：辅助静安约 1.5 万㎡、浦东约 2 万㎡；不写成已引进。
- 森马产业园约 2 万㎡：方案拟定；吴泾公开约 22 万㎡，2 万㎡按签约切片写。
- 2026.05 AI 峰会：主办方代表并主持圆桌，不是付费落地。
- 2024.06 锦天城：主持海外不动产战略布局。

## 硬约束

一线投资计至 2021；副教授级高工归职称；房商会不写现任秘书长（官网为李祥）；不写身份证住址配偶；不把 REITs / 保租房收储写成业绩；中南保留江阴白鹭湾同股同权与 2014 西南拓展；不把集团已披露企业写成个人名单；不说已经招满。

## 文件

- `output/胡继刚-简历-优化版.html` / `.pdf` / `.docx`
- `output/胡继刚-简历-一页精华.html` / `.pdf`
- `output/2026已签约产业合作-面试口径.md`
"""
    path = OUT / "简历优化说明.md"
    path.write_text(text, encoding="utf-8")
    return path


def write_interview() -> Path:
    text = """# 2026 已签约产业合作 · 面试口径

**闸门：有甲方盖章原件，国资更新和产业导入岗可以开始投；没有文件，先别对外发第八版简历。**

印刷面只写负责 / 推动 / 辅助 / 方案拟定。下面四条是面试会被追问的内容。甲方抬头、签约日、角色以盖章原件为准；仓库里目前只有草案，日期空白。**不要当场编日期。** 原件没带在身上，就说「合同原件在包里 / 会后补扫描件」，不要用「待核」二字。

---

## 1. 四份合同：甲方抬头、签约日、角色

填空表。面试前用原件抄一遍，抄完再投。

| 项目 | 甲方抬头（原件） | 签约日（原件） | 你的角色 | 草案里看到的、不能当原件用 |
| --- | --- | --- | --- | --- |
| 创智汇约 6600㎡ |  |  | 负责 · 联合招商 / 平台方（科企联侧） | 草案甲乙方是杨浦同谱会 × 杨浦科企联；业主更像轻舟，不是创智天地。日期空白。 |
| 东方枢纽约 143 万㎡ |  |  | 推动 · 平台方协同 | 内部方案曾按复旦中心 / 市科企联 / 东方枢纽三方活动合作写。不是枢纽集团编制。 |
| 华为汽车 · 静安约 1.5 万㎡ |  |  | 辅助 · 顾问 / 联合招商（以原件为准） | 业主侧草案指向冠松静安永和社区 075b-07、C6、地上约 1.5 万㎡。日期空白。 |
| 华为汽车 · 浦东约 2 万㎡ |  |  | 辅助 | 仓库没有与静安同级的盖章对应。面积按签约切片说，主体以原件为准。 |
| 森马产业园约 2 万㎡ |  |  | 方案拟定 · 顾问 | 草案抬头写过「森马集团股份有限公司」；上市主体是浙江森马服饰股份有限公司，项目公司是上海森马服饰有限公司。先看原件用哪一个。 |

角色三个词怎么选：

- **顾问**：出定位、方案、对接名单，不负责盖租赁章。
- **联合招商**：和业主/平台一起招，佣金或分成以合同为准。
- **平台方**：科企联 / 中心侧协同，不是业主，也不是入驻企业。

被问「合同在哪」：拿出原件或扫描件。拿不出，就不要把这一版发出去。

---

## 2. 华为两条：哪家主体、推进到哪一步

**开口第一句：我没有引进华为。简历写的是辅助。**

- 招商标的曾经按华为车 BU / 鸿蒙智行做锚定设想，也可以是合作车企通道。面试时只说原件和当时纪要里写死的那一家，不要三家一起说。
- 静安：辅助一栋约 1.5 万㎡ 的拿地—规划—定位—招商协同。公开检索几乎看不到「静安华为汽车园区」。上海华为公开重心在青浦练秋湖。对方用公开报道追问，就承认：公开重心不在静安，所以我只写辅助，不写落地。
- 浦东：约 2 万㎡ 按签约切片写。没有原件就不要补充园区名称、不要说已经拿地完工。
- 现在推进到：规划与定位协同、对接名单阶段。不是入驻交割，不是政府签约仪式已经办完。
- 禁止句：「我把华为引进来了」「华为已经入驻静安/浦东」「这是华为上海第二个总部」。

---

## 3. 东方枢纽：143 万㎡是哪一块、和枢纽集团什么关系

三套数字不要混成一句「我负责 143 万㎡开发」。

| 数字 | 是什么 | 不是什么 |
| --- | --- | --- |
| 约 143 万㎡ | 简历上的**产业对接范围**（签约切片） | 不是你开发、代建或销售的建筑总量 |
| 公开约 133 万㎡ | 东站建设/工程的公开报道口径 | 不是你的个人 KPI |
| 内部方案里出现过的 A 片区办公约 30 万㎡ | 某一类办公产品 | 不能拿来否定 143，也不能拿来对外改口 |

和枢纽集团的关系：推动与产业方对接，促成联动招商。不是枢纽集团编制内招商负责人，不是拿地主体。

京东、波士顿科学是枢纽集团对外口径，不是你的招商名单。被点名就说：那是业主/集团已披露的方向，不是我签下来的客户。

被问「为什么简历 143、报道 133」：范围定义不同。133 是东站公开建设体量；143 是委托对接范围，以合同附图为准。不要说「媒体写少了」。

---

## 4. 创智汇、森马：定位招商权，还是已经有入驻

**已签约的是定位与招商委托，不是入驻清单。没有经手企业名单，就不要报已经招满。**

创智汇：

- 杨浦五角场创智汇，约 6600㎡（草案口径：3F 约 2850 + 5F 约 3670）。
- 不说创智天地。创智天地是另一块牌子，对不上署名。
- 负责 AI 与 IP 定位招商，辅助向智造、出海、人工智能转型。
- 园区历史上出现过的租户（例如公开检索能看到的旧品牌）不是你招来的，不要认领。

森马：

- 吴泾改扩建公开约 22 万㎡、投资约 15 亿。那是业主项目全盘。
- 你的签约切片是约 2 万㎡ 工业空间的定位与招商方案拟定。
- 方案拟定 ≠ 已经招满。有入驻名单再补；没有就停在方案和委托。

---

## 30 秒开场（可背）

万科上海区域投资副总经理，2011 到 2021 一线投拓，国资合作、代建、不良、轻资产、物流园，证明能把结构做成。2021 年起任复旦住房政策研究中心秘书长。2026 年四宗已签约产业合作：创智汇负责定位招商，东方枢纽推动产业对接，华为汽车两园区辅助全流程协同，森马产业园做方案拟定。写的是委托，不是入驻。一线投资计到 2021。

---

## 被追问时的停句

- 合同抬头和日期，以原件为准，我不口头估。
- 这四条是定位招商权和方案权，不是已经招满。
- 公开报道和签约切片不是同一个口径，我按切片说，按报道认公开事实。
- 华为这条听成引进，就是听错了；我的动词是辅助。
"""
    path = OUT / "2026已签约产业合作-面试口径.md"
    path.write_text(text, encoding="utf-8")
    return path


def export_pdf(html_path: Path, pdf_path: Path) -> None:
    import subprocess
    import tempfile
    import time

    html_uri = html_path.resolve().as_uri()
    if pdf_path.exists():
        pdf_path.unlink()
    with tempfile.TemporaryDirectory(prefix="chrome-resume-") as tmp:
        cmd = [
            "google-chrome",
            "--headless=new",
            "--no-sandbox",
            "--disable-gpu",
            "--disable-dev-shm-usage",
            "--disable-extensions",
            "--disable-background-networking",
            f"--user-data-dir={tmp}",
            "--no-first-run",
            "--no-pdf-header-footer",
            "--virtual-time-budget=8000",
            f"--print-to-pdf={pdf_path}",
            html_uri,
        ]
        try:
            subprocess.run(cmd, check=True, timeout=20)
        except subprocess.TimeoutExpired:
            time.sleep(0.5)
            if not pdf_path.exists() or pdf_path.stat().st_size < 1000:
                raise
            print(f"chrome 打印后未退出，但已生成：{pdf_path.name}")


def main():
    docx_path = build_docx()
    html_full, html_brief = write_html()
    notes = write_notes()
    interview = write_interview()
    pdf_full = OUT / "胡继刚-简历-优化版.pdf"
    pdf_brief = OUT / "胡继刚-简历-一页精华.pdf"
    export_pdf(html_full, pdf_full)
    export_pdf(html_brief, pdf_brief)
    print(f"docx: {docx_path}")
    print(f"html: {html_full}")
    print(f"html_brief: {html_brief}")
    print(f"pdf: {pdf_full}")
    print(f"pdf_brief: {pdf_brief}")
    print(f"notes: {notes}")
    print(f"interview: {interview}")


if __name__ == "__main__":
    main()
