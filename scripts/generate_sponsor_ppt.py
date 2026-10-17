#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""王德峰老师公开讲座赞助方案，对外商务幻灯片。与 Word 同一口径。"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "output"
PATH = OUT / "王德峰老师公开讲座_赞助方案_对外商务版.pptx"

NAVY = RGBColor(0x1A, 0x3A, 0x5C)
NAVY_DEEP = RGBColor(0x12, 0x2A, 0x44)
GOLD = RGBColor(0x8B, 0x69, 0x14)
GOLD_SOFT = RGBColor(0xC4, 0xA3, 0x5A)
CREAM = RGBColor(0xF7, 0xF3, 0xEA)
INK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE4, 0xDB, 0xC8)
CARD = RGBColor(0xFF, 0xFC, 0xF7)
SOFT = RGBColor(0xEE, 0xF3, 0xF8)

W = Inches(13.333)
H = Inches(7.5)
FONT = "微软雅黑"
SLIDES = 8


def set_run_font(run, size, bold, color):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    r_pr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        node = r_pr.find(qn(f"a:{tag}"))
        if node is None:
            node = etree.SubElement(r_pr, qn(f"a:{tag}"))
        node.set("typeface", FONT)


def no_line(shape):
    shape.line.fill.background()


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    no_line(shape)
    return shape


def add_round(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    no_line(shape)
    if shape.adjustments:
        shape.adjustments[0] = 0.08
    return shape


def _box(slide, left, top, width, height, anchor):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    anchor_map = {
        MSO_ANCHOR.TOP: "t",
        MSO_ANCHOR.MIDDLE: "ctr",
        MSO_ANCHOR.BOTTOM: "b",
    }
    frame._txBody.bodyPr.set("anchor", anchor_map[anchor])
    return frame


def add_text(slide, left, top, width, height, text, size, bold, color, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    frame = _box(slide, left, top, width, height, anchor)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    set_run_font(run, size, bold, color)
    return frame


def add_lines(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    frame = _box(slide, left, top, width, height, anchor)
    for index, (text, size, bold, color, space_before) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(space_before)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = text
        set_run_font(run, size, bold, color)
    return frame


def chrome(slide, title, page):
    add_rect(slide, 0, 0, W, H, CREAM)
    add_rect(slide, 0, 0, W, Inches(0.08), NAVY)
    add_text(slide, Inches(0.48), Inches(0.24), Inches(10.5), Inches(0.46), title, 26, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.5), Inches(0.74), Inches(1.15), Inches(0.035), GOLD)
    add_text(
        slide,
        Inches(0.48),
        Inches(7.12),
        Inches(9.2),
        Inches(0.28),
        "王德峰老师公开讲座  ·  赞助方案",
        11,
        False,
        MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        Inches(11.3),
        Inches(7.12),
        Inches(1.5),
        Inches(0.28),
        f"{page}  /  {SLIDES}",
        11,
        False,
        MUTED,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def style_cell(cell, text, size, bold, color, fill, align=PP_ALIGN.CENTER):
    cell.text = ""
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    set_run_font(run, size, bold, color)


def set_cell_borders(cell, color="E4DBC8", width_pt=0.75):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    lines = tc_pr.find(qn("a:lnL"))
    # Table cell borders live on tcPr as lnL/lnR/lnT/lnB in the c namespace? 
    # python-pptx uses a:lnL inside tcPr for table cells in the drawingml table model.
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        existing = tc_pr.find(qn(f"a:{edge}"))
        if existing is not None:
            tc_pr.remove(existing)
        node = etree.SubElement(tc_pr, qn(f"a:{edge}"))
        node.set("w", str(int(width_pt * 12700)))
        node.set("cap", "flat")
        node.set("cmpd", "sng")
        node.set("algn", "ctr")
        fill = etree.SubElement(node, qn("a:solidFill"))
        srgb = etree.SubElement(fill, qn("a:srgbClr"))
        srgb.set("val", color)
        etree.SubElement(node, qn("a:prstDash")).set("val", "solid")
    del lines


def add_table(slide, left, top, width, height, headers, rows, col_widths, font=12):
    table_shape = slide.shapes.add_table(1 + len(rows), len(headers), left, top, width, height)
    table = table_shape.table
    table.rows[0].height = Inches(0.38)
    for index, col_width in enumerate(col_widths):
        table.columns[index].width = col_width
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        style_cell(cell, header, font, True, WHITE, NAVY)
        set_cell_borders(cell, "1A3A5C")
    for r_index, row in enumerate(rows):
        table.rows[r_index + 1].height = Inches(0.36)
        fill = SOFT if r_index % 2 else WHITE
        for c_index, value in enumerate(row):
            cell = table.cell(r_index + 1, c_index)
            align = PP_ALIGN.LEFT if c_index == 0 else PP_ALIGN.CENTER
            style_cell(cell, value, font, c_index == 0, NAVY if c_index == 0 else INK, fill, align)
            set_cell_borders(cell)
    return table


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, NAVY)
    add_rect(slide, 0, 0, W, Inches(0.12), GOLD)
    add_rect(slide, 0, Inches(7.38), W, Inches(0.12), GOLD)
    add_text(slide, Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.4), "赞 助 方 案", 18, True, GOLD_SOFT, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        Inches(0.8),
        Inches(2.05),
        Inches(11.7),
        Inches(1.15),
        "王德峰老师公开讲座",
        44,
        True,
        WHITE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, Inches(6.05), Inches(3.35), Inches(1.25), Inches(0.035), GOLD)
    add_lines(
        slide,
        Inches(0.8),
        Inches(3.6),
        Inches(11.7),
        Inches(1.3),
        [
            ("2026年10月31日（周六）9:00–12:00", 20, False, RGBColor(0xE8, 0xE0, 0xD0), 0),
            ("场地另行通知    ·    约 150 座", 20, False, RGBColor(0xE8, 0xE0, 0xD0), 8),
        ],
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        Inches(0.8),
        Inches(6.35),
        Inches(11.7),
        Inches(0.4),
        "联合主办    上海市杨浦区科技企业联合会  ×  首乾书院",
        16,
        False,
        GOLD_SOFT,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_event(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "这场讲座", 2)
    facts = [
        ("主讲", "王德峰，复旦大学哲学学院教授"),
        ("时间", "2026年10月31日（周六）9:00–12:00"),
        ("场地", "具体地点待定，确定后另行通知"),
        ("规模", "按约 150 座准备。对外售票 128 张"),
        ("主办", "上海市杨浦区科技企业联合会、首乾书院"),
        ("听众", "企业负责人、管理者，以及关心哲学与中国思想的公众"),
    ]
    top = 1.05
    for index, (label, value) in enumerate(facts):
        y = Inches(top + index * 0.72)
        add_round(slide, Inches(0.5), y, Inches(12.35), Inches(0.64), CARD)
        add_rect(slide, Inches(0.5), y, Inches(0.08), Inches(0.64), GOLD if index % 2 == 0 else NAVY)
        add_text(slide, Inches(0.85), y, Inches(1.5), Inches(0.64), label, 16, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(2.5), y, Inches(10.0), Inches(0.64), value, 16, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(
        slide,
        Inches(0.55),
        Inches(5.5),
        Inches(12.2),
        Inches(1.4),
        "品牌出现在报名页、现场和活动照片中。讲座内容不冠名，不出售授课录像。\n现场不设展位，也不做产品推介。一家企业选择一档。",
        15,
        False,
        MUTED,
    )


def slide_agenda(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "上午议程", 3)
    rows = [
        ("08:30–09:00", "入场签到", "听众看到背景板和报名页上的标识。首席支持可在签到处放一份资料"),
        ("09:00–09:10", "开场", "主持人鸣谢。首席支持不超过 60 秒致辞，战略支持单独念到名称"),
        ("09:10–10:40", "公开讲座", "王德峰老师主讲。此段不插播企业内容"),
        ("10:40–11:20", "现场交流", "听众提问。赞助席位按票档就座，不单独设置提问权"),
        ("11:20–12:00", "签书与合影", "前排票听众签书或签名。企业使用现场照片，不使用授课画面"),
    ]
    top = 1.15
    for index, (when, name, how) in enumerate(rows):
        y = Inches(top + index * 1.05)
        add_round(slide, Inches(0.5), y, Inches(2.55), Inches(0.9), NAVY)
        add_text(slide, Inches(0.62), y, Inches(2.3), Inches(0.9), when, 16, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_round(slide, Inches(3.2), y, Inches(9.65), Inches(0.9), CARD)
        add_text(slide, Inches(3.45), y + Inches(0.08), Inches(9.2), Inches(0.36), name, 16, True, NAVY)
        add_text(slide, Inches(3.45), y + Inches(0.42), Inches(9.2), Inches(0.4), how, 13, False, INK)


def slide_tiers(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "六档里选一档", 4)
    cards = [
        ("首席支持", "68,000 元", "1 家", ["主标识", "60 秒致辞", "前排票 4 张"]),
        ("战略支持", "48,000 元", "1 家", ["次主标识", "单独口播", "前排票 2 张"]),
        ("合作支持", "38,000 元", "2 家", ["第二序列标识", "口播名称", "前排票 2 张"]),
        ("品牌支持", "28,000 元", "2 家", ["背景板品牌列", "集体鸣谢", "中排票 2 张"]),
        ("鸣谢支持", "18,000 元", "3 家", ["鸣谢区名称", "集体鸣谢", "中排票 2 张"]),
        ("友好支持", "9,800 元", "4 家", ["报名页名单", "列入鸣谢", "后排票 2 张"]),
    ]
    card_w = 3.95
    card_h = 2.35
    gap_x = 0.18
    gap_y = 0.16
    origin_x = 0.5
    origin_y = 1.05
    for index, (name, price, count, points) in enumerate(cards):
        col = index % 3
        row = index // 3
        x = Inches(origin_x + col * (card_w + gap_x))
        y = Inches(origin_y + row * (card_h + gap_y))
        add_round(slide, x, y, Inches(card_w), Inches(card_h), CARD)
        add_text(slide, x + Inches(0.22), y + Inches(0.16), Inches(2.3), Inches(0.38), name, 16, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(
            slide,
            x + Inches(2.3),
            y + Inches(0.16),
            Inches(1.4),
            Inches(0.38),
            count,
            13,
            False,
            MUTED,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text(slide, x + Inches(0.22), y + Inches(0.58), Inches(3.5), Inches(0.5), price, 26, True, GOLD, anchor=MSO_ANCHOR.MIDDLE)
        add_lines(
            slide,
            x + Inches(0.22),
            y + Inches(1.18),
            Inches(3.45),
            Inches(1.0),
            [(point, 14, False, INK, 3) for point in points],
        )
    add_text(
        slide,
        Inches(0.5),
        Inches(6.72),
        Inches(12.3),
        Inches(0.32),
        "首席同一行业只此一家。战略不与首席同业。两家合作支持之间避开直接竞品。",
        13,
        False,
        MUTED,
    )


def slide_matrix(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "权益对照", 5)
    headers = ["权益", "首席", "战略", "合作", "品牌", "鸣谢", "友好"]
    rows = [
        ["金额", "6.8万", "4.8万", "3.8万", "2.8万", "1.8万", "0.98万"],
        ["名额", "1 家", "1 家", "2 家", "2 家", "3 家", "4 家"],
        ["报名页", "主位置", "次主位置", "第二序列", "品牌列", "鸣谢列", "名单"],
        ["背景板", "主标识", "次主标识", "有", "品牌列", "鸣谢区", "不单列"],
        ["开场", "口播+60秒", "单独口播", "口播名称", "集体鸣谢", "集体鸣谢", "列入名单"],
        ["门票", "前排 4 张", "前排 2 张", "前排 2 张", "中排 2 张", "中排 2 张", "后排 2 张"],
        ["签到资料", "1 份", "—", "—", "—", "—", "—"],
        ["现场照片", "可用", "可用", "可用", "可用", "可用", "可用"],
        ["授课录像", "不含", "不含", "不含", "不含", "不含", "不含"],
        ["同业", "独家", "不与首席同业", "避开竞品", "不限", "不限", "不限"],
    ]
    widths = [Inches(1.7)] + [Inches(1.72)] * 6
    add_table(slide, Inches(0.48), Inches(1.05), Inches(12.02), Inches(5.15), headers, rows, widths, font=13)
    add_text(
        slide,
        Inches(0.5),
        Inches(6.55),
        Inches(12.3),
        Inches(0.48),
        "门票从现有票档划出，不另加座。六档全部确认时，划出前排 10 张、中排 10 张、后排 8 张。未确认的档位不预留空座。",
        13,
        False,
        MUTED,
    )


def slide_presence(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "现场怎么出现", 6)
    rows = [
        ("报名页", "讲座信息下方，按首席、战略、合作、品牌、鸣谢、友好排列"),
        ("海报", "底部赞助列。讲座名称保持「王德峰老师公开讲座」"),
        ("签到台", "仅首席支持可放一份资料"),
        ("背景板", "主办双方名称在上，赞助标识在下。场地确定后，场地支持单独列明"),
        ("开场", "鸣谢不超过 3 分钟。首席致辞不超过 60 秒"),
        ("讲座过程", "不插播广告，不发放课程资料"),
        ("结束后", "企业获得现场照片。照片不含可单独传播的授课片段"),
    ]
    top = 1.05
    for index, (place, what) in enumerate(rows):
        y = Inches(top + index * 0.78)
        add_round(slide, Inches(0.5), y, Inches(2.3), Inches(0.68), NAVY)
        add_text(slide, Inches(0.62), y, Inches(2.06), Inches(0.68), place, 15, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_round(slide, Inches(2.95), y, Inches(9.9), Inches(0.68), CARD)
        add_text(slide, Inches(3.2), y, Inches(9.45), Inches(0.68), what, 15, False, INK, anchor=MSO_ANCHOR.MIDDLE)


def slide_limits(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "边界与时间", 7)
    add_round(slide, Inches(0.45), Inches(1.05), Inches(6.35), Inches(5.45), CARD)
    add_text(slide, Inches(0.7), Inches(1.2), Inches(5.9), Inches(0.4), "方案里没有的", 18, True, NAVY)
    limits = [
        "不以企业名称替换讲座名称",
        "不提供授课录像、回放和课程剪辑",
        "不安排老师单独代言或出席企业内场",
        "不在现场销售其他课程",
        "赞助款不退，缺席不折成门票",
        "不写成协办单位或联合主办",
    ]
    add_lines(
        slide,
        Inches(0.7),
        Inches(1.75),
        Inches(5.85),
        Inches(4.8),
        [(f"·  {item}", 15, False, INK, 10) for item in limits],
    )
    times = [
        ("场地确定后", "另行通知具体地点"),
        ("10月22日", "确认档位、标识、致辞人"),
        ("10月24日", "付清赞助款，交门票名单"),
        ("10月28日", "定稿背景板和鸣谢词"),
        ("10月31日", "9:00 讲座开始"),
        ("11月5日前", "发送现场照片"),
    ]
    for index, (when, what) in enumerate(times):
        y = Inches(1.05 + index * 0.98)
        add_round(slide, Inches(7.0), y, Inches(5.85), Inches(0.88), CARD if index % 2 == 0 else SOFT)
        add_text(slide, Inches(7.2), y + Inches(0.06), Inches(5.45), Inches(0.34), when, 14, True, GOLD)
        add_text(slide, Inches(7.2), y + Inches(0.4), Inches(5.45), Inches(0.38), what, 15, False, INK)


def slide_confirm(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, "选定后回传", 8)
    add_text(
        slide,
        Inches(0.5),
        Inches(0.95),
        Inches(12.3),
        Inches(0.4),
        "勾选一档回复。需要盖章的，用同名 Word 文末确认表。付款账户在确认函中告知。",
        15,
        False,
        INK,
    )
    options = [
        ("首席支持", "68,000 元", "1 家"),
        ("战略支持", "48,000 元", "1 家"),
        ("合作支持", "38,000 元", "2 家"),
        ("品牌支持", "28,000 元", "2 家"),
        ("鸣谢支持", "18,000 元", "3 家"),
        ("友好支持", "9,800 元", "4 家"),
    ]
    for index, (name, price, count) in enumerate(options):
        col = index % 2
        row = index // 2
        x = Inches(0.5 + col * 4.35)
        y = Inches(1.55 + row * 0.85)
        add_round(slide, x, y, Inches(4.15), Inches(0.72), CARD)
        add_text(slide, x + Inches(0.16), y, Inches(0.45), Inches(0.72), "□", 20, False, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + Inches(0.6), y, Inches(1.7), Inches(0.72), name, 15, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + Inches(2.25), y, Inches(1.75), Inches(0.72), f"{price}  {count}", 13, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    add_round(slide, Inches(9.15), Inches(1.55), Inches(3.7), Inches(3.4), NAVY)
    add_lines(
        slide,
        Inches(9.35),
        Inches(1.75),
        Inches(3.3),
        Inches(3.05),
        [
            ("请一并告知", 16, True, GOLD_SOFT, 0),
            ("公司名称", 15, False, WHITE, 12),
            ("行业与对外名称", 15, False, WHITE, 6),
            ("出席人、职务、手机", 15, False, WHITE, 6),
            ("发票抬头和税号", 15, False, WHITE, 6),
            ("授权人签字", 15, False, WHITE, 6),
        ],
    )
    add_text(
        slide,
        Inches(0.5),
        Inches(4.35),
        Inches(8.4),
        Inches(1.2),
        "首席支持请注明致辞人。\n10月22日后确认的鸣谢、友好支持，不保证补印已送印的海报。",
        14,
        False,
        MUTED,
    )
    add_rect(slide, Inches(0.5), Inches(5.7), Inches(12.35), Inches(1.2), NAVY)
    add_text(slide, Inches(0.75), Inches(5.82), Inches(11.8), Inches(0.4), "上海市杨浦区科技企业联合会", 16, True, WHITE)
    add_text(slide, Inches(0.75), Inches(6.25), Inches(11.8), Inches(0.45), "联系人：胡继刚　　联合主办：首乾书院", 15, False, GOLD_SOFT)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "王德峰老师公开讲座赞助方案"
    prs.core_properties.subject = "对外商务版"
    prs.core_properties.author = "上海市杨浦区科技企业联合会"
    slide_cover(prs)
    slide_event(prs)
    slide_agenda(prs)
    slide_tiers(prs)
    slide_matrix(prs)
    slide_presence(prs)
    slide_limits(prs)
    slide_confirm(prs)
    OUT.mkdir(parents=True, exist_ok=True)
    prs.save(PATH)
    return PATH


if __name__ == "__main__":
    print(build())
