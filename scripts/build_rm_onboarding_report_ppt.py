#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成《客户经理新员工培养阶段工作汇报》PPT。

风格：简约商务、正式清爽；主色为紫 / 黄 / 橙；无装饰配图。
页数：5 页（封面 + 调研条形图 + 情绪饼图 + 合格率增量规则空位 + 团队激励规则空位）。

运行：
    python3 scripts/build_rm_onboarding_report_ppt.py
输出：
    deliverables/客户经理新员工培养阶段工作汇报.pptx
"""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "客户经理新员工培养阶段工作汇报.pptx"

# ---------------------------------------------------------------------------
# 视觉主题：紫、黄、橙 · 分行领导汇报
# ---------------------------------------------------------------------------
PURPLE_INK = RGBColor(0x2B, 0x16, 0x4A)
PURPLE = RGBColor(0x5B, 0x3A, 0x9A)
PURPLE_MID = RGBColor(0x7A, 0x58, 0xB8)
PURPLE_SOFT = RGBColor(0xED, 0xE7, 0xF6)
PURPLE_PALE = RGBColor(0xF6, 0xF2, 0xFB)
PURPLE_LINE = RGBColor(0xE4, 0xDC, 0xF0)
YELLOW = RGBColor(0xF0, 0xC4, 0x1A)
YELLOW_DEEP = RGBColor(0xC9, 0x96, 0x10)
YELLOW_SOFT = RGBColor(0xFF, 0xF6, 0xD4)
ORANGE = RGBColor(0xE8, 0x78, 0x2A)
ORANGE_DEEP = RGBColor(0xC4, 0x56, 0x16)
ORANGE_SOFT = RGBColor(0xFF, 0xEF, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x2A, 0x24, 0x35)
MUTED = RGBColor(0x6E, 0x65, 0x7C)
TRACK = RGBColor(0xEE, 0xEA, 0xF4)

FONT = "微软雅黑"
TOTAL = 5
FOOTER = "人力资源部  ·  客户经理新员工培养阶段工作汇报  ·  内部材料"

# 六大帮带角色：5 分制，按得分从高到低
ROLE_SCORES = [
    ("业务技能传授", 4.62, "优势"),
    ("合规风险把关", 4.41, "优势"),
    ("日常答疑陪伴", 4.18, ""),
    ("客户拓展陪访", 3.72, ""),
    ("目标节奏辅导", 3.45, "薄弱"),
    ("心态激励疏导", 3.16, "薄弱"),
]
SCORE_MAX = 5.0
BAR_COLORS = [PURPLE, PURPLE_MID, YELLOW, ORANGE, ORANGE, ORANGE_DEEP]

# 五类工作状态占比（合计 100%）
MOOD_SHARE = [
    ("充实成长", 34, PURPLE),
    ("干劲十足", 25, YELLOW),
    ("繁忙但低效", 17, ORANGE),
    ("迷茫不清楚", 14, PURPLE_MID),
    ("疲惫内耗、压力大", 10, ORANGE_DEEP),
]


def _set_font(run, size: float, color: RGBColor, bold: bool = False, font: str = FONT) -> None:
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        if node is None:
            node = rPr.makeelement(qn(tag), {})
            rPr.append(node)
        node.set("typeface", font)


def add_rect(slide, x, y, w, h, color, *, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def add_round(slide, x, y, w, h, color, *, adj=0.08):
    sp = add_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    try:
        sp.adjustments[0] = adj
    except Exception:
        pass
    return sp


def add_oval(slide, x, y, w, h, color):
    return add_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.OVAL)


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text: str,
    *,
    size=16,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    font=FONT,
):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        _set_font(run, size, color, bold, font)
    return tb


def content_chrome(slide, title: str, page: int) -> None:
    """内容页顶栏 + 页脚。"""
    add_rect(slide, 0, 0, SW, Inches(0.08), PURPLE)
    add_rect(slide, 0, 0, Inches(0.10), SH, YELLOW)
    add_rect(slide, Inches(0.10), 0, Inches(0.05), SH, ORANGE)
    add_text(
        slide,
        Inches(0.55),
        Inches(0.22),
        Inches(11.0),
        Inches(0.22),
        "人力资源部  |  新员工培养",
        size=11,
        color=PURPLE_MID,
        bold=True,
    )
    add_text(
        slide,
        Inches(0.55),
        Inches(0.44),
        Inches(11.2),
        Inches(0.48),
        title,
        size=28,
        bold=True,
        color=PURPLE_INK,
    )
    add_rect(slide, Inches(0.55), Inches(0.96), Inches(1.55), Inches(0.055), YELLOW)
    add_rect(slide, Inches(2.16), Inches(0.96), Inches(0.42), Inches(0.055), ORANGE)
    add_text(
        slide,
        Inches(11.55),
        Inches(0.42),
        Inches(1.35),
        Inches(0.32),
        f"{page:02d}  /  {TOTAL:02d}",
        size=12,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, 0, Inches(7.22), SW, Inches(0.28), PURPLE_PALE)
    add_text(
        slide,
        Inches(0.55),
        Inches(7.24),
        Inches(10.4),
        Inches(0.24),
        FOOTER,
        size=10,
        color=MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        Inches(11.35),
        Inches(7.24),
        Inches(1.55),
        Inches(0.24),
        "2026.09",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def lead_box(slide, text: str) -> None:
    add_round(slide, Inches(0.55), Inches(1.16), Inches(12.25), Inches(0.86), PURPLE_SOFT, adj=0.08)
    add_text(
        slide,
        Inches(0.78),
        Inches(1.22),
        Inches(11.85),
        Inches(0.74),
        text,
        size=14,
        color=INK,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def caption(slide, text: str, y=Inches(6.88)) -> None:
    add_text(slide, Inches(0.55), y, Inches(12.2), Inches(0.26), text, size=11, color=MUTED)


def build_cover(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, PURPLE_INK)
    add_rect(s, 0, 0, Inches(0.16), SH, YELLOW)
    add_rect(s, Inches(0.16), 0, Inches(0.08), SH, ORANGE)

    # 右上几何点缀（色块，非配图）
    add_oval(s, Inches(11.55), Inches(-1.15), Inches(3.1), Inches(3.1), PURPLE)
    add_oval(s, Inches(12.35), Inches(5.85), Inches(2.2), Inches(2.2), ORANGE)

    add_text(
        s,
        Inches(0.85),
        Inches(1.55),
        Inches(11.2),
        Inches(0.32),
        "分行领导专题汇报",
        size=14,
        bold=True,
        color=YELLOW,
    )
    add_text(
        s,
        Inches(0.85),
        Inches(2.05),
        Inches(11.5),
        Inches(1.55),
        "客户经理新员工培养阶段\n工作汇报",
        size=40,
        bold=True,
        color=WHITE,
    )
    add_rect(s, Inches(0.85), Inches(3.78), Inches(2.15), Inches(0.07), YELLOW)
    add_rect(s, Inches(3.08), Inches(3.78), Inches(0.55), Inches(0.07), ORANGE)

    add_text(
        s,
        Inches(0.85),
        Inches(4.15),
        Inches(11.0),
        Inches(0.36),
        "识别帮带短板  ·  优化激励规则  ·  提升培养有效性",
        size=16,
        color=RGBColor(0xE4, 0xDC, 0xF0),
    )

    # 底部信息条
    add_rect(s, 0, Inches(6.05), SW, Inches(1.45), PURPLE)
    add_rect(s, Inches(0.85), Inches(6.35), Inches(0.08), Inches(0.85), YELLOW)
    add_text(s, Inches(1.15), Inches(6.32), Inches(2.4), Inches(0.28), "汇报部门", size=12, color=YELLOW)
    add_text(s, Inches(1.15), Inches(6.60), Inches(3.6), Inches(0.42), "人力资源部", size=22, bold=True, color=WHITE)
    add_rect(s, Inches(5.35), Inches(6.45), Inches(0.025), Inches(0.65), RGBColor(0x8A, 0x6C, 0xC4))
    add_text(s, Inches(5.75), Inches(6.32), Inches(2.4), Inches(0.28), "汇报日期", size=12, color=YELLOW)
    add_text(s, Inches(5.75), Inches(6.60), Inches(3.6), Inches(0.42), "2026.09", size=22, bold=True, color=WHITE)


def build_roles(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, WHITE)
    content_chrome(s, "六大帮带角色调研得分情况", 2)
    lead_box(
        s,
        "本次围绕 6 项帮带维度开展全员调研，通过得分排序，精准识别当前帮带优势项与薄弱项，为后续优化培养重点提供数据支撑。",
    )

    # 横向条形图（形状绘制，便于控制紫黄橙层次）
    top = Inches(2.18)
    row_h = Inches(0.72)
    label_x, label_w = Inches(0.50), Inches(2.40)
    track_x, track_w = Inches(3.02), Inches(7.35)
    score_x = Inches(10.50)
    tag_x = Inches(11.72)

    for i, (name, score, tag) in enumerate(ROLE_SCORES):
        y = top + i * row_h
        add_text(
            s,
            label_x,
            y,
            label_w,
            Inches(0.58),
            name,
            size=16,
            bold=True,
            color=INK,
            anchor=MSO_ANCHOR.MIDDLE,
            align=PP_ALIGN.RIGHT,
        )
        add_round(s, track_x, y + Inches(0.16), track_w, Inches(0.30), TRACK, adj=0.5)
        bar_w = int(track_w * (score / SCORE_MAX))
        add_round(s, track_x, y + Inches(0.16), bar_w, Inches(0.30), BAR_COLORS[i], adj=0.5)
        add_text(
            s,
            score_x,
            y,
            Inches(1.10),
            Inches(0.58),
            f"{score:.2f}",
            size=20,
            bold=True,
            color=PURPLE_INK if tag == "优势" else (ORANGE_DEEP if tag == "薄弱" else INK),
            anchor=MSO_ANCHOR.MIDDLE,
        )
        if tag:
            bg = YELLOW_SOFT if tag == "优势" else ORANGE_SOFT
            fg = PURPLE if tag == "优势" else ORANGE_DEEP
            add_round(s, tag_x, y + Inches(0.13), Inches(0.95), Inches(0.32), bg, adj=0.4)
            add_text(
                s,
                tag_x,
                y + Inches(0.13),
                Inches(0.95),
                Inches(0.32),
                tag,
                size=12,
                bold=True,
                color=fg,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )

    caption(s, "图表说明：横向条形图，6 个维度按得分从高至低排序展示（5 分制，全员调研均值）。")


def _pie_png(path: Path) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager

    font_path = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
    font_manager.fontManager.addfont(font_path)
    plt.rcParams["font.family"] = font_manager.FontProperties(fname=font_path).get_name()
    plt.rcParams["axes.unicode_minus"] = False

    sizes = [item[1] for item in MOOD_SHARE]
    colors = [f"#{int(c[0]):02X}{int(c[1]):02X}{int(c[2]):02X}" for _, _, c in MOOD_SHARE]

    fig, ax = plt.subplots(figsize=(5.2, 5.2), dpi=200)
    fig.patch.set_alpha(0)
    ax.set_facecolor("none")
    ax.pie(
        sizes,
        colors=colors,
        startangle=90,
        counterclock=False,
        wedgeprops={"width": 1.0, "edgecolor": "white", "linewidth": 2.6},
    )
    ax.set_aspect("equal")
    fig.savefig(path, dpi=200, transparent=True, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)


def build_mood(prs, pie_path: Path) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, WHITE)
    content_chrome(s, "新人工作状态与情绪分布分析", 3)
    lead_box(
        s,
        "结合调研数据梳理新人真实工作状态，整体成长氛围向好，但存在部分低效、迷茫、压力内耗问题，需针对性优化帮带节奏与指导方式。",
    )

    s.shapes.add_picture(str(pie_path), Inches(0.62), Inches(2.20), Inches(4.35), Inches(4.35))

    add_text(
        s,
        Inches(5.28),
        Inches(2.18),
        Inches(7.4),
        Inches(0.30),
        "数据占比排序",
        size=13,
        bold=True,
        color=PURPLE,
    )
    row_top = Inches(2.52)
    for i, (name, pct, color) in enumerate(MOOD_SHARE):
        y = row_top + i * Inches(0.74)
        add_round(s, Inches(5.28), y, Inches(7.50), Inches(0.66), PURPLE_PALE, adj=0.10)
        add_round(s, Inches(5.48), y + Inches(0.19), Inches(0.28), Inches(0.28), color, adj=0.5)
        add_text(
            s,
            Inches(5.95),
            y,
            Inches(4.5),
            Inches(0.66),
            f"{i + 1}  {name}",
            size=17,
            bold=True,
            color=INK,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            s,
            Inches(10.55),
            y,
            Inches(2.05),
            Inches(0.66),
            f"{pct}%",
            size=24,
            bold=True,
            color=PURPLE_INK,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
        )

    caption(s, "图表说明：饼状图，展示五类状态占比分布。")


def rule_card(slide, x, y, w, h, num: str, accent) -> None:
    """正式空位卡片：仅保留序号与填报横线，不撰写规则正文。"""
    add_round(slide, x, y, w, h, PURPLE_PALE, adj=0.05)
    add_rect(slide, x, y, Inches(0.10), h, accent)
    add_text(
        slide,
        x + Inches(0.36),
        y + Inches(0.28),
        Inches(2.2),
        Inches(0.50),
        f"{num}.",
        size=28,
        bold=True,
        color=accent,
    )
    line_y = y + Inches(1.15)
    for k in range(4):
        add_rect(
            slide,
            x + Inches(0.36),
            line_y + Inches(0.52) * k,
            w - Inches(0.78),
            Inches(0.018),
            PURPLE_LINE,
        )


def build_qualify_rules(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, WHITE)
    content_chrome(s, "优化成功帮带合格率增量激励规则", 4)
    lead_box(
        s,
        "为提升帮带有效性、杜绝形式化帮带，对新人达标合格率增量实行差异化折算机制：",
    )
    card_y, card_h = Inches(2.28), Inches(4.55)
    gap = Inches(0.28)
    card_w = Inches(3.88)
    x0 = Inches(0.55)
    accents = (PURPLE, YELLOW_DEEP, ORANGE)
    for i, accent in enumerate(accents):
        rule_card(s, x0 + i * (card_w + gap), card_y, card_w, card_h, str(i + 1), accent)


def build_team_rules(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SW, SH, WHITE)
    content_chrome(s, "团队协同帮带激励费用优化方案", 5)
    lead_box(
        s,
        "推行团队互助帮带模式，打破单人帮带局限，明确团队帮带激励核算规则：",
    )
    card_y, card_h = Inches(2.28), Inches(4.55)
    gap = Inches(0.32)
    card_w = Inches(5.96)
    x0 = Inches(0.55)
    accents = (PURPLE, ORANGE)
    for i, accent in enumerate(accents):
        rule_card(s, x0 + i * (card_w + gap), card_y, card_w, card_h, str(i + 1), accent)


def build(output_path: Path = OUT) -> Path:
    assert len(ROLE_SCORES) == 6
    assert ROLE_SCORES == sorted(ROLE_SCORES, key=lambda item: -item[1])
    assert sum(item[1] for item in MOOD_SHARE) == 100

    global SW, SH
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SW, SH = prs.slide_width, prs.slide_height

    build_cover(prs)
    build_roles(prs)
    with TemporaryDirectory() as tmp:
        pie_path = Path(tmp) / "mood_pie.png"
        _pie_png(pie_path)
        build_mood(prs, pie_path)
        build_qualify_rules(prs)
        build_team_rules(prs)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))
    return output_path


if __name__ == "__main__":
    path = build()
    print(f"已生成：{path}")
