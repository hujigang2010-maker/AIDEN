#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将工作经历原文整理为 Excel、Word 与预览网页。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output"

# 按用户提供的原文录入，不改写单位名称、职务和主要工作内容。
RECORDS: list[dict[str, str]] = [
    {
        "no": "1",
        "start": "2015-09",
        "end": "2015-10",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海睿泽股权投资管理有限公司",
        "dept_title": "工程部/资料员、施工员",
        "work": "施工资料及现场管理",
    },
    {
        "no": "2",
        "start": "2015-12",
        "end": "2017-02",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "新城控股集团股份有限公司上海第一分公司",
        "dept_title": "工程部/高级经理",
        "work": "工程管理级施工协调",
    },
    {
        "no": "3",
        "start": "2017-03",
        "end": "2017-12",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海万科长宁置业有限公司",
        "dept_title": "工程部/总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "4",
        "start": "2018-01",
        "end": "2019-03",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "万科（上海）实业发展有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "5",
        "start": "2019-04",
        "end": "2020-02",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海万居企业管理有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "6",
        "start": "2020-03",
        "end": "2020-05",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "金科（上海）建筑设计有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "7",
        "start": "2020-06",
        "end": "2020-10",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海金科文化旅游发展集团有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "8",
        "start": "2020-11",
        "end": "2020-12",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海欣驭文化传播有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "9",
        "start": "2021-01",
        "end": "2021-02",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海弘阳汇商业管理有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "10",
        "start": "2021-03",
        "end": "2021-08",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海融之旅文化旅游发展有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "11",
        "start": "2021-09",
        "end": "2022-07",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海社通实业有限公司",
        "dept_title": "工程部总监",
        "work": "工程施工技术管理",
    },
    {
        "no": "12",
        "start": "2022-08",
        "end": "2023-04",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海苏科企业管理有限公司",
        "dept_title": "工程部副总经理",
        "work": "钢结构施工技术管理",
    },
    {
        "no": "13",
        "start": "2023-05",
        "end": "2023-08",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海容钲建筑科技有限公司",
        "dept_title": "工程部副总经理",
        "work": "工程施工技术管理",
    },
    {
        "no": "14",
        "start": "2023-10",
        "end": "2024-07",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海徐汇绿地商业管理有限公司",
        "dept_title": "工程部副总经理",
        "work": "工程施工技术管理",
    },
    {
        "no": "15",
        "start": "2024-08",
        "end": "2025-02",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海惟嘉建筑工程有限公司",
        "dept_title": "工程部副总经理",
        "work": "工程施工技术管理",
    },
    {
        "no": "16",
        "start": "2025-03",
        "end": "2025-12",
        "country": "中华人民共和国",
        "region": "江苏省南京市",
        "org": "江苏省建筑装饰设计研究院有限公司",
        "dept_title": "工程部副总经理",
        "work": "工程技术及项目管理",
    },
    {
        "no": "17",
        "start": "2026-01",
        "end": "2026-04",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海宝龙商办企业管理有限公司",
        "dept_title": "工程部副总经理",
        "work": "工程技术及项目管理",
    },
    {
        "no": "18",
        "start": "2026-05",
        "end": "今",
        "country": "中华人民共和国",
        "region": "上海市",
        "org": "上海贝迪创建科技有限公司",
        "dept_title": "工程部/高级工程师",
        "work": "工程技术及项目管理",
    },
]

MAIN_HEADERS = ["序号", "起止时间", "国家地区", "国家省市", "工作单位", "部门或职务", "主要工作内容"]
SPLIT_HEADERS = [
    "序号",
    "起始年月",
    "截止年月",
    "国家地区",
    "国家省市",
    "工作单位",
    "部门",
    "职务",
    "主要工作内容",
    "部门或职务（原文）",
]


def period(record: dict[str, str]) -> str:
    return f"{record['start']} 至 {record['end']}"


def month_index(value: str) -> int | None:
    if value == "今":
        return None
    year, month = value.split("-")
    return int(year) * 12 + int(month)


def format_year_month(index: int) -> str:
    year, month = divmod(index - 1, 12)
    return f"{year:04d}-{month + 1:02d}"


def find_gaps(records: list[dict[str, str]] | None = None) -> list[str]:
    """相邻两段之间未覆盖的月份。起止均按所给年月原样比较。"""
    rows = records if records is not None else RECORDS
    notes: list[str] = []
    for previous, current in zip(rows, rows[1:]):
        end = month_index(previous["end"])
        start = month_index(current["start"])
        if end is None or start is None or start <= end + 1:
            continue
        missing = [format_year_month(index) for index in range(end + 1, start)]
        joined = "、".join(missing)
        notes.append(f"第{previous['no']}段与第{current['no']}段之间空缺 {joined}")
    return notes


def split_dept_title(text: str) -> tuple[str, str]:
    """把“部门或职务”拆成部门、职务。含斜杠时按斜杠拆；否则“工程部”开头的记入部门。"""
    if "/" in text:
        dept, title = text.split("/", 1)
        return dept.strip(), title.strip()
    if text.startswith("工程部"):
        return "工程部", text[len("工程部") :].strip()
    return "", text.strip()


def gap_note() -> str:
    gaps = find_gaps()
    if not gaps:
        return "各段起止时间首尾相接，原文未提供的月份未作补写。"
    detail = "；".join(gaps)
    return f"原文未覆盖的月份保持空白，未作补写：{detail}。"


def main_row(record: dict[str, str]) -> list[str]:
    return [
        record["no"],
        period(record),
        record["country"],
        record["region"],
        record["org"],
        record["dept_title"],
        record["work"],
    ]


def split_row(record: dict[str, str]) -> list[str]:
    dept, title = split_dept_title(record["dept_title"])
    return [
        record["no"],
        record["start"],
        record["end"],
        record["country"],
        record["region"],
        record["org"],
        dept,
        title,
        record["work"],
        record["dept_title"],
    ]


def build_workbook() -> Workbook:
    wb = Workbook()
    wb.properties.title = "工作经历一览表"
    wb.properties.subject = "工作经历"
    _build_main_sheet(wb.active)
    _build_split_sheet(wb.create_sheet("分列填写"))
    return wb


def _thin_border() -> Border:
    line = Side(style="thin", color="1F2933")
    return Border(left=line, right=line, top=line, bottom=line)


def _apply_page(ws, fit_height: int) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.55, bottom=0.5, header=0.25, footer=0.25)
    ws.oddHeader.left.text = "工作经历一览表"
    ws.oddHeader.left.font = "宋体"
    ws.oddHeader.left.size = 10
    ws.oddFooter.left.text = "按提供材料原文录入"
    ws.oddFooter.left.font = "宋体"
    ws.oddFooter.left.size = 9
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"
    ws.oddFooter.right.font = "宋体"
    ws.oddFooter.right.size = 9
    ws.sheet_properties.tabColor = "1F4E79"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.print_options.horizontalCentered = True


def _write_banner(ws, last_col: int, subtitle: str) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(1, 1, "工作经历一览表")
    title.font = Font(name="宋体", size=18, bold=True, color="1F4E79")
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    note = ws.cell(2, 1, subtitle)
    note.font = Font(name="宋体", size=10, color="334155")
    note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 36


def _write_table(ws, headers: list[str], rows: list[list[str]], header_row: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="宋体", size=11, bold=True, color="FFFFFF")
    alt_fill = PatternFill("solid", fgColor="F4F7FB")
    value_font = Font(name="宋体", size=10, color="1F2933")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    border = _thin_border()
    left_columns = {5, 6, 7} if len(headers) == 7 else {6, 8, 9, 10}

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[header_row].height = 24
    ws.freeze_panes = f"A{header_row + 1}"

    for offset, values in enumerate(rows):
        excel_row = header_row + 1 + offset
        fill = alt_fill if offset % 2 else None
        for col, value in enumerate(values, start=1):
            cell = ws.cell(excel_row, col, int(value) if col == 1 else value)
            cell.font = value_font
            cell.alignment = left if col in left_columns else center
            cell.border = border
            if fill is not None:
                cell.fill = fill
        ws.row_dimensions[excel_row].height = 28

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
    ws.print_title_rows = f"1:{header_row}"
    last_row = header_row + len(rows)
    ws.print_area = f"A1:{get_column_letter(len(headers))}{last_row}"


def _build_main_sheet(ws) -> None:
    ws.title = "工作经历"
    headers = MAIN_HEADERS
    widths = [8, 22, 18, 16, 42, 26, 24]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    subtitle = (
        f"共 {len(RECORDS)} 段。国家地区均为中华人民共和国；截止时间“今”表示至今。"
        f"{gap_note()}"
    )
    _write_banner(ws, len(headers), subtitle)
    _write_table(ws, headers, [main_row(record) for record in RECORDS], 3)
    _apply_page(ws, 1)


def _build_split_sheet(ws) -> None:
    headers = SPLIT_HEADERS
    widths = [8, 14, 12, 18, 16, 42, 14, 18, 24, 26]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    subtitle = (
        "本表便于分栏填写。原文含“/”时按斜杠拆成部门与职务；"
        "原文写作“工程部总监”“工程部副总经理”时，部门记为“工程部”，其余记为职务。"
        "第一张表保留原文，不做拆分。"
    )
    _write_banner(ws, len(headers), subtitle)
    _write_table(ws, headers, [split_row(record) for record in RECORDS], 3)
    _apply_page(ws, 1)
    ws.sheet_properties.tabColor = "0F6E56"


def _set_run_font(run, name="宋体", size=10.5, bold=False, color=None) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), hex_color)
    shade.set(qn("w:val"), "clear")
    tc_pr.append(shade)


def _valign(cell, value="center") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    align = OxmlElement("w:vAlign")
    align.set(qn("w:val"), value)
    tc_pr.append(align)


def _write_doc_cell(cell, text: str, *, bold=False, size=9, align="center", fill=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold, color=RGBColor(0xFF, 0xFF, 0xFF) if fill == "1F4E79" else None)
    if fill:
        _shade(cell, fill)
    _valign(cell)


def _set_col_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    total = int(sum(widths_cm) * 567)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tbl_pr.addnext(grid)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 567)))
        grid.append(col)

    for row in table.rows:
        tr = row._tr
        tr_pr = tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("工作经历一览表")
    _set_run_font(run, size=9, color=RGBColor(0x64, 0x74, 0x8B))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("按提供材料原文录入")
    _set_run_font(run, size=9, color=RGBColor(0x64, 0x74, 0x8B))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("工作经历一览表")
    _set_run_font(run, size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    intro.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = intro.add_run(
        f"共 {len(RECORDS)} 段。国家地区均为中华人民共和国；截止时间“今”表示至今。{gap_note()}"
    )
    _set_run_font(run, size=9, color=RGBColor(0x33, 0x41, 0x55))

    widths = [1.3, 3.5, 2.8, 2.7, 7.4, 4.4, 5.2]
    table = doc.add_table(rows=1 + len(RECORDS), cols=len(MAIN_HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_col_widths(table, widths)

    header_row = table.rows[0]._tr
    header_pr = header_row.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    header_pr.append(repeat)

    for col, text in enumerate(MAIN_HEADERS):
        _write_doc_cell(table.cell(0, col), text, bold=True, size=9, fill="1F4E79")
    left_cols = {4, 5, 6}
    for row_index, record in enumerate(RECORDS, start=1):
        values = main_row(record)
        fill = "F4F7FB" if row_index % 2 == 0 else None
        for col, value in enumerate(values):
            _write_doc_cell(
                table.cell(row_index, col),
                value,
                size=8.5,
                align="left" if col in left_cols else "center",
                fill=fill,
            )

    split_title = doc.add_paragraph()
    split_title.paragraph_format.space_before = Pt(12)
    split_title.paragraph_format.space_after = Pt(2)
    run = split_title.add_run("部门与职务分列")
    _set_run_font(run, size=12, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    split_note = doc.add_paragraph()
    split_note.paragraph_format.space_after = Pt(4)
    run = split_note.add_run(
        "下表供分栏填写使用。原文含“/”时按斜杠拆分；"
        "“工程部总监”“工程部副总经理”的部门记为“工程部”，其余记为职务。"
    )
    _set_run_font(run, size=9, color=RGBColor(0x33, 0x41, 0x55))

    split_headers = ["序号", "起始年月", "截止年月", "国家省市", "工作单位", "部门", "职务", "主要工作内容"]
    split_widths = [1.2, 2.2, 1.8, 2.6, 7.6, 2.2, 3.2, 6.5]
    split = doc.add_table(rows=1 + len(RECORDS), cols=len(split_headers))
    split.style = "Table Grid"
    split.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_col_widths(split, split_widths)
    split_header = split.rows[0]._tr
    split_header_pr = split_header.get_or_add_trPr()
    split_header_pr.append(OxmlElement("w:tblHeader"))
    for col, text in enumerate(split_headers):
        _write_doc_cell(split.cell(0, col), text, bold=True, size=8, fill="1F4E79")
    for row_index, record in enumerate(RECORDS, start=1):
        dept, title_text = split_dept_title(record["dept_title"])
        values = [
            record["no"],
            record["start"],
            record["end"],
            record["region"],
            record["org"],
            dept,
            title_text,
            record["work"],
        ]
        fill = "F4F7FB" if row_index % 2 == 0 else None
        for col, value in enumerate(values):
            _write_doc_cell(
                split.cell(row_index, col),
                value,
                size=8,
                align="left" if col in {4, 7} else "center",
                fill=fill,
            )
    return doc


def build_preview_html() -> str:
    def cell(text: str, kind: str) -> str:
        return f'<td class="{kind}">{text}</td>'

    body_rows = []
    for record in RECORDS:
        values = main_row(record)
        kinds = ["c", "c", "c", "c", "l", "l", "l"]
        tds = "".join(cell(value, kind) for value, kind in zip(values, kinds))
        body_rows.append(f"<tr>{tds}</tr>")
    header = "".join(f"<th>{name}</th>" for name in MAIN_HEADERS)
    note = (
        f"共 {len(RECORDS)} 段。国家地区均为中华人民共和国；截止时间“今”表示至今。{gap_note()}"
    )
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8" />
<title>工作经历一览表</title>
<style>
  body {{
    margin: 24px;
    background: #fff;
    color: #1f2933;
    font-family: "WenQuanYi Micro Hei", "Noto Sans CJK SC", "Microsoft YaHei", sans-serif;
  }}
  h1 {{
    margin: 0 0 8px;
    text-align: center;
    color: #1f4e79;
    font-size: 28px;
    font-weight: 700;
  }}
  p {{
    margin: 0 0 14px;
    font-size: 14px;
    line-height: 1.6;
    color: #334155;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    table-layout: fixed;
  }}
  th, td {{
    border: 1px solid #1f2933;
    padding: 8px 6px;
    font-size: 13px;
    line-height: 1.45;
    vertical-align: middle;
  }}
  th {{
    background: #1f4e79;
    color: #fff;
    font-weight: 700;
    text-align: center;
  }}
  td.c {{ text-align: center; }}
  td.l {{ text-align: left; }}
  tbody tr:nth-child(even) td {{ background: #f4f7fb; }}
  col.c1 {{ width: 6%; }}
  col.c2 {{ width: 14%; }}
  col.c3 {{ width: 12%; }}
  col.c4 {{ width: 11%; }}
  col.c5 {{ width: 26%; }}
  col.c6 {{ width: 16%; }}
  col.c7 {{ width: 15%; }}
</style>
</head>
<body>
  <h1>工作经历一览表</h1>
  <p>{note}</p>
  <table>
    <colgroup>
      <col class="c1" /><col class="c2" /><col class="c3" /><col class="c4" />
      <col class="c5" /><col class="c6" /><col class="c7" />
    </colgroup>
    <thead><tr>{header}</tr></thead>
    <tbody>
      {"".join(body_rows)}
    </tbody>
  </table>
</body>
</html>
"""


def write_outputs(directory: Path | None = None) -> dict[str, Path]:
    target = directory or OUT_DIR
    target.mkdir(parents=True, exist_ok=True)
    xlsx = target / "工作经历一览表.xlsx"
    docx = target / "工作经历一览表.docx"
    html = target / "工作经历一览表.html"
    build_workbook().save(xlsx)
    build_document().save(docx)
    html.write_text(build_preview_html(), encoding="utf-8")
    return {"xlsx": xlsx, "docx": docx, "html": html}


def main() -> None:
    paths = write_outputs()
    for path in paths.values():
        print(path)


if __name__ == "__main__":
    main()
