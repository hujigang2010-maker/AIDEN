#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成王德峰老师公开讲座的对外策划案，以及海报用议程。"""

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generate_wangdefeng_proposal as g  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "output"
DOC_PATH = OUT / "王德峰老师公开讲座_对外活动方案.docx"
POSTER_PATH = OUT / "王德峰老师公开讲座_对外议程海报.png"
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"

AGENDA = [
    ("08:40–09:10", "入场签到"),
    ("09:10–09:20", "开场"),
    ("09:20–10:50", "王德峰老师公开讲座"),
    ("10:50–11:20", "现场交流"),
    ("11:20–11:40", "合影留念"),
]


def build_doc():
    doc = g.new_doc()
    g.setup_section(doc, "对外活动方案  ·  王德峰老师公开讲座  ·  2026年10月31日")

    g.add_para(
        doc,
        "对外活动方案",
        size=14,
        bold=True,
        color=g.GOLD,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=4,
        space_after=4,
    )
    g.add_para(
        doc,
        "王德峰老师公开讲座",
        size=22,
        bold=True,
        color=g.NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    g.add_para(
        doc,
        "2026年10月31日（周六）上午  ·  上海",
        size=12,
        color=g.NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
    )
    g.add_quote_box(
        doc,
        "给听众的一句话",
        "一场两个半至三小时的线下公开讲座。王德峰老师主讲，面向愿意把哲学、中国思想与当下生活放在一起听的人。",
    )

    g.add_h1(doc, "一、活动信息")
    g.add_table(
        doc,
        ["项目", "对外口径"],
        [
            ["活动名称", "王德峰老师公开讲座"],
            ["时间", "2026年10月31日（周六）上午，约 2.5–3 小时"],
            ["地点", "上海。具体场地以报名页为准"],
            ["形式", "线下公开讲座，一个主题，持票入场"],
            ["联合主办", "上海市杨浦区科技企业联合会、首乾书院"],
            ["标准票", "399 元。票档为 399 / 499 / 599 / 699 / 999 元"],
        ],
        [3.8, 12.6],
    )
    g.add_para(
        doc,
        "场地进场钟点若有调整，只改时间，不改环节顺序。海报和报名页使用同一套议程。",
        size=10.5,
        color=g.MUTED,
        space_before=6,
    )

    g.add_h1(doc, "二、嘉宾")
    g.add_para(
        doc,
        "王德峰，复旦大学哲学学院教授。长期从事哲学教学，也常以公开讲座的方式，向校园以外的听众讲述哲学与人生、中国思想与现代生活。本场是一场面向公众的讲座，不是系列课程的课堂。",
        size=11,
        first_line=True,
    )

    g.add_h1(doc, "三、海报议程")
    g.add_para(
        doc,
        "以下议程直接用于对外海报和报名页。",
        size=11,
        space_after=6,
    )
    g.add_table(
        doc,
        ["时间", "环节", "说明"],
        [
            ["08:40–09:10", "入场签到", "凭票入场，按票档分区就座"],
            ["09:10–09:20", "开场", "介绍到场嘉宾与本场安排"],
            ["09:20–10:50", "公开讲座", "王德峰老师主讲，约 90 分钟"],
            ["10:50–11:20", "现场交流", "围绕讲座内容互动。前排票含提问资格"],
            ["11:20–11:40", "合影留念", "有序合影。前排票含签名或合影"],
        ],
        [3.6, 3.4, 9.4],
    )

    g.add_h1(doc, "四、票务与现场权益")
    g.add_table(
        doc,
        ["票档", "价格", "听众获得"],
        [
            ["标准票", "399 元", "讲座正座，本场对外主价格"],
            ["进阶票", "499 元", "讲座正座"],
            ["优选票", "599 元", "讲座正座"],
            ["前区票", "699 元", "前区座位"],
            ["前排票", "999 元", "前排座位，并含提问、签名或合影中的一项"],
        ],
        [3.2, 3.0, 10.2],
        emphasize_col=1,
    )
    g.add_para(
        doc,
        "报名、座位和发票以正式报名页为准。海报上的价格与报名页保持一致。",
        size=10.5,
        color=g.MUTED,
        space_before=6,
    )

    g.add_h1(doc, "五、听众")
    g.add_bullets(
        doc,
        [
            "关心哲学、中国思想与现实生活的普通听众。",
            "企业经营者、管理者，以及希望在业务之外保留一段思想时间的人。",
            "在校师生和文化爱好者。",
        ],
    )
    g.add_para(
        doc,
        "会场按公开讲座安排：听得清、进得去、结束时间明确。不设课程作业，也不在现场改成其他课程的招生说明会。",
        size=11,
        first_line=True,
    )

    g.add_h1(doc, "六、联合主办")
    g.add_table(
        doc,
        ["单位", "在本场中的角色"],
        [
            ["上海市杨浦区科技企业联合会", "联合主办。负责活动组织、报名、现场与城市听众的召集"],
            ["首乾书院", "联合主办。负责王德峰老师出席、讲座内容，以及书院听众的通知"],
        ],
        [6.4, 10.0],
    )
    g.add_para(
        doc,
        "海报、报名页和现场背景同时出现双方名称。对外表述为联合主办，不使用「独家」或「承办方特邀」等容易让听众误解主办关系的说法。",
        size=11,
        first_line=True,
    )

    g.add_h1(doc, "七、海报上怎么写")
    g.add_para(
        doc,
        "海报信息从上到下只有这些，避免把报名规则写成小字说明。",
        size=11,
        space_after=4,
    )
    g.add_bullets(
        doc,
        [
            "眉题：联合主办　上海市杨浦区科技企业联合会 × 首乾书院",
            "主标题：王德峰老师公开讲座",
            "时间地点：2026年10月31日（周六）上午　上海",
            "中部：上面这五段议程",
            "价格：标准票 399 元",
            "落款：持票入场　具体场地与报名以报名页为准",
        ],
    )

    g.add_para(
        doc,
        "本方案只用于对外介绍这场讲座。报名页上线后，议程、场地和价格以报名页为准。",
        size=10,
        color=g.MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=16,
    )
    doc.save(DOC_PATH)
    return DOC_PATH


def _font(size):
    return ImageFont.truetype(FONT_PATH, size)


def _center(draw, text, y, font, fill, width):
    box = draw.textbbox((0, 0), text, font=font)
    x = (width - (box[2] - box[0])) / 2
    draw.text((x, y), text, font=font, fill=fill)
    return box[3] - box[1]


def build_poster():
    w, h = 1080, 1680
    navy = (26, 58, 92)
    gold = (166, 124, 46)
    cream = (248, 244, 236)
    ink = (36, 36, 36)
    muted = (90, 90, 90)
    white = (255, 252, 247)

    img = Image.new("RGB", (w, h), cream)
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, w, 520), fill=navy)
    draw.rectangle((0, 512, w, 520), fill=gold)

    y = 72
    _center(draw, "联合主办", y, _font(28), gold, w)
    y = 124
    _center(draw, "上海市杨浦区科技企业联合会  ×  首乾书院", y, _font(26), (232, 224, 208), w)
    y = 230
    _center(draw, "王德峰老师", y, _font(92), white, w)
    y = 350
    _center(draw, "公 开 讲 座", y, _font(48), gold, w)
    y = 440
    _center(draw, "2026年10月31日（周六）上午    上海", y, _font(30), (232, 224, 208), w)

    card_top = 580
    draw.rounded_rectangle((70, card_top, w - 70, 1280), radius=18, fill=white, outline=(226, 216, 198), width=2)
    _center(draw, "议  程", card_top + 36, _font(28), gold, w)

    row_y = card_top + 100
    for time_text, title in AGENDA:
        draw.text((120, row_y), time_text, font=_font(32), fill=navy)
        draw.text((460, row_y), title, font=_font(34), fill=ink)
        row_y += 108

    y = 1340
    _center(draw, "标准票  399 元", y, _font(48), navy, w)
    y = 1410
    _center(draw, "票档  399  /  499  /  599  /  699  /  999 元", y, _font(26), muted, w)
    y = 1480
    _center(draw, "持票入场", y, _font(28), gold, w)
    y = 1540
    _center(draw, "具体场地与报名，以报名页为准", y, _font(24), muted, w)
    y = 1600
    draw.rectangle((420, y, 660, y + 4), fill=gold)

    img.save(POSTER_PATH, "PNG")
    return POSTER_PATH


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    print(f"已生成策划案：{build_doc()}")
    print(f"已生成海报：{build_poster()}")


if __name__ == "__main__":
    main()
