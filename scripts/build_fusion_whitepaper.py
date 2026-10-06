# -*- coding: utf-8 -*-
"""将白皮书 Markdown 转为 Word 与 PDF。

依赖：pandoc、LibreOffice（soffice）、python-docx。
用法：python3 scripts/build_fusion_whitepaper.py
"""

import subprocess
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
NAME = "十五五氢能与核聚变能产业链白皮书"
SRC = ROOT / "whitepaper" / f"{NAME}.md"
OUT = ROOT / "whitepaper" / "下载版本"
DOCX = OUT / f"{NAME}.docx"

CN_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


def set_font(style_or_run, size=None, bold=None, color=None, east=CN_FONT, latin="Arial"):
    font = style_or_run.font
    font.name = latin
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color:
        font.color.rgb = RGBColor.from_string(color)
    el = style_or_run.element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), latin)
    rfonts.set(qn("w:eastAsia"), east)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def add_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "8EA9C1")
        borders.append(b)
    tbl_pr.append(borders)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def polish(path):
    doc = Document(path)
    styles = {s.name.lower(): s for s in doc.styles}
    for name in ("normal", "body text", "first paragraph", "compact", "block text"):
        if name in styles:
            set_font(styles[name], size=10.5)
    sizes = {"title": 22, "subtitle": 14, "heading 1": 18, "heading 2": 15, "heading 3": 12.5}
    for name, size in sizes.items():
        if name in styles:
            set_font(styles[name], size=size, bold=True, color="0B3D6B")
    if "source code" in styles:
        set_font(styles["source code"], size=8.5, latin=CODE_FONT)

    for p in doc.paragraphs:
        for run in p.runs:
            if run.style is not None and run.style.name == "Verbatim Char":
                set_font(run, size=8.5, latin=CODE_FONT)

    for table in doc.tables:
        add_borders(table)
        for r, row in enumerate(table.rows):
            for cell in row.cells:
                if r == 0:
                    shade(cell, "DCE8F3")
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_font(run, size=9, bold=True if r == 0 else None)

    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(64)
        section.top_margin = section.bottom_margin = Pt(64)
    doc.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["pandoc", str(SRC), "-f", "gfm", "-t", "docx", "-o", str(DOCX)],
        check=True,
    )
    polish(DOCX)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(DOCX)],
        check=True,
        stdout=subprocess.DEVNULL,
    )
    (OUT / f"{NAME}.md").write_bytes(SRC.read_bytes())
    for f in sorted(OUT.iterdir()):
        print(f.name, f.stat().st_size)


if __name__ == "__main__":
    main()
