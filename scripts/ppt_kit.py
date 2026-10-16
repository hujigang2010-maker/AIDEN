# -*- coding: utf-8 -*-
"""PPT 版式工具。坐标单位为英寸。"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x0E, 0x24, 0x38)
NAVY_MID = RGBColor(0x1B, 0x3A, 0x56)
INK = RGBColor(0x1C, 0x27, 0x33)
GOLD = RGBColor(0xC4, 0xA1, 0x5A)
GOLD_DEEP = RGBColor(0x8C, 0x6E, 0x2F)
TEAL = RGBColor(0x1F, 0x6A, 0x7A)
ICE = RGBColor(0xF3, 0xF6, 0xF9)
MIST = RGBColor(0xE6, 0xED, 0xF3)
SAND = RGBColor(0xFB, 0xF6, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5C, 0x6B, 0x78)
LINE = RGBColor(0xD5, 0xDE, 0xE6)
LIGHT_GOLD = RGBColor(0xF8, 0xF1, 0xDE)

FONT = "Microsoft YaHei"
FOOTER = "潮起潮落 TIDELINE  ×  北外滩尚9·一滴水    内部研讨稿"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def rgb_of(hex_color):
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_font(run, size, color=INK, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    run.font.language_id = MSO_LANGUAGE_ID.SIMPLIFIED_CHINESE
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def add_rect(slide, x, y, w, h, fill, line=None, line_pt=0.75):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_pt)
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = "" if text is None else str(text)
    set_font(r, size, color, bold)
    return tb


def add_lines(slide, x, y, w, h, lines, size=14, color=INK, bold=False,
              spacing=6, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        p.line_spacing = 1.08
        if isinstance(line, str):
            r = p.add_run()
            r.text = line
            set_font(r, size, color, bold)
        else:
            for text, sz, col, bd in line:
                r = p.add_run()
                r.text = text
                set_font(r, sz, col, bd)
    return tb


def bullets(slide, x, y, w, h, items, size=14, spacing=7, color=INK):
    lines = []
    for it in items:
        if isinstance(it, tuple):
            lead, rest = it
            lines.append([
                ("—  ", size, GOLD_DEEP, True),
                (lead, size, NAVY, True),
                ("  " + rest, size, color, False),
            ])
        else:
            lines.append([
                ("—  ", size, GOLD_DEEP, True),
                (it, size, color, False),
            ])
    return add_lines(slide, x, y, w, h, lines, size=size, spacing=spacing)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_page_slot(slide, color=RGBColor(0xC5, 0xD0, 0xDA)):
    tb = add_text(
        slide, 11.35, 7.15, 1.55, 0.26, "00", 11, color, True,
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name="PAGE",
    )
    return tb


def footer(slide, color=MUTED):
    add_text(slide, 0.42, 7.15, 10.6, 0.26, FOOTER, 10, color, anchor=MSO_ANCHOR.MIDDLE)


def header(slide, kicker, title):
    if len(title) > 34:
        raise RuntimeError("标题过长：" + title)
    add_rect(slide, 0, 0, 13.333, 0.86, NAVY)
    add_rect(slide, 0, 0.86, 13.333, 0.045, GOLD)
    add_text(slide, 0.42, 0.05, 10.8, 0.24, kicker, 11, GOLD, True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.42, 0.30, 12.2, 0.50, title, 22, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    footer(slide)
    add_page_slot(slide)


def callout(slide, x, y, w, h, text, fill=SAND, size=12, color=INK, bold=False):
    add_rect(slide, x, y, w, h, fill)
    add_rect(slide, x, y, 0.06, h, GOLD)
    lines = text.split("\n") if isinstance(text, str) else text
    add_lines(
        slide, x + 0.18, y + 0.08, w - 0.30, h - 0.12, lines,
        size=size, color=color, bold=bold, spacing=4, anchor=MSO_ANCHOR.MIDDLE,
    )


def card(slide, x, y, w, h, title, body, kicker=None, fill=WHITE, bar=GOLD, title_size=15, body_size=12):
    add_rect(slide, x, y, w, h, fill, LINE, 0.75)
    add_rect(slide, x, y, w, 0.045, bar)
    cursor = y + 0.12
    if kicker:
        add_text(slide, x + 0.14, cursor, w - 0.28, 0.24, kicker, 11, GOLD_DEEP, True)
        cursor += 0.24
    add_text(slide, x + 0.14, cursor, w - 0.28, 0.36, title, title_size, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
    cursor += 0.38
    if isinstance(body, list) or (isinstance(body, str) and "\n" in body):
        lines = body if isinstance(body, list) else body.split("\n")
        add_lines(slide, x + 0.14, cursor, w - 0.28, y + h - cursor - 0.08, lines, body_size, INK, spacing=2)
    elif body:
        add_text(slide, x + 0.14, cursor, w - 0.28, y + h - cursor - 0.08, body, body_size, INK)


def place_grid(items, cols, top, card_h, gap_x=0.14, gap_y=0.12, left=0.42, total_w=12.50):
    width = (total_w - gap_x * (cols - 1)) / cols
    boxes = []
    for i, item in enumerate(items):
        r, c = divmod(i, cols)
        x = left + c * (width + gap_x)
        y = top + r * (card_h + gap_y)
        boxes.append((x, y, width, card_h, item))
    return boxes


def _set_cell_border(cell, color_hex="E1E7EE", width_emu="6350"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for old in tcPr.findall(qn(f"a:{edge}")):
            tcPr.remove(old)
        ln = etree.fromstring(
            f'<a:{edge} xmlns:a="{A_NS}" w="{width_emu}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{color_hex}"/></a:solidFill>'
            f'<a:prstDash val="solid"/></a:{edge}>'
        )
        tcPr.append(ln)


def _strip_table_style(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("bandCol", "0")
    for el in list(tblPr):
        if el.tag.endswith("tableStyleId"):
            tblPr.remove(el)


def write_cell(cell, text, size, color, bold, fill, align):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.07)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.text = "" if text is None else str(text)
    tf = cell.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            set_font(r, size, color, bold)
    _set_cell_border(cell)


def add_table(slide, x, y, col_ws, rows, font=12, aligns=None, row_h=0.34, header_h=None):
    """rows 含表头。col_ws 为英寸宽度列表。返回表格底部 y。"""
    ncol = len(col_ws)
    nrow = len(rows)
    if any(len(r) != ncol for r in rows):
        raise RuntimeError("表格列数不一致")
    header_h = header_h or row_h
    total_h = header_h + row_h * (nrow - 1)
    bottom = y + total_h
    if bottom > 7.02:
        raise RuntimeError(f"表格超出页面：bottom={bottom:.2f}，行数={nrow}")
    width = sum(col_ws)
    table_shape = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(width), Inches(total_h))
    table = table_shape.table
    _strip_table_style(table)
    table.rows[0].height = Inches(header_h)
    for i in range(1, nrow):
        table.rows[i].height = Inches(row_h)
    for i, w in enumerate(col_ws):
        table.columns[i].width = Inches(w)
    if aligns is None:
        aligns = [PP_ALIGN.LEFT] * ncol
    for ri, row in enumerate(rows):
        is_header = ri == 0
        is_total = (not is_header) and str(row[0]).startswith("合计")
        for ci, val in enumerate(row):
            if is_header:
                fill, color, bold = NAVY, WHITE, True
            elif is_total:
                fill, color, bold = LIGHT_GOLD, NAVY, True
            else:
                fill = ICE if ri % 2 == 0 else WHITE
                color, bold = INK, ci == 0
            write_cell(cell=table.cell(ri, ci), text=val, size=font, color=color,
                       bold=bold, fill=fill, align=aligns[ci])
    return bottom


def note(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def stamp_pages(prs):
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.name == "PAGE" and shape.has_text_frame:
                shape.text_frame.paragraphs[0].runs[0].text = f"{i:02d}  /  {total:02d}"
    return total
