#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""王德峰老师公开讲座赞助方案，对外商务版。不含分成与内部谈判。"""

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generate_wangdefeng_proposal as g  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "output"
PATH = OUT / "王德峰老师公开讲座_赞助方案_对外商务版.docx"

TIERS = [
    ["首席支持", "68,000 元", "1 家", "首席支持伙伴"],
    ["战略支持", "48,000 元", "1 家", "战略支持伙伴"],
    ["合作支持", "38,000 元", "2 家", "合作支持伙伴"],
    ["品牌支持", "28,000 元", "2 家", "品牌支持伙伴"],
    ["鸣谢支持", "18,000 元", "3 家", "鸣谢伙伴"],
    ["友好支持", "9,800 元", "4 家", "友好支持伙伴"],
]


def build():
    doc = g.new_doc()
    g.setup_section(doc, "赞助方案  ·  对外商务版  ·  2026年10月17日")

    g.add_para(
        doc,
        "赞助方案",
        size=14,
        bold=True,
        color=g.GOLD,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=2,
        space_after=2,
    )
    g.add_para(
        doc,
        "王德峰老师公开讲座",
        size=22,
        bold=True,
        color=g.NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    g.add_para(
        doc,
        "2026年10月31日（周六）9:00–12:00　　场地另行通知　　约 150 座",
        size=12,
        color=g.NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )
    g.add_quote_box(
        doc,
        "给赞助企业",
        "一场面向企业负责人和公众听众的上午讲座。品牌出现在报名页、现场和活动照片中。讲座内容不冠名，不出售授课录像。一家企业选择一档。",
    )

    g.add_h1(doc, "一、活动")
    g.add_table(
        doc,
        ["项目", "说明"],
        [
            ["主讲", "王德峰，复旦大学哲学学院教授"],
            ["形式", "线下公开讲座，约 90 分钟主讲，随后现场交流"],
            ["主办", "上海市杨浦区科技企业联合会、首乾书院"],
            ["时间", "2026年10月31日（周六）9:00–12:00"],
            ["场地", "具体地点待定，确定后另行通知。若由机构提供场地，场地支持单独鸣谢，不占本方案现金赞助席位"],
            ["规模", "按约 150 座准备。对外售票 128 张，其余为工作人员、嘉宾和已安排席位"],
            ["听众", "企业负责人、管理者，以及关心哲学与中国思想的公众"],
        ],
        [3.4, 13.0],
    )
    g.add_para(
        doc,
        "场地确定后如座位有变化，对外出售的票数随之调整。已经确认的赞助企业，门票张数仍然保留。企业随讲座出现在同一批听众面前。现场不设展位，也不做产品推介。",
        size=11,
        first_line=True,
        space_before=6,
    )

    agenda_title = g.add_h1(doc, "二、上午议程")
    agenda_title.paragraph_format.page_break_before = True
    agenda_title.paragraph_format.keep_with_next = True
    g.add_table(
        doc,
        ["时间", "环节", "赞助企业在场的方式"],
        [
            ["08:30–09:00", "入场签到", "看到背景板和报名页标识。首席支持可放一份资料"],
            ["09:00–09:10", "开场", "主持人鸣谢。首席可致辞 60 秒，战略单独念到名称"],
            ["09:10–10:40", "公开讲座", "王德峰老师主讲。此段不插播企业内容"],
            ["10:40–11:20", "现场交流", "按票档就座，不单独设置提问权"],
            ["11:20–12:00", "签书与合影", "前排票签书或签名。企业可用现场照片，不用授课画面"],
        ],
        [3.3, 3.2, 9.9],
    )

    g.add_h1(doc, "三、六档赞助")
    g.add_para(
        doc,
        "六档里选一档。金额为人民币，含本页所列权益，不含企业自行制作的礼品和额外广告。六档不必同时招满，背景板按实际确认的企业排列。",
        size=11,
        first_line=True,
    )
    g.add_table(
        doc,
        ["档位", "金额", "家数", "对外称谓"],
        TIERS,
        [3.6, 3.6, 3.2, 6.0],
        emphasize_col=1,
    )
    fit = g.add_para(doc, "各档适合的企业", size=12, bold=True, color=g.NAVY, space_before=2, space_after=4)
    fit.paragraph_format.page_break_before = True
    fit.paragraph_format.keep_with_next = True
    g.add_table(
        doc,
        ["档位", "适合这样的企业"],
        [
            ["首席支持", "要主标识，并在开场致辞 1 分钟"],
            ["战略支持", "位置仅次于首席，开场被单独介绍"],
            ["合作支持", "要前排席和标识，可与非竞品同时出现"],
            ["品牌支持", "背景板要有标识，预算低于合作支持"],
            ["鸣谢支持", "名称出现在鸣谢区"],
            ["友好支持", "列入支持名单，并带两张后排票"],
        ],
        [3.6, 12.8],
    )

    h = g.add_h1(doc, "四、权益对照")
    h.paragraph_format.keep_with_next = True
    g.add_table(
        doc,
        ["权益", "首席", "战略", "合作", "品牌", "鸣谢", "友好"],
        [
            ["金额", "6.8万", "4.8万", "3.8万", "2.8万", "1.8万", "0.98万"],
            ["名额", "1 家", "1 家", "2 家", "2 家", "3 家", "4 家"],
            ["报名页", "主位置", "次主位置", "第二序列", "品牌列", "鸣谢列", "名单"],
            ["背景板", "主标识", "次主标识", "有", "品牌列", "鸣谢区", "不单列"],
            ["开场", "口播+60秒", "单独口播", "口播名称", "集体鸣谢", "集体鸣谢", "列入名单"],
            ["门票", "前排 4 张", "前排 2 张", "前排 2 张", "中排 2 张", "中排 2 张", "后排 2 张"],
            ["签到资料", "1 份", "—", "—", "—", "—", "—"],
            ["现场照片", "可用", "可用", "可用", "可用", "可用", "可用"],
            ["授课录像", "不含", "不含", "不含", "不含", "不含", "不含"],
            ["同业", "独家", "错开首席", "互避竞品", "不限", "不限", "不限"],
        ],
        [2.2, 2.36, 2.36, 2.36, 2.36, 2.36, 2.4],
    )
    g.add_para(
        doc,
        "门票从本场对应票档中划出，不另外加座。前排票当前售价 999 元，中排票 699 元，后排票 499 元。若六档全部确认，划出前排票 10 张、中排票 10 张、后排票 8 张。未确认的档位不预留空座。赞助企业未使用的门票不折现、不转让给本场以外的活动。",
        size=10.5,
        color=g.MUTED,
        space_before=4,
    )
    g.add_para(
        doc,
        "首席支持同一行业只此一家。战略支持不与首席支持同行业。两家合作支持之间不接受直接竞品。品牌支持、鸣谢支持和友好支持不设同业限制。",
        size=10.5,
        color=g.MUTED,
        space_before=2,
    )

    h = g.add_h1(doc, "五、现场怎么出现")
    h.paragraph_format.keep_with_next = True
    g.add_table(
        doc,
        ["位置", "观众看到什么"],
        [
            ["报名页", "讲座信息下方，按首席、战略、合作、品牌、鸣谢、友好的顺序"],
            ["海报", "底部赞助列。讲座名称保持「王德峰老师公开讲座」"],
            ["签到台", "仅首席支持可放一份资料"],
            ["背景板", "主办双方名称，其下为赞助标识。场地确定后，场地支持单独列明"],
            ["开场", "鸣谢不超过 3 分钟。首席致辞不超过 60 秒，战略支持单独念到名称"],
            ["讲座过程", "不插播广告，不发放课程资料"],
            ["结束后", "企业可获得现场照片。照片不含可单独传播的授课片段"],
        ],
        [3.6, 12.8],
    )

    h = g.add_h1(doc, "六、不包含的内容")
    h.paragraph_format.keep_with_next = True
    g.add_bullets(
        doc,
        [
            "不以赞助企业名称替换讲座名称，也不把讲座改成产品发布会。",
            "不提供王德峰老师授课录像、回放和课程剪辑。",
            "不安排老师为赞助企业单独代言、录祝贺视频或出席企业内场。若另有需要，须老师书面同意后单签。",
            "不在座位区和签到动线销售其他课程。",
            "赞助款不兑换成可退的门票，企业缺席不退还赞助款。",
            "不把赞助企业写成协办单位或联合主办。联合主办只有主办双方。",
        ],
    )

    h = g.add_h1(doc, "七、时间")
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.page_break_before = True
    g.add_table(
        doc,
        ["日期", "事项"],
        [
            ["场地确定后", "把具体地点另行通知已确认的赞助企业和听众"],
            ["10月22日", "确认档位、标识和致辞人。此后海报按已确认名单印制"],
            ["10月24日", "付清赞助款，门票名单交给主办方"],
            ["10月28日", "定稿背景板和开场鸣谢词"],
            ["10月31日 9:00", "讲座开始"],
            ["11月5日前", "发送现场照片"],
        ],
        [3.6, 12.8],
    )
    g.add_para(
        doc,
        "10月22日后仍可谈鸣谢支持和友好支持，但不再保证出现在已送印的海报上，现场鸣谢和报名页仍然保留。",
        size=10.5,
        color=g.MUTED,
        space_before=6,
    )

    h = g.add_h1(doc, "八、企业确认")
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.page_break_before = True
    g.add_para(
        doc,
        "企业勾选一档，填好后回传。主办方据此发出确认函，写明金额、权益和门票名单的提交日期。授权人书面回复或盖章后生效。付款账户在确认函中告知。",
        size=11,
        first_line=True,
    )
    form = g.add_table(
        doc,
        ["项目", "企业填写"],
        [
            ["公司名称", ""],
            ["所选档位", "□ 首席支持 68,000 元"],
            ["所选档位", "□ 战略支持 48,000 元"],
            ["所选档位", "□ 合作支持 38,000 元"],
            ["所选档位", "□ 品牌支持 28,000 元"],
            ["所选档位", "□ 鸣谢支持 18,000 元"],
            ["所选档位", "□ 友好支持 9,800 元"],
            ["行业", "用于判断是否与已确认企业构成直接竞品"],
            ["对外名称", "印在报名页和现场鸣谢中的名称"],
            ["出席人", "姓名、职务、手机。首席支持请注明致辞人"],
            ["发票抬头", "名称、税号"],
            ["授权人", "签字：　　　　　　日期：　　年　　月　　日"],
        ],
        [3.6, 12.8],
    )
    label = form.cell(2, 0)
    for row_index in range(3, 8):
        label = label.merge(form.cell(row_index, 0))
    g.cell_text(
        form.cell(2, 0),
        "所选档位",
        bold=True,
        size=10,
        color=g.NAVY,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        fill="FFFFFF",
    )
    g.add_para(doc, "主办方", size=12, bold=True, color=g.NAVY, space_before=8, space_after=2)
    g.add_para(doc, "上海市杨浦区科技企业联合会　　联系人：胡继刚", size=11, space_after=2)
    g.add_para(doc, "联合主办：首乾书院", size=11, space_after=2)
    g.add_para(
        doc,
        "本方案只用于赞助洽谈。场地确定后另行通知。票档或议程若调整，已确认企业的标识、鸣谢和门票数量不变。",
        size=10,
        color=g.MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=10,
    )
    doc.save(PATH)
    return PATH


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print(build())
