#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成星火社区 OPC 申请材料：DOCX + HTML + PDF。"""

from __future__ import annotations

import html
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from weasyprint import HTML

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
sys.path.insert(0, str(Path(__file__).resolve().parent))
import content as C  # noqa: E402

NAVY = RGBColor(0x1A, 0x3A, 0x5C)
TEAL = RGBColor(0x1F, 0x6B, 0x5C)
GOLD = RGBColor(0xB8, 0x8A, 0x2E)
INK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"
FONT_FALLBACK = "宋体"


def set_run_font(run, size=11, bold=False, color=INK, font=FONT):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), FONT_FALLBACK)
    rFonts.set(qn("w:cs"), font)


def shade(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="C5D0DC"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def v_align(cell, val="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), val)
    tcPr.append(va)


def page_setup(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    return sec


def header_footer(sec, header_text: str):
    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run(header_text)
    set_run_font(run, size=9, color=GOLD)
    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("复兴岛星火社区入驻申请材料  ·  仅供官方审核使用  ·  第 ")
    set_run_font(run, size=8, color=GRAY)
    for kind, val in (("begin", None), ("instr", " PAGE "), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
            fp._p.append(el)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            fp._p.append(el)
    run2 = fp.add_run(" 页")
    set_run_font(run2, size=8, color=GRAY)


def p(
    doc,
    text,
    size=11,
    bold=False,
    color=INK,
    align="left",
    space_after=8,
    space_before=0,
    first_line=None,
    italic=False,
):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.32
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    para.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    run.italic = italic
    return para


def h1(doc, text):
    para = p(doc, text, size=16, bold=True, color=NAVY, space_before=16, space_after=10)
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "B88A2E")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return para


def h2(doc, text):
    return p(doc, text, size=13, bold=True, color=TEAL, space_before=12, space_after=6)


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.clear()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.25
    run = para.add_run(text)
    set_run_font(run, size=11, color=INK)
    return para


def write_cell(cell, text, *, bold=False, size=10, color=INK, fill=None, align="left"):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.space_before = Pt(2)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade(cell, fill)
    set_cell_border(cell)
    v_align(cell, "center")


def add_table(doc, rows, header=True, col_widths=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i].cells[c_i]
            is_header = header and r_i == 0
            write_cell(
                cell,
                str(val),
                bold=is_header,
                size=9.5 if not is_header else 10,
                color=WHITE if is_header else INK,
                fill="1A3A5C" if is_header else ("F7F1E4" if r_i % 2 == 0 else "FFFFFF"),
                align="center" if is_header else "left",
            )
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def cover_block(doc, title, subtitle, extra_lines, note=""):
    p(doc, "复兴岛星火社区  ·  OPC入驻申请", size=12, bold=True, color=GOLD, align="center", space_after=6)
    p(doc, title, size=22, bold=True, color=NAVY, align="center", space_after=6)
    p(doc, subtitle, size=14, bold=True, color=TEAL, align="center", space_after=10)
    p(doc, C.BP_TAG, size=11, color=GRAY, align="center", space_after=16)
    for line in extra_lines:
        p(doc, line, size=12, align="center", space_after=4)
    p(doc, f"{C.DOC_DATE}  {C.DOC_REVISION}", size=12, align="center", space_before=12, space_after=4)
    if note:
        p(doc, note, size=9, color=GRAY, align="center", space_before=10)


def build_guide() -> Document:
    doc = Document()
    sec = page_setup(doc)
    header_footer(sec, "星火社区申请  ·  内部填写说明（勿上传）")
    cover_block(
        doc,
        "填写说明（内部）",
        "只上传02个人简介和03业务计划书",
        [
            f"申请人：{C.PERSON['姓名']}",
            f"项目：{C.PROJECT['项目名称']}",
            f"赛道：{C.PROJECT['赛道']}　　盈利模式：{C.PROJECT['盈利模式勾选']}",
        ],
        note="本稿含填表备忘和履历核验，不要作为申请附件上传。",
    )
    h1(doc, "一、提交什么")
    add_table(
        doc,
        [
            ("文件", "用途"),
            ("02 个人简介", "上传「个人介绍／创始人简介」"),
            ("03 业务计划书", "上传「商业计划书／项目计划」"),
            ("04 字段粘贴稿", "自己对着系统复制，不要整份上传"),
            ("01 本说明", "内部备忘，不要上传"),
        ],
    )
    h1(doc, "二、申请表最终勾选")
    add_table(doc, [("字段", "建议填写"), *C.CORE_FORM_CHOICES])
    h1(doc, "三、填表注意")
    for line in C.FILL_NOTES:
        bullet(doc, line)
    h1(doc, "四、履历核验（仅内部）")
    for line in C.INTERNAL_VERIFY:
        bullet(doc, line)
    return doc


def build_bio() -> Document:
    doc = Document()
    sec = page_setup(doc)
    header_footer(sec, "星火社区申请  ·  个人简介")
    cover_block(
        doc,
        "个人简介",
        C.PERSON["姓名"],
        [
            "复旦大学住房政策研究中心秘书长",
            "上海市杨浦区科技企业联合会执行会长",
            "AIDEN产业智能体创始人",
        ],
    )
    h1(doc, "一、简介")
    for para in C.BIO_FORMAL_PARAS:
        p(doc, para, align="justify", first_line=0.74)
    h1(doc, "二、学习与工作简历")
    for line in C.BIO_RESUME_LINES:
        p(doc, line, size=11, space_after=3)
    h1(doc, "三、社会职务")
    for line in C.BIO_SOCIAL:
        bullet(doc, line)
    h1(doc, "四、公开活动（附出处）")
    for line in C.BIO_PUBLIC:
        bullet(doc, line)
    return doc


def _bp_extra_tables(doc, title: str):
    if title.startswith("三、"):
        add_table(doc, C.PRODUCT_TABLE)
    if title.startswith("四、"):
        add_table(doc, C.PROGRESS_TABLE)
    if title.startswith("五、"):
        add_table(doc, C.IO_TABLE)
    if title.startswith("七、"):
        add_table(doc, C.REVENUE_TABLE)
        add_table(doc, C.ASSUME_TABLE)
        add_table(doc, C.PRICE_TABLE)
        add_table(doc, C.COST_TABLE)
    if title.startswith("八、"):
        add_table(doc, C.AGENT_TABLE)
    if title.startswith("十、"):
        add_table(doc, C.NEED_TABLE)


def build_bp() -> Document:
    doc = Document()
    sec = page_setup(doc)
    header_footer(sec, "星火社区申请  ·  业务计划书")
    cover_block(
        doc,
        C.BP_TITLE,
        C.BP_SUBTITLE,
        [
            f"申请人：{C.PERSON['姓名']}",
            f"赛道：{C.PROJECT['赛道']}",
            f"首期原型：{C.PROJECT['首期原型']}（正在搭建）",
            f"拟定主体：{C.PROJECT['拟定字号']}（待核名）",
        ],
    )
    h1(doc, "目录")
    for title, _ in C.BP_SECTIONS:
        p(doc, title, size=11, space_after=3, color=NAVY)
    for title, paras in C.BP_SECTIONS:
        h1(doc, title)
        for para in paras:
            p(doc, para, align="justify", first_line=0.74)
        _bp_extra_tables(doc, title)
    return doc


def build_form() -> Document:
    doc = Document()
    sec = page_setup(doc)
    header_footer(sec, "星火社区申请  ·  字段粘贴稿（内部）")
    cover_block(
        doc,
        "申请表最终答案",
        "首页只保留与系统字段对应的一稿",
        [
            f"姓名：{C.PERSON['姓名']}",
            f"项目：{C.PROJECT['项目名称']}",
            "赛道=AI智能体开发；已注册=否；营收=50万以下；盈利模式=技术开发+服务收费",
        ],
        note="内部填表用。证件号码和详细门牌请在系统按原件填写。",
    )
    h1(doc, "一、下拉框／选项（最终）")
    add_table(doc, [("字段", "建议填写"), *C.CORE_FORM_CHOICES])
    h1(doc, "二、短字段（最终）")
    add_table(doc, [("字段", "粘贴内容", "备注"), *C.FORM_FIELDS], col_widths=[4.2, 8.2, 3.8])
    h1(doc, "三、长文本最终稿（已按上限计字）")
    for title, text, limit in C.FORM_FINAL_LONG:
        n = C.char_count(text)
        h2(doc, f"{title}　·　{n}字／上限{limit}字")
        p(doc, text, size=10.5, align="justify", space_after=10)
    h1(doc, "四、备选稿（仅当系统字数上限不同）")
    p(doc, "不要把本节上传，也不要在有明确上限时用超限备选替换最终稿。", size=10.5, color=GRAY)
    for title, text in C.FORM_ALT_LONG:
        h2(doc, f"{title}　·　{C.char_count(text)}字")
        p(doc, text, size=10.5, align="justify", space_after=10)
    h1(doc, "五、内部备忘")
    for line in C.FILL_NOTES + C.INTERNAL_VERIFY:
        bullet(doc, line)
    return doc


CSS = """
@page { size: A4; margin: 18mm 16mm 18mm 16mm;
  @bottom-center { content: "复兴岛星火社区入驻申请材料  ·  " counter(page); font-size: 9px; color: #666; }
}
html { font-family: "WenQuanYi Micro Hei", "Noto Sans CJK SC", "Source Han Sans SC", sans-serif; color: #222; }
body { font-size: 11.5pt; line-height: 1.55; }
h1 { color: #1A3A5C; font-size: 18pt; border-bottom: 2px solid #B88A2E; padding-bottom: 4px; margin-top: 22px; }
h2 { color: #1F6B5C; font-size: 13.5pt; margin-top: 16px; }
.cover { text-align: center; margin: 24px 0 28px; }
.kicker { color: #B88A2E; font-weight: 700; letter-spacing: .08em; }
.title { color: #1A3A5C; font-size: 26pt; margin: 8px 0; }
.sub { color: #1F6B5C; font-size: 14pt; font-weight: 700; }
.tag { color: #666; margin: 8px 0 16px; }
.meta { margin: 4px 0; }
.note { color: #666; font-size: 10pt; margin-top: 16px; }
p { margin: 0 0 8pt; text-align: justify; }
ul { margin: 4pt 0 10pt 1.2em; }
li { margin-bottom: 4pt; }
table { width: 100%; border-collapse: collapse; margin: 8px 0 16px; font-size: 10pt; table-layout: fixed; }
th { background: #1A3A5C; color: #fff; padding: 6px 8px; text-align: left; }
td { border: 1px solid #C5D0DC; padding: 6px 8px; vertical-align: top; overflow-wrap: break-word; word-break: break-word; }
tr:nth-child(even) td { background: #F7F1E4; }
.paste { background: #f7f8fa; border: 1px solid #e2e6ea; padding: 8px 10px; white-space: pre-wrap; }
"""


def table_html(rows, header=True) -> str:
    out = ["<table>"]
    for i, row in enumerate(rows):
        tag = "th" if header and i == 0 else "td"
        out.append("<tr>" + "".join(f"<{tag}>{html.escape(str(c)).replace(chr(10), '<br/>')}</{tag}>" for c in row) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def cover_html(title, subtitle, lines, note=""):
    metas = "".join(f'<p class="meta">{html.escape(x)}</p>' for x in lines)
    note_html = f'<p class="note">{html.escape(note)}</p>' if note else ""
    return f"""
<div class="cover">
  <div class="kicker">复兴岛星火社区  ·  OPC入驻申请</div>
  <div class="title">{html.escape(title)}</div>
  <div class="sub">{html.escape(subtitle)}</div>
  <div class="tag">{html.escape(C.BP_TAG)}</div>
  {metas}
  <p class="meta">{html.escape(C.DOC_DATE)}  {html.escape(C.DOC_REVISION)}</p>
  {note_html}
</div>
"""


def html_guide() -> str:
    body = cover_html("填写说明（内部）", "只上传02个人简介和03业务计划书", [
        f"申请人：{C.PERSON['姓名']}",
        f"项目：{C.PROJECT['项目名称']}",
    ], "本稿不要作为申请附件上传。")
    body += "<h1>一、提交什么</h1>" + table_html([
        ("文件", "用途"),
        ("02 个人简介", "上传个人介绍"),
        ("03 业务计划书", "上传商业计划书"),
        ("04 字段粘贴稿", "自己填表，不整份上传"),
        ("01 本说明", "内部备忘，不上传"),
    ])
    body += "<h1>二、申请表最终勾选</h1>" + table_html([("字段", "建议填写"), *C.CORE_FORM_CHOICES])
    body += "<h1>三、填表注意</h1><ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in C.FILL_NOTES) + "</ul>"
    body += "<h1>四、履历核验</h1><ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in C.INTERNAL_VERIFY) + "</ul>"
    return body


def html_bio() -> str:
    body = cover_html("个人简介", C.PERSON["姓名"], [
        "复旦大学住房政策研究中心秘书长",
        "上海市杨浦区科技企业联合会执行会长",
        "AIDEN产业智能体创始人",
    ])
    body += "<h1>一、简介</h1>" + "".join(f"<p>{html.escape(x)}</p>" for x in C.BIO_FORMAL_PARAS)
    body += "<h1>二、学习与工作简历</h1>" + "".join(f"<p>{html.escape(x)}</p>" for x in C.BIO_RESUME_LINES)
    body += "<h1>三、社会职务</h1><ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in C.BIO_SOCIAL) + "</ul>"
    body += "<h1>四、公开活动（附出处）</h1><ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in C.BIO_PUBLIC) + "</ul>"
    return body


def html_bp() -> str:
    body = cover_html(C.BP_TITLE, C.BP_SUBTITLE, [
        f"申请人：{C.PERSON['姓名']}",
        f"赛道：{C.PROJECT['赛道']}",
        f"首期原型：{C.PROJECT['首期原型']}（正在搭建）",
        f"拟定主体：{C.PROJECT['拟定字号']}（待核名）",
    ])
    body += "<h1>目录</h1><ul>" + "".join(f"<li>{html.escape(t)}</li>" for t, _ in C.BP_SECTIONS) + "</ul>"
    for title, paras in C.BP_SECTIONS:
        body += f"<h1>{html.escape(title)}</h1>" + "".join(f"<p>{html.escape(x)}</p>" for x in paras)
        if title.startswith("三、"):
            body += table_html(C.PRODUCT_TABLE)
        if title.startswith("四、"):
            body += table_html(C.PROGRESS_TABLE)
        if title.startswith("五、"):
            body += table_html(C.IO_TABLE)
        if title.startswith("七、"):
            body += table_html(C.REVENUE_TABLE) + table_html(C.ASSUME_TABLE) + table_html(C.PRICE_TABLE) + table_html(C.COST_TABLE)
        if title.startswith("八、"):
            body += table_html(C.AGENT_TABLE)
        if title.startswith("十、"):
            body += table_html(C.NEED_TABLE)
    return body


def html_form() -> str:
    body = cover_html("申请表最终答案", "首页只保留与系统字段对应的一稿", [
        f"姓名：{C.PERSON['姓名']}",
        f"项目：{C.PROJECT['项目名称']}",
        "赛道=AI智能体开发；盈利模式=技术开发+服务收费",
    ], "内部填表用。证件号码和详细门牌请在系统按原件填写。")
    body += "<h1>一、下拉框／选项（最终）</h1>" + table_html([("字段", "建议填写"), *C.CORE_FORM_CHOICES])
    body += "<h1>二、短字段（最终）</h1>" + table_html([("字段", "粘贴内容", "备注"), *C.FORM_FIELDS])
    body += "<h1>三、长文本最终稿</h1>"
    for title, text, limit in C.FORM_FINAL_LONG:
        n = C.char_count(text)
        body += f"<h2>{html.escape(title)}  ·  {n}字／上限{limit}字</h2><p class='paste'>{html.escape(text)}</p>"
    body += "<h1>四、备选稿</h1>"
    for title, text in C.FORM_ALT_LONG:
        body += f"<h2>{html.escape(title)}  ·  {C.char_count(text)}字</h2><p class='paste'>{html.escape(text)}</p>"
    body += "<h1>五、内部备忘</h1><ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in C.FILL_NOTES + C.INTERNAL_VERIFY) + "</ul>"
    return body


def wrap_html(title: str, body: str) -> str:
    return f"<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{CSS}</style></head><body>{body}</body></html>"


DOCS = [
    ("01_填写说明与口径备忘", "填写说明（内部）", build_guide, html_guide),
    ("02_个人简介", "个人简介", build_bio, html_bio),
    ("03_业务计划书", "业务计划书", build_bp, html_bp),
    ("04_申请表可粘贴字段", "申请表最终答案", build_form, html_form),
]


def save_docx(doc: Document, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def save_pdf(html_doc: str, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html_doc, base_url=str(ROOT)).write_pdf(str(path))


def write_markdown_form(path: Path):
    lines = [
        "# 星火社区申请表 · 最终答案（内部填表）",
        "",
        f"{C.DOC_DATE} {C.DOC_REVISION}。只把本节复制进系统。02、03分开上传。证件号和门牌不要从本文复制。",
        "",
        "## 下拉框／选项",
        "",
        "| 字段 | 建议填写 |",
        "| --- | --- |",
    ]
    for k, v in C.CORE_FORM_CHOICES:
        lines.append(f"| {k} | {v} |")
    lines += ["", "## 短字段", "", "| 字段 | 粘贴内容 | 备注 |", "| --- | --- | --- |"]
    for k, v, note in C.FORM_FIELDS:
        vv = str(v).replace("|", "\\|").replace("\n", "<br>")
        lines.append(f"| {k} | {vv} | {note} |")
    lines += ["", "## 长文本最终稿", ""]
    for title, text, limit in C.FORM_FINAL_LONG:
        n = C.char_count(text)
        flag = "超限" if n > limit else "未超限"
        lines += [f"### {title}（{n}字／上限{limit}字 · {flag}）", "", text, ""]
    lines += ["", "## 备选稿（系统上限不同时再用）", ""]
    for title, text in C.FORM_ALT_LONG:
        lines += [f"### {title}（{C.char_count(text)}字）", "", text, ""]
    lines += ["", "## 内部备忘", ""] + [f"- {x}" for x in C.FILL_NOTES + C.INTERNAL_VERIFY]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_all():
    OUT.mkdir(parents=True, exist_ok=True)
    for old in OUT.glob("星火社区_00_全套材料_上传用*"):
        old.unlink()
        print(f"已删除 {old.name}")
    written = []
    for stem, title, docx_fn, html_fn in DOCS:
        docx_path = OUT / f"星火社区_{stem}.docx"
        pdf_path = OUT / f"星火社区_{stem}.pdf"
        html_path = OUT / f"星火社区_{stem}.html"
        save_docx(docx_fn(), docx_path)
        html_doc = wrap_html(title, html_fn())
        html_path.write_text(html_doc, encoding="utf-8")
        save_pdf(html_doc, pdf_path)
        written.extend([docx_path, pdf_path, html_path])
        print(f"已生成 {docx_path.name} / {pdf_path.name}")
    md_path = OUT / "星火社区_04_申请表可粘贴字段.md"
    write_markdown_form(md_path)
    print(f"已生成 {md_path.name}")
    written.append(md_path)
    return written


if __name__ == "__main__":
    build_all()
