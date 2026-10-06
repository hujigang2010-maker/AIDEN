#!/usr/bin/env python3
"""将 manuscript/ 下的 Markdown 书稿合并，并生成 Word 版《以韩为鉴》。

输出：
- output/以韩为鉴_书稿全文.md
- output/以韩为鉴_书稿.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MS_DIR = ROOT / "manuscript"
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)

TITLE = "以韩为鉴"
SUBTITLE = "压缩式现代化之后的就业、民生与周期"
TAGLINE = "一个追赶型经济体的高光与代价 · 供中国借鉴"
BODY_FONT = "宋体"
HEAD_FONT = "黑体"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x7F, 0x8C, 0x8D)


def set_run_font(run, name=BODY_FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(paragraph, text, size=10.5, color=None):
    """处理 **加粗** 与 `代码` 两种行内标记。"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        content = content.replace("`", "")
        run = paragraph.add_run(content)
        set_run_font(run, size=size, bold=bold, color=color)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)
    for level, size in ((1, 20), (2, 15), (3, 12.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = HEAD_FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(14 if level < 3 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_run_font(run, size=9, color=GREY)


def add_static_toc(doc, files):
    """静态目录：列出篇与章，便于在任何阅读器中直接查看。"""
    for f in files:
        if f.name.startswith("00_"):
            continue
        for line in f.read_text(encoding="utf-8").splitlines():
            m = re.match(r"^(#{1,2})\s+(.*)", line.strip())
            if not m:
                continue
            level, title = len(m.group(1)), m.group(2)
            if level == 2 and not title.startswith(("第", "附录")):
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.2
            if level == 1:
                p.paragraph_format.space_before = Pt(6)
                set_run_font(p.add_run(title), HEAD_FONT, 11, True, NAVY)
            else:
                p.paragraph_format.left_indent = Cm(0.8)
                set_run_font(p.add_run(title), BODY_FONT, 10)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "（在 Word 中右键此处选择“更新域”即可生成目录）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, placeholder, fld_end):
        run._r.append(el)


def add_cover(doc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(TITLE), HEAD_FONT, 44, True, NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(SUBTITLE), HEAD_FONT, 18, False, NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(TAGLINE), BODY_FONT, 12, False, GREY)
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("书稿 · 2026 年 10 月"), BODY_FONT, 11, False, GREY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_cell_shading(cell, hex_color):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], text, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
        for run in cell.paragraphs[0].runs:
            run.bold = True
        set_cell_shading(cell, "1F3A5F")
    for r_idx, row in enumerate(body):
        cells = table.add_row().cells
        for i in range(len(header)):
            text = row[i] if i < len(row) else ""
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], text, size=9)
            if r_idx % 2 == 1:
                set_cell_shading(cells[i], "EEF3F8")
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(1)
    doc.add_paragraph()


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def render_markdown(doc, text, base_dir, state):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [parse_table_row(b) for b in block if not re.match(r"^\|\s*:?-{3,}", b.strip())]
            add_table(doc, rows)
            continue
        img = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if img:
            caption, rel = img.groups()
            path = (base_dir / rel).resolve()
            if path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Cm(15.5))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run_font(cap.add_run(caption), size=9, color=GREY)
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)", stripped)
        if heading:
            level = len(heading.group(1))
            if level == 1:
                if state["h1_seen"]:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                state["h1_seen"] = True
            doc.add_heading(heading.group(2), level=level)
            i += 1
            continue
        if stripped == "---":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run("◆ ◆ ◆"), size=10, color=GREY)
            i += 1
            continue
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, stripped.lstrip("> ").strip(), size=10, color=NAVY)
            i += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, f"{numbered.group(1)}. {numbered.group(2)}")
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        add_inline(p, stripped)
        i += 1


def main():
    files = sorted(MS_DIR.glob("[0-9][0-9]_*.md"))
    combined = []
    for f in files:
        combined.append(f.read_text(encoding="utf-8").strip())
    full_md = "\n\n".join(combined).replace("](figures/", "](../manuscript/figures/")
    (OUT_DIR / "以韩为鉴_书稿全文.md").write_text(
        f"# {TITLE}\n\n> {SUBTITLE}——{TAGLINE}\n\n" + full_md + "\n", encoding="utf-8"
    )

    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(26), Cm(18.4)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2))
    section.top_margin = section.bottom_margin = Cm(2.2)
    setup_styles(doc)
    add_cover(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("目  录"), HEAD_FONT, 18, True, NAVY)
    add_static_toc(doc, files)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    set_run_font(p.add_run("带页码目录："), BODY_FONT, 9, False, GREY)
    add_toc(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    add_page_number_footer(body_section)
    state = {"h1_seen": False}
    for f in files:
        render_markdown(doc, f.read_text(encoding="utf-8"), MS_DIR, state)

    out = OUT_DIR / "以韩为鉴_书稿.docx"
    doc.save(out)
    chars = sum(len(re.sub(r"\s", "", t)) for t in combined)
    print(f"已生成 {out.name}，Markdown 源文件 {len(files)} 个，正文约 {chars} 字")


if __name__ == "__main__":
    main()
