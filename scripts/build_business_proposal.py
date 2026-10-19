#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成发给新谷湾的商务合作方案。

只保留已确定的合作内容、费用和进度。
"""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft YaHei"
W, H = 13.333, 7.5
TOTAL = 8

NAVY = "16324F"
INK = "1C2833"
MUTED = "5E6D7A"
TEAL = "1F5C57"
TEAL_SOFT = "E6F1EF"
GOLD = "8C6A2F"
BG = "F6F3EC"
WHITE = "FFFFFF"
CARD = "FFFCF7"
LINE = "E4DACC"
PAPER = "EFEAE1"


def rgb(value):
    return RGBColor.from_string(value)


def font_run(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rpr, qn(tag))
        node.set("typeface", FONT)


def rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    return shape


def paint(slide, color):
    shape = rect(slide, 0, 0, W, H, color)
    tree = slide.shapes._spTree
    element = shape._element
    tree.remove(element)
    tree.insert(2, element)


def write(slide, paragraphs, x, y, w, h, anchor="t"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    frame._txBody.bodyPr.set("anchor", anchor)
    for index, para in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = para.get("align", PP_ALIGN.LEFT)
        paragraph.space_before = Pt(para.get("before", 0))
        paragraph.space_after = Pt(para.get("after", 0))
        if "spacing" in para:
            paragraph.line_spacing = para["spacing"]
        parts = para["text"].split("\n")
        for part_index, part in enumerate(parts):
            if part_index:
                etree.SubElement(paragraph._p, qn("a:br"))
            run = paragraph.add_run()
            run.text = part
            font_run(run, para.get("size", 15), para.get("color", INK), para.get("bold", False))
    return box


def t(text, size=15, color=INK, bold=False, align=PP_ALIGN.LEFT, before=0, after=0, spacing=1.08):
    return {
        "text": text,
        "size": size,
        "color": color,
        "bold": bold,
        "align": align,
        "before": before,
        "after": after,
        "spacing": spacing,
    }


def blank(prs, color=BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint(slide, color)
    return slide


def footer(slide, page):
    rect(slide, 0.5, 7.02, 12.33, 0.01, LINE)
    write(slide, [t("上海市杨浦区科技企业联合会  ·  合作方案", 11, MUTED)], 0.5, 7.1, 9.2, 0.28)
    write(
        slide,
        [t(f"{page:02d}  /  {TOTAL:02d}", 11, MUTED, align=PP_ALIGN.RIGHT)],
        10.4,
        7.1,
        2.43,
        0.28,
    )


def header(slide, kicker, title):
    rect(slide, 0, 0, W, 0.07, TEAL)
    write(slide, [t(kicker, 12, GOLD, True)], 0.52, 0.22, 12.2, 0.26)
    write(slide, [t(title, 26, NAVY, True, spacing=1.0)], 0.5, 0.48, 12.3, 0.48)
    rect(slide, 0.52, 1.06, 0.86, 0.03, GOLD)


def card(slide, x, y, w, h, fill=CARD, line=LINE):
    return rect(slide, x, y, w, h, fill, line)


def slide_cover(prs):
    slide = blank(prs, NAVY)
    rect(slide, 0, 0, 0.14, H, GOLD)
    write(slide, [t("合作方案", 14, "E4C98A", True)], 0.72, 1.35, 8, 0.32)
    write(
        slide,
        [
            t("科技企业成长与", 40, WHITE, True, spacing=1.0),
            t("产业协同合作", 40, WHITE, True, before=6, spacing=1.0),
        ],
        0.7,
        1.85,
        12,
        1.8,
    )
    rect(slide, 0.74, 4.0, 1.4, 0.035, GOLD)
    write(
        slide,
        [
            t("致  成都新谷投资集团  ·  上海新谷湾", 18, "F4F7F8"),
            t("上海市杨浦区科技企业联合会", 18, "F4F7F8", before=8),
        ],
        0.72,
        4.25,
        11,
        0.9,
    )
    write(slide, [t("2026年10月", 14, "C9D4DE")], 0.72, 6.35, 4, 0.3)


def slide_offer(prs):
    slide = blank(prs)
    header(slide, "01    合作内容", "三项合作同时启动")
    footer(slide, 2)
    items = [
        ("01", "副会长单位", "10,000 元 / 年", "以上海新谷湾投资有限公司入会，任期一年。"),
        ("02", "三个月试点", "20,000 元", "交付企业名单、投资机构名单、到场对接、一场闭门会和台账。"),
        ("03", "服务站挂牌", "含在试点内", "科技企业成长与产业协同服务站设在上海新谷国际，挂牌一年。"),
    ]
    for index, (num, title, price, body) in enumerate(items):
        y = 1.32 + index * 1.55
        card(slide, 0.5, y, 12.33, 1.42)
        rect(slide, 0.5, y, 0.08, 1.42, TEAL)
        write(slide, [t(num, 18, GOLD, True)], 0.85, y + 0.28, 0.8, 0.36)
        write(slide, [t(title, 20, NAVY, True)], 1.8, y + 0.26, 4.2, 0.4)
        write(slide, [t(body, 15, INK)], 1.8, y + 0.76, 6.6, 0.42)
        write(slide, [t(price, 18, TEAL, True, align=PP_ALIGN.RIGHT)], 8.6, y + 0.46, 3.9, 0.42)
    write(
        slide,
        [t("企业与新谷签订租赁合同后，按该合同一个月基础租金结算成交服务费。", 14, MUTED)],
        0.55,
        6.1,
        12.2,
        0.4,
    )


def slide_delivery(prs):
    slide = blank(prs)
    header(slide, "02    试点交付", "三个月内交付四项")
    footer(slide, 3)
    items = [
        ("01", "企业名单", "从《杨浦一百》和联合会科技企业资源中提交有办公需求的企业。"),
        ("02", "投资机构名单", "提交投资阶段与企业相符的投资机构。"),
        ("03", "到场对接", "企业和投资机构到上海新谷国际实地交流，联合会记录到场情况。"),
        ("04", "闭门会与台账", "试点期内举办一场闭门会。会后提交对接台账。"),
    ]
    for index, (num, title, body) in enumerate(items):
        col, row = index % 2, index // 2
        x, y = 0.5 + col * 6.4, 1.35 + row * 2.6
        card(slide, x, y, 6.2, 2.4)
        write(slide, [t(num, 16, GOLD, True)], x + 0.28, y + 0.28, 1.2, 0.32)
        write(slide, [t(title, 22, NAVY, True)], x + 0.28, y + 0.72, 5.6, 0.42)
        write(slide, [t(body, 15, INK, spacing=1.15)], x + 0.28, y + 1.3, 5.6, 0.85)


def slide_station(prs):
    slide = blank(prs)
    header(slide, "03    服务站", "挂牌一年，作为对接场地")
    footer(slide, 4)
    card(slide, 0.5, 1.35, 12.33, 1.7, NAVY, NAVY)
    write(slide, [t("名称", 12, "E4C98A", True)], 0.78, 1.55, 3, 0.26)
    write(slide, [t("科技企业成长与产业协同服务站", 26, WHITE, True)], 0.78, 1.95, 11.5, 0.5)
    write(slide, [t("上海新谷国际    ·    挂牌一年    ·    与试点同步挂牌", 15, "D5DEE6")], 0.78, 2.55, 11.5, 0.32)
    rows = [
        ("地点", "上海新谷国际"),
        ("场地", "新谷提供活动场地和接待"),
        ("组织", "联合会负责邀请、主持和记录"),
        ("活动", "试点期内举办一场闭门会，服务费已含本场"),
    ]
    for index, (label, value) in enumerate(rows):
        y = 3.28 + index * 0.88
        card(slide, 0.5, y, 12.33, 0.78)
        write(slide, [t(label, 15, TEAL, True)], 0.78, y + 0.22, 1.6, 0.34)
        write(slide, [t(value, 16, NAVY)], 2.6, y + 0.2, 9.8, 0.38)


def slide_roles(prs):
    slide = blank(prs)
    header(slide, "04    分工", "双方各司其职")
    footer(slide, 5)
    columns = [
        (
            "联合会",
            [
                "办理副会长入会",
                "提交企业名单和投资机构名单",
                "组织到场对接和闭门会",
                "提交对接台账和试点成果",
                "租赁合同签署后出具成交服务费结算单",
            ],
        ),
        (
            "新谷",
            [
                "指定一名试点对接人",
                "提供在租房源、面积和租金",
                "提供服务站场地并安排接待",
                "参加闭门会",
                "按合同约定支付会费、试点费和成交服务费",
            ],
        ),
    ]
    for index, (title, lines) in enumerate(columns):
        x = 0.5 + index * 6.4
        card(slide, x, 1.35, 6.2, 5.4)
        rect(slide, x, 1.35, 6.2, 0.7, TEAL if index == 0 else NAVY)
        write(slide, [t(title, 20, WHITE, True)], x + 0.3, 1.52, 5.5, 0.4)
        for line_index, line in enumerate(lines):
            y = 2.35 + line_index * 0.78
            write(slide, [t(f"{line_index + 1}", 14, GOLD, True)], x + 0.32, y, 0.4, 0.32)
            write(slide, [t(line, 16, INK)], x + 0.85, y, 5.0, 0.55)


def slide_fees(prs):
    slide = blank(prs)
    header(slide, "05    费用", "三项费用，一次说清")
    footer(slide, 6)
    columns = [3.4, 3.5, 5.43]
    headers = ["项目", "金额", "支付时间"]
    data = [
        ["副会长会费", "10,000 元", "入会时支付，有效期一年"],
        ["试点服务费", "20,000 元", "协议签署后支付"],
        ["成交服务费", "1 个月基础租金", "每份租赁合同签署后支付"],
    ]
    x0, y0, height = 0.5, 1.45, 0.72
    cursor = x0
    for width, label in zip(columns, headers):
        rect(slide, cursor, y0, width, height, NAVY)
        write(slide, [t(label, 15, WHITE, True)], cursor + 0.22, y0 + 0.2, width - 0.36, 0.34)
        cursor += width
    for row_index, row in enumerate(data):
        y = y0 + height + row_index * 0.85
        fill = WHITE if row_index % 2 == 0 else PAPER
        cursor = x0
        for col_index, (width, label) in enumerate(zip(columns, row)):
            rect(slide, cursor, y, width, 0.85, fill, LINE)
            write(
                slide,
                [t(label, 16, NAVY if col_index == 0 else INK, bold=(col_index == 0))],
                cursor + 0.22,
                y + 0.24,
                width - 0.36,
                0.4,
            )
            cursor += width
    card(slide, 0.5, 5.15, 12.33, 1.55, TEAL_SOFT, "C9DDD9")
    write(slide, [t("说明", 14, TEAL, True)], 0.78, 5.35, 11.5, 0.28)
    write(
        slide,
        [
            t("基础租金指租赁合同载明的月租金，不含物业费。", 16, INK, before=6),
            t("试点服务费已包含一场闭门会。", 16, INK, before=4),
        ],
        0.78,
        5.7,
        11.7,
        0.8,
    )


def slide_schedule(prs):
    slide = blank(prs)
    header(slide, "06    进度", "签署后九十日完成")
    footer(slide, 7)
    rows = [
        ("第 10 日", "双方对接人到位。新谷提交房源清单，联合会提交首批企业名单和投资机构名单。"),
        ("第 20 日", "完成入会。服务站挂牌。"),
        ("第 40 日", "完成首批企业和投资机构到场对接。"),
        ("第 60 日", "举办闭门会，提交对接台账。"),
        ("第 90 日", "提交试点成果。"),
    ]
    for index, (when, task) in enumerate(rows):
        y = 1.35 + index * 1.05
        card(slide, 0.5, y, 12.33, 0.92)
        rect(slide, 0.5, y, 0.08, 0.92, TEAL)
        write(slide, [t(when, 16, NAVY, True)], 0.85, y + 0.28, 1.8, 0.36)
        write(slide, [t(task, 16, INK)], 2.9, y + 0.26, 9.6, 0.42)


def slide_confirm(prs):
    slide = blank(prs)
    header(slide, "07    请确认", "确认三项后，联合会准备协议")
    footer(slide, 8)
    items = [
        ("01", "入会主体", "上海新谷湾投资有限公司"),
        ("02", "试点对接人", "姓名、职务、手机"),
        ("03", "挂牌日期", "服务站在上海新谷国际的挂牌日"),
    ]
    for index, (num, title, body) in enumerate(items):
        x = 0.5 + index * 4.2
        card(slide, x, 1.4, 4.02, 2.7)
        write(slide, [t(num, 16, GOLD, True)], x + 0.24, 1.6, 3.5, 0.3)
        write(slide, [t(title, 22, NAVY, True)], x + 0.24, 2.1, 3.5, 0.45)
        write(slide, [t(body, 16, INK)], x + 0.24, 2.75, 3.5, 0.9)
    card(slide, 0.5, 4.35, 12.33, 2.35, NAVY, NAVY)
    write(slide, [t("联系人", 13, "E4C98A", True)], 0.82, 4.6, 11, 0.28)
    write(slide, [t("朱震", 28, WHITE, True)], 0.82, 5.05, 11, 0.5)
    write(slide, [t("上海市杨浦区科技企业联合会  法人", 16, "D5DEE6")], 0.82, 5.65, 11, 0.36)


def build():
    presentation = Presentation()
    presentation.slide_width = Inches(W)
    presentation.slide_height = Inches(H)
    presentation.core_properties.title = "科技企业成长与产业协同合作方案"
    presentation.core_properties.subject = "上海市杨浦区科技企业联合会致成都新谷投资集团"
    presentation.core_properties.author = "上海市杨浦区科技企业联合会"
    presentation.core_properties.category = "合作方案"
    for builder in (
        slide_cover,
        slide_offer,
        slide_delivery,
        slide_station,
        slide_roles,
        slide_fees,
        slide_schedule,
        slide_confirm,
    ):
        builder(presentation)
    if len(presentation.slides) != TOTAL:
        raise SystemExit(f"页数不符：{len(presentation.slides)} != {TOTAL}")
    output = "/workspace/output/新谷湾合作方案-商务版.pptx"
    presentation.save(output)
    print(output)


if __name__ == "__main__":
    build()
