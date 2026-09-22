# -*- coding: utf-8 -*-
"""生成泰隆银行 2026 年四季度参访与高端小局执行表。

单场费用硬顶 5000 元；目标是存量客户满意和新客户获取。
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
NAVY = "0A1E36"
LIGHT = "E6EEF3"
ZEBRA = "F2F6F9"
OK = "E6F4EC"
WARN = "F8F0E2"
RED_SOFT = "FDECEC"
MUTED = "6A8490"
WHITE = "FFFFFF"

FONT = "微软雅黑"
thin = Side(style="thin", color="B7C7CF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(name=FONT, size=11, bold=True, color=WHITE)
BODY_FONT = Font(name=FONT, size=10.5)
TITLE_FONT = Font(name=FONT, size=14, bold=True, color=NAVY)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
LIGHT_FILL = PatternFill("solid", fgColor=LIGHT)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
OK_FILL = PatternFill("solid", fgColor=OK)
WARN_FILL = PatternFill("solid", fgColor=WARN)
RED_FILL = PatternFill("solid", fgColor=RED_SOFT)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()


def style_sheet(ws, title, headers, rows, widths, note=None, status_col=None, row_height=36):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    c = ws.cell(1, 1, title)
    c.font = TITLE_FONT
    c.fill = LIGHT_FILL
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36

    for j, h in enumerate(headers, 1):
        cell = ws.cell(2, j, h)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[2].height = 26

    for i, row in enumerate(rows, 3):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, v)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER
            if (i - 3) % 2 == 1 and (status_col is None or j != status_col):
                cell.fill = ZEBRA_FILL
            if status_col and j == status_col and isinstance(v, str):
                if any(k in v for k in ("合格", "已完成", "已确认", "必须")):
                    cell.fill = OK_FILL
                elif any(k in v for k in ("待填", "待启动", "待确认", "待邀")):
                    cell.fill = WARN_FILL
                elif any(k in v for k in ("超支", "禁止", "高")):
                    cell.fill = RED_FILL
        ws.row_dimensions[i].height = row_height

    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if note:
        r = len(rows) + 4
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        cell = ws.cell(r, 1, note)
        cell.font = Font(name=FONT, size=10, italic=True, color=MUTED)
        cell.alignment = WRAP
        ws.row_dimensions[r].height = 48
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(rows) + 2}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return len(rows) + 2


def add_list(ws, cell_range, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)


# ---------------------------------------------------------------------------
# 00 总览
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "00-总览"
style_sheet(
    ws,
    "2026年10—12月客户参访与高端小局 · 执行总览（单场不超过5000元）",
    ["模块", "执行口径", "核对数字", "谁负责", "状态"],
    [
        ["费用", "每场决算 ≤5000 元，超了减项，不追加", "7 场合计硬顶 35000 元", "项目经理", "必须"],
        ["场次", "10/22 起，周四下午；12 月最密", "参访 4 场 + 高端小局 3 场", "项目经理", "待确认"],
        ["存量客户", "客户经理陪同，次日问感受", "累计不少于 60 人次，满意度 ≥90%", "客户经理", "待启动"],
        ["新客户", "提名邀请，5 个工作日内首联", "到场不少于 25 家，首联率 100%", "客户经理", "待启动"],
        ["场地", "银行贵宾室或企业现场，不租酒店", "场地费目标为 0", "会务", "待启动"],
        ["礼品与嘉宾", "默认无伴手礼、无付费嘉宾、无摄影团队", "有例外须合规书面同意且仍 ≤5000", "银行合规", "必须"],
        ["取舍", "陪同排不开时，11/26 并入 12/17", "12 月 3 场不减", "项目经理", "待确认"],
        ["合规", "不承诺额度、利率、放款和领导出席", "话术按 14 表执行", "银行合规", "必须"],
    ],
    [14, 42, 36, 14, 12],
    note="今天为 2026-09-22。8 月、9 月上旬场次不再回补。本表不构成授信或业务承诺。",
    status_col=5,
    row_height=32,
)

# ---------------------------------------------------------------------------
# 01 日历
# ---------------------------------------------------------------------------
ws = wb.create_sheet("01-七场日历")
style_sheet(
    ws,
    "七场日历｜周四下午 13:30—16:30｜客户人数不含工作人员",
    ["序", "日期", "星期", "类型", "主题", "客户人数", "存量", "新客", "目的", "场地", "预算上限", "状态"],
    [
        [1, "2026-10-22", "周四", "精品参访", "智能制造或具身智能企业", 16, 12, 4, "让老客户感到被重视", "企业现场", 5000, "待确认"],
        [2, "2026-11-05", "周四", "高端小局", "重点客户私享问诊", 12, 12, 0, "把经营问题听清楚", "银行贵宾室", 5000, "待确认"],
        [3, "2026-11-19", "周四", "精品参访", "链主企业或AI应用现场", 16, 10, 6, "老客户带新客户", "企业现场", 5000, "待确认"],
        [4, "2026-11-26", "周四", "高端小局", "意向新客闭门局", 12, 4, 8, "新客户建立信任", "银行贵宾室", 5000, "待确认"],
        [5, "2026-12-03", "周四", "精品参访", "跨境窗口或园区", 16, 10, 6, "把出海变成看得见", "园区或机构现场", 5000, "待确认"],
        [6, "2026-12-10", "周四", "精品参访", "客户点名的标杆企业", 18, 18, 0, "年底答谢", "企业现场", 5000, "待确认"],
        [7, "2026-12-17", "周四", "高端小局", "年度收官与次年预约", 16, 10, 6, "感谢并约下一次", "银行贵宾室", 5000, "待确认"],
    ],
    [6, 14, 8, 12, 28, 12, 10, 10, 24, 18, 12, 12],
    note="相邻两场名单默认不重复。一位存量客户全季度被邀请不超过 3 次。11/26 若陪同排不开，并入 12/17，12 月场次不减。",
    status_col=12,
    row_height=30,
)

# ---------------------------------------------------------------------------
# 02 预算模板
# ---------------------------------------------------------------------------
ws = wb.create_sheet("02-预算模板")
style_sheet(
    ws,
    "三套预算模板（元）｜单项可调，合计不得超过 5000",
    ["费用项", "参访16人", "高端小局12人", "收官小局16人", "能不能动", "说明"],
    [
        ["场地", 0, 0, 0, "固定为0", "企业现场或银行贵宾室，不租酒店"],
        ["交通", 1800, 0, 0, "参访可调", "商务车或公共交通，按实际票据"],
        ["茶水或茶叙", 600, 1800, 2400, "可减不可加过线", "小局按约150元/人，不上宴请"],
        ["物料", 400, 300, 400, "可减", "行程单、胸卡、服务卡，数字材料优先"],
        ["应急", 400, 400, 400, "不用不花", "只在现场发生时动用"],
        ["建议合计", 3200, 2500, 3200, "低于硬顶", "余量不主动花完"],
        ["硬顶", 5000, 5000, 5000, "不可突破", "伴手礼、付费嘉宾、摄影团队默认不做"],
    ],
    [16, 14, 16, 16, 16, 42],
    note="若合规书面同意采购伴手礼，从余量列支，整场仍不得超过 5000 元。",
    row_height=30,
)

# ---------------------------------------------------------------------------
# 03 决算闸门
# ---------------------------------------------------------------------------
ws = wb.create_sheet("03-决算闸门")
ws.merge_cells("A1:L1")
ws["A1"] = "单场决算闸门｜填写实际金额后自动判定｜任一科目或合计大于 5000 即超支"
ws["A1"].font = TITLE_FONT
ws["A1"].fill = LIGHT_FILL
ws["A1"].alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 36

gate_headers = ["场次", "日期", "类型", "场地", "交通", "茶歇", "物料", "应急", "其他", "合计", "判定", "超支后的动作"]
for j, h in enumerate(gate_headers, 1):
    cell = ws.cell(2, j, h)
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
    cell.border = BORDER
ws.row_dimensions[2].height = 26

gate_rows = [
    ("10/22 参访", "2026-10-22", "精品参访"),
    ("11/5 小局", "2026-11-05", "高端小局"),
    ("11/19 参访", "2026-11-19", "精品参访"),
    ("11/26 小局", "2026-11-26", "高端小局"),
    ("12/3 参访", "2026-12-03", "精品参访"),
    ("12/10 参访", "2026-12-10", "精品参访"),
    ("12/17 收官", "2026-12-17", "高端小局"),
]
for i, (name, date, kind) in enumerate(gate_rows):
    r = 3 + i
    ws.cell(r, 1, name).font = BODY_FONT
    ws.cell(r, 2, date).font = BODY_FONT
    ws.cell(r, 3, kind).font = BODY_FONT
    for col in range(1, 13):
        cell = ws.cell(r, col)
        cell.border = BORDER
        cell.alignment = CENTER
        cell.font = BODY_FONT
    ws.cell(r, 10, f"=SUM(D{r}:I{r})")
    ws.cell(r, 11, f'=IF(COUNT(D{r}:I{r})=0,"待填",IF(J{r}>5000,"超支","合格"))')
    ws.cell(r, 12, f'=IF(K{r}="超支","下一场改为零采购","")')
    ws.row_dimensions[r].height = 24

ws.cell(10, 1, "七场合计")
ws["A10"].font = Font(name=FONT, size=11, bold=True, color=NAVY)
ws["J10"] = "=SUM(J3:J9)"
ws["K10"] = '=IF(J10>35000,"超全年硬顶","全年硬顶内")'
for col in range(1, 13):
    ws.cell(10, col).border = BORDER
    ws.cell(10, col).font = Font(name=FONT, size=11, bold=True)
    ws.cell(10, col).fill = LIGHT_FILL
ws.row_dimensions[10].height = 26

red_font = Font(name=FONT, size=10.5, bold=True, color="9C2F2F")
green_font = Font(name=FONT, size=10.5, bold=True, color="1F7A4D")
ws.conditional_formatting.add("K3:K9", FormulaRule(formula=['K3="超支"'], fill=RED_FILL, font=red_font))
ws.conditional_formatting.add("K3:K9", FormulaRule(formula=['K3="合格"'], fill=OK_FILL, font=green_font))
ws.conditional_formatting.add("J3:J9", CellIsRule(operator="greaterThan", formula=["5000"], fill=RED_FILL))

dv_money = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dv_money.error = "金额不能为负"
dv_money.errorTitle = "金额"
dv_money.add("D3:I9")
ws.add_data_validation(dv_money)

note = ws.cell(12, 1, "填写 D–I 列实际支出。合计大于 5000 时判定为超支，下一场只保留零成本场地和自带茶水。其他列请不要写入礼品，除非合规已书面同意。")
note.font = Font(name=FONT, size=10, italic=True, color=MUTED)
ws.merge_cells("A12:L12")
ws.row_dimensions[12].height = 36

widths = [16, 14, 12, 10, 10, 10, 10, 10, 10, 12, 14, 28]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A3"
ws.auto_filter.ref = "A2:L9"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddHeader.left.text = "泰隆银行客户参访决算闸门"
ws.print_title_rows = "1:2"

# ---------------------------------------------------------------------------
# 04 首场倒排
# ---------------------------------------------------------------------------
ws = wb.create_sheet("04-首场倒排")
style_sheet(
    ws,
    "第一场 10月22日参访倒排｜自 9月23日启动",
    ["日期", "星期", "动作", "交付物", "负责人角色", "姓名", "状态"],
    [
        ["2026-09-23", "周三", "确认 7 场框架和单场 5000 元硬顶", "会议纪要", "四方联系人", "", "待启动"],
        ["2026-09-24", "周四", "每位客户经理提交重点客户 8 家、意向新客 5 家", "客户名单", "客户经理", "", "待启动"],
        ["2026-09-28", "周一", "储备至少 3 家零场地费、愿意接待的企业", "候选企业名单", "项目办/科企联", "", "待启动"],
        ["2026-09-30", "周三", "合规确认邀约话术；默认不采购伴手礼", "口径确认", "银行合规", "", "待启动"],
        ["2026-10-08", "周四", "国庆后敲定 10/22 被访企业、线路和客户名单", "任务书一页", "项目经理", "", "待启动"],
        ["2026-10-09", "周五", "客户经理一对一发出邀请，不在公开群海投", "邀约记录", "客户经理", "", "待启动"],
        ["2026-10-14", "周三", "确认名单：存量约 12 人、新客约 4 人", "出席名单", "项目经理", "", "待启动"],
        ["2026-10-16", "周五", "客户经理陪同表；服务卡印制不超过物料预算", "陪同表", "项目办", "", "待启动"],
        ["2026-10-19", "周一", "预走访：动线 90 分钟、无拍摄区、安全口", "预走访记录", "会务", "", "待启动"],
        ["2026-10-21", "周三", "电话确认到场；核对交通安排仍在预算内", "到场确认", "客户经理", "", "待启动"],
        ["2026-10-22", "周四", "13:30—16:30 参访；离场前完成服务卡", "服务卡", "全体陪同", "", "待启动"],
        ["2026-10-23", "周五", "次日单独问感受，记录满意/一般/不满意", "回访记录", "客户经理", "", "待启动"],
        ["2026-10-27", "周二", "台账建档；不满意事项项目经理已知情", "台账", "项目办", "", "待启动"],
        ["2026-10-29", "周四", "新客户完成第一次联系", "联系记录", "客户经理", "", "待启动"],
    ],
    [14, 8, 48, 16, 16, 12, 12],
    note="10 月 1 日至 7 日为国庆假期，不安排客户活动。",
    status_col=7,
    row_height=28,
)

# ---------------------------------------------------------------------------
# 05 后续节点
# ---------------------------------------------------------------------------
ws = wb.create_sheet("05-后续节点")
style_sheet(
    ws,
    "第2至第7场关键节点｜邀请一律一对一，不公开海投",
    ["节点", "11/5 小局", "11/19 参访", "11/26 新客局", "12/3 参访", "12/10 答谢参访", "12/17 收官", "负责人"],
    [
        ["锁定场地/企业", "2026-10-15", "2026-10-30", "2026-11-05", "2026-11-12", "2026-11-20", "2026-11-26", "项目经理"],
        ["名单确认", "2026-10-22", "2026-11-05", "2026-11-12", "2026-11-19", "2026-11-26", "2026-12-03", "客户经理"],
        ["发出邀请", "2026-10-23", "2026-11-06", "2026-11-13", "2026-11-20", "2026-11-27", "2026-12-04", "客户经理"],
        ["到场确认", "2026-11-04", "2026-11-18", "2026-11-25", "2026-12-02", "2026-12-09", "2026-12-16", "客户经理"],
        ["活动日", "2026-11-05", "2026-11-19", "2026-11-26", "2026-12-03", "2026-12-10", "2026-12-17", "全体"],
        ["次日满意度", "2026-11-06", "2026-11-20", "2026-11-27", "2026-12-04", "2026-12-11", "2026-12-18", "客户经理"],
        ["新客第5个工作日", "2026-11-12", "2026-11-26", "2026-12-03", "2026-12-10", "2026-12-17", "2026-12-24", "客户经理"],
    ],
    [18, 14, 14, 14, 14, 16, 14, 12],
    note="11/19 与 11/26 间隔一周，两场名单不要重复。新客局若取消，把已邀请的新客户改约到 12/17，并在当天说明改期原因。",
    row_height=28,
)

# ---------------------------------------------------------------------------
# 06 即日行动
# ---------------------------------------------------------------------------
ws = wb.create_sheet("06-即日行动")
style_sheet(
    ws,
    "即日行动清单｜9月23日至10月22日",
    ["日期", "动作", "完成标准", "负责人", "完成"],
    [
        ["2026-09-23", "四方确认费用硬顶和 7 场安排", "纪要写明单场 ≤5000、12 月不减场", "", "待启动"],
        ["2026-09-24", "客户经理提交名单", "每人：存量重点 8 家、意向新客 5 家", "", "待启动"],
        ["2026-09-28", "参访企业储备", "3 家企业同意免费接待并接受拍摄限制", "", "待启动"],
        ["2026-09-30", "话术与礼品口径", "合规书面确认：默认无伴手礼、不承诺授信", "", "待启动"],
        ["2026-10-08", "第一场立项", "企业、线路、16 位客户、预算表同时定稿", "", "待启动"],
        ["2026-10-09", "发出邀请", "每一家都有客户经理邀约记录", "", "待启动"],
        ["2026-10-19", "预走访", "动线、茶水、交通报价都落在 5000 元内", "", "待启动"],
        ["2026-10-22", "举办第一场", "服务卡与当天决算草稿同时收回", "", "待启动"],
        ["2026-10-23", "满意度回访", "出席客户均有满意/一般/不满意记录", "", "待启动"],
    ],
    [14, 36, 42, 14, 12],
    note="负责人一栏在 9 月 23 日会上填姓名。没有姓名的事项不开始。",
    status_col=5,
)

# ---------------------------------------------------------------------------
# 07 客户出席
# ---------------------------------------------------------------------------
ws = wb.create_sheet("07-客户出席回访")
headers = ["场次", "日期", "客户经理", "企业名称", "联系人", "手机", "客户类型", "邀约状态", "是否出席", "满意度", "一句原因", "不满意回访", "备注"]
sample = ["10/22 参访", "2026-10-22", "（示例，请覆盖）", "某某科技", "", "", "存量", "待邀", "", "", "", "", "示例行"]
empty = [[""] * len(headers) for _ in range(80)]
style_sheet(
    ws,
    "客户出席与回访｜一场一行一位客户｜用来统计满意度和人次",
    headers,
    [sample] + empty,
    [14, 12, 14, 22, 12, 14, 12, 12, 12, 12, 28, 14, 18],
    note="客户类型：存量 / 新客 / 转介。转介指老客户带来的新企业，计入新客成果。不满意的，2 日内必须电话回访并让项目经理知情。",
    row_height=22,
)
add_list(ws, "A3:A200", '"10/22 参访,11/5 小局,11/19 参访,11/26 新客局,12/3 参访,12/10 答谢参访,12/17 收官"')
add_list(ws, "G3:G200", '"存量,新客,转介"')
add_list(ws, "H3:H200", '"待邀,已邀,确认到场,拒绝,候补"')
add_list(ws, "I3:I200", '"是,否"')
add_list(ws, "J3:J200", '"满意,一般,不满意"')
add_list(ws, "L3:L200", '"无需,已回访,未回访"')

# ---------------------------------------------------------------------------
# 08 新客跟进
# ---------------------------------------------------------------------------
ws = wb.create_sheet("08-新客跟进")
headers = ["场次", "企业名称", "来源", "陪同老客户", "客户经理", "出席", "离场对接人", "第5个工作日", "是否已联系", "联系日期", "客户原话要点", "下一步", "状态"]
sample = ["11/26 新客局", "（示例，请覆盖）", "转介", "", "", "", "", "", "", "", "", "5个工作日内拜访或电话", "待联系"]
empty = [[""] * len(headers) for _ in range(40)]
style_sheet(
    ws,
    "新客户跟进｜只记录新客和转介｜5 个工作日内必须有联系记录",
    headers,
    [sample] + empty,
    [14, 20, 10, 14, 12, 8, 12, 14, 12, 12, 28, 22, 12],
    note="不在本表记录授信结论。联系结果只写：已联系、约到下次、暂无需求、联系不上。",
    status_col=13,
    row_height=22,
)
add_list(ws, "A3:A80", '"10/22 参访,11/5 小局,11/19 参访,11/26 新客局,12/3 参访,12/10 答谢参访,12/17 收官"')
add_list(ws, "C3:C80", '"提名,转介,活动相识"')
add_list(ws, "F3:F80", '"是,否"')
add_list(ws, "I3:I80", '"是,否"')
add_list(ws, "M3:M80", '"待联系,已联系,约到下次,暂无需求,联系不上"')

# ---------------------------------------------------------------------------
# 09 服务事项
# ---------------------------------------------------------------------------
ws = wb.create_sheet("09-服务事项")
headers = ["事项编号", "场次", "企业名称", "客户类型", "客户要办的一件事", "客户经理", "承诺反馈日", "状态", "反馈内容", "是否涉及融资"]
sample = ["1022-001", "10/22 参访", "（示例）", "存量", "想了解订单周转可以准备哪些材料", "", "2026-11-05", "待反馈", "", "是"]
empty = [[""] * len(headers) for _ in range(40)]
style_sheet(
    ws,
    "服务事项｜每家客户最多先记一件事｜10 个工作日内给第一次反馈",
    headers,
    [sample] + empty,
    [12, 14, 18, 12, 36, 12, 14, 12, 28, 14],
    note="涉及融资的，只反馈材料准备或是否进入银行正常流程，不写额度、利率和放款时间。",
    row_height=24,
)
add_list(ws, "D3:D80", '"存量,新客,转介"')
add_list(ws, "H3:H80", '"待反馈,已反馈,转介,暂不适配,客户暂缓"')
add_list(ws, "J3:J80", '"是,否"')

# ---------------------------------------------------------------------------
# 10 KPI
# ---------------------------------------------------------------------------
ws = wb.create_sheet("10-指标")
ws.merge_cells("A1:F1")
ws["A1"] = "指标自动汇总｜出席与满意度来自「07-客户出席回访」，新客联系来自「08-新客跟进」"
ws["A1"].font = TITLE_FONT
ws["A1"].fill = LIGHT_FILL
ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 36
for j, h in enumerate(["指标", "目标", "当前数字", "怎么数", "是否达到", "说明"], 1):
    cell = ws.cell(2, j, h)
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
    cell.border = BORDER

kpis = [
    ("存量客户出席人次", 60, '=COUNTIFS(\'07-客户出席回访\'!G:G,"存量",\'07-客户出席回访\'!I:I,"是")', "类型=存量 且 出席=是"),
    ("新客户出席", 25, '=COUNTIFS(\'07-客户出席回访\'!G:G,"新客",\'07-客户出席回访\'!I:I,"是")+COUNTIFS(\'07-客户出席回访\'!G:G,"转介",\'07-客户出席回访\'!I:I,"是")', "新客+转介，且出席=是"),
    ("转介到场", 8, '=COUNTIFS(\'07-客户出席回访\'!G:G,"转介",\'07-客户出席回访\'!I:I,"是")', "类型=转介 且 出席=是"),
    ("已填满意度份数", 1, '=COUNTIF(\'07-客户出席回访\'!J:J,"满意")+COUNTIF(\'07-客户出席回访\'!J:J,"一般")+COUNTIF(\'07-客户出席回访\'!J:J,"不满意")', "有满意度记录的人数"),
    ("满意份数", 1, '=COUNTIF(\'07-客户出席回访\'!J:J,"满意")', "满意度=满意"),
    ("满意度", 0.9, '=IF(C6=0,"",C7/C6)', "满意份数 / 已填份数"),
    ("新客已联系", 1, '=COUNTIF(\'08-新客跟进\'!I:I,"是")', "新客表中已联系=是"),
    ("新客应联系", 1, '=COUNTIF(\'08-新客跟进\'!F:F,"是")', "新客表中出席=是"),
    ("新客联系率", 1, '=IF(C10=0,"",C9/C10)', "已联系 / 应联系"),
]
for i, (name, target, formula, how) in enumerate(kpis):
    r = 3 + i
    ws.cell(r, 1, name).font = BODY_FONT
    ws.cell(r, 2, target).font = BODY_FONT
    ws.cell(r, 3, formula).font = BODY_FONT
    ws.cell(r, 4, how).font = BODY_FONT
    if name == "满意度":
        ws.cell(r, 2, "≥90%")
        ws.cell(r, 5, '=IF(C8="","待填",IF(C8>=0.9,"达到","未达到"))')
        ws.cell(r, 3).number_format = "0%"
    elif name == "新客联系率":
        ws.cell(r, 2, "100%")
        ws.cell(r, 5, '=IF(C11="","待填",IF(C11>=1,"达到","未达到"))')
        ws.cell(r, 3).number_format = "0%"
    elif name in ("已填满意度份数", "满意份数", "新客已联系", "新客应联系"):
        ws.cell(r, 2, "—")
        ws.cell(r, 5, "—")
    else:
        ws.cell(r, 5, f'=IF(C{r}="","待填",IF(C{r}>=B{r},"达到","未达到"))')
    ws.cell(r, 6, "").font = BODY_FONT
    for col in range(1, 7):
        ws.cell(r, col).border = BORDER
        ws.cell(r, col).alignment = WRAP
    ws.row_dimensions[r].height = 28

ws.conditional_formatting.add("E3:E11", FormulaRule(formula=['E3="达到"'], fill=OK_FILL))
ws.conditional_formatting.add("E3:E11", FormulaRule(formula=['E3="未达到"'], fill=RED_FILL))

ws.cell(12, 1, "满意度目标看比例，不看绝对份数。新客 5 日内联系率 = 已联系 / 应联系，目标 100%。示例行请覆盖，否则会被算进数字。")
ws["A12"].font = Font(name=FONT, size=10, italic=True, color=MUTED)
ws.merge_cells("A12:F12")
ws.row_dimensions[12].height = 32
for i, w in enumerate([22, 12, 18, 36, 12, 28], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A3"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------------------
# 11 责任
# ---------------------------------------------------------------------------
ws = wb.create_sheet("11-责任与人力")
style_sheet(
    ws,
    "责任与小场人力｜先填姓名，没有姓名的岗位视为未到位",
    ["事项或岗位", "人数", "拍板人", "办理人", "姓名", "备份", "一场要做完的事"],
    [
        ["项目经理", "1", "项目经理", "项目经理", "", "", "控预算、控时间、处理不满意升级"],
        ["客户经理陪同", "4—8", "客户经理", "客户经理", "", "", "邀请、陪同、次日回访、新客首联"],
        ["会务", "1", "项目经理", "会务", "", "", "场地、交通、茶水、物料，报价不超 5000"],
        ["记录", "1", "项目经理", "记录", "", "", "服务卡收回，当天录入台账"],
        ["合规窗口", "1", "银行合规", "银行合规", "", "", "审邀请话术；不必全程在场"],
        ["被访企业引导", "1", "被访企业", "被访企业", "", "", "路线、安全、无拍摄区"],
        ["参访企业落实", "—", "项目经理", "科企联/项目办", "", "", "书面确认免费接待和拍摄边界"],
        ["预算不超 5000", "—", "项目经理", "会务", "", "", "决算录入闸门表"],
    ],
    [18, 10, 14, 16, 14, 12, 42],
    note="小局可以不设被访企业引导。客户经理人数按自己客户的数量增减，但不把客户交给不认识的人陪同。",
    row_height=30,
)

# ---------------------------------------------------------------------------
# 12 检查表
# ---------------------------------------------------------------------------
ws = wb.create_sheet("12-现场检查")
style_sheet(
    ws,
    "现场检查｜每场复制使用｜场次：　　日期：　　项目经理：",
    ["时点", "检查项", "适用", "完成", "姓名"],
    [
        ["立项", "预算表合计 ≤5000，且无付费嘉宾、无摄影团队", "全部", "否", ""],
        ["立项", "被访企业书面同意免费接待，写明无拍摄区", "参访", "否", ""],
        ["邀请前", "话术不含额度、利率、放款、领导保证出席", "全部", "否", ""],
        ["T-5", "交通和茶歇报价复核，超了已减项", "全部", "否", ""],
        ["T-1", "每位客户都有陪同的客户经理", "全部", "否", ""],
        ["当天开始", "13:30 开始；告知不收费、不承诺贷款", "全部", "否", ""],
        ["参访中", "动线 90 分钟内，引导在场", "参访", "否", ""],
        ["离场前", "每家客户有一张服务卡，客户点头", "全部", "否", ""],
        ["离场前", "新客户知道对接人姓名", "有新客时", "否", ""],
        ["当天", "票据装袋，准备录入决算闸门", "全部", "否", ""],
        ["次日", "出席客户都有满意度", "全部", "否", ""],
        ["2日内", "不满意客户已电话，项目经理已知情", "如有", "否", ""],
    ],
    [12, 48, 12, 10, 14],
    note="完成栏改为“是”才算过。参访场次中标成“参访”的项目不能空。",
    row_height=26,
)
add_list(ws, "D3:D30", '"是,否"')

# ---------------------------------------------------------------------------
# 13 风险
# ---------------------------------------------------------------------------
ws = wb.create_sheet("13-风险与边界")
style_sheet(
    ws,
    "风险与边界｜开场先说不收费、不承诺贷款",
    ["风险", "出现时怎么处理", "谁处理", "程度"],
    [
        ["报价超过 5000 元", "减交通、茶歇或物料；减完仍超过就换零成本场地", "项目经理", "高"],
        ["被访企业临时取消", "改到银行贵宾室，同主题小局，已确认客户照常接待", "项目经理", "高"],
        ["客户问能批多少", "回答由银行独立审批；不给区间、不给时间", "客户经理", "高"],
        ["有人现场拉业务或收费", "当场制止并请离，取消后续资格", "项目经理", "高"],
        ["客户不满意", "2 日内客户经理电话；项目经理记录原因", "客户经理", "中"],
        ["新客 5 日内没人联系", "项目经理当天改派另一位客户经理", "项目经理", "高"],
        ["同一客户被连续邀请", "相邻场次默认不重复，季度不超过 3 次", "项目经理", "中"],
        ["想加伴手礼", "先过合规；总额仍须 ≤5000，否则不做", "银行合规", "中"],
        ["拍摄到不该拍的区域", "当场删除，书面说明", "会务", "高"],
    ],
    [24, 48, 14, 10],
    note="金融业务由泰隆银行独立尽调、审批。活动方不代替银行作授信判断。",
    status_col=4,
    row_height=30,
)

# ---------------------------------------------------------------------------
# 14 话术
# ---------------------------------------------------------------------------
ws = wb.create_sheet("14-话术")
style_sheet(
    ws,
    "可直接使用的话术｜方括号处替换成当天信息",
    ["场景", "怎么说", "不要补上的话"],
    [
        ["邀请老客户", "【姓名】总，10月22日下午我们安排了一场小范围企业参访，半天，我陪您一起去。人不超二十位，不收费。您看这天是否方便？", "不要说保证见到某位领导，不要说去了有额度。"],
        ["邀请新客户", "我们和复旦、科企联一起，请少数企业去现场看一家公司，银行客户经理会陪同。不收费，也不在现场谈贷款结果。如果您愿意来，我负责接您，活动后 5 个工作日内再跟您联系一次。", "不要说内部指标、隐形名额、送礼开户。"],
        ["老客户带新", "如果您有一位合适的同行或合作伙伴，可以一起来。名额有限，我先替您留一个位子。", "不要承诺对方一定能成为我行客户。"],
        ["开场", "今天下午只做一件事：把大家带来的问题记清楚。活动不收费。贷款能否办理，由泰隆银行按自己的流程审批，现场不作承诺。", "不要介绍具体产品价格和通过率。"],
        ["离场", "今天我为您记下的是【一件事】。下一步由我来办，最晚【日期】前给您回复。您看是不是这一件？", "不要当场加第二件、第三件把客户留住。"],
        ["次日回访", "昨天的参访，哪一点对您有用，哪一点浪费了您的时间？我想听一句实在的。", "不要群发同一段话。"],
        ["不满意", "昨天那件事没安排好，是我们的问题。我今天先处理【具体动作】，处理完再告诉您。", "不要解释成客户误会了。"],
        ["新客首次联系", "谢谢您昨天来。我是您的对接人【姓名】。您当时提到【一件事】，我这周把可以准备的资料或可介绍的同事跟您说清楚。是否办理业务，还是要走银行正常流程。", "不要在电话里做审批判断。"],
        ["暂不适合", "以目前这件需求，银行这条线暂时对不上，原因是【客观原因】。我可以介绍【另一条合规渠道】，是否使用由您决定。您仍在我们的客户名单里，下一场合适的活动我再单独邀您。", "不要说再等等来拖延，不要留下负面评价。"],
    ],
    [16, 62, 36],
    note="对外提到泰隆银行名称、标识和案例前，先经银行指定窗口看过。",
    row_height=48,
)

# ---------------------------------------------------------------------------
# 打印与页眉
# ---------------------------------------------------------------------------
for ws in wb.worksheets:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.oddHeader.center.text = "泰隆银行 · 2026年四季度客户参访"
    ws.oddFooter.left.text = "单场不超过5000元"
    ws.oddFooter.right.text = "第 &P 页"
    ws.page_setup.orientation = "landscape"
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5
    ws.sheet_properties.tabColor = "0A1E36"

wb["03-决算闸门"].sheet_properties.tabColor = "C98A2A"
wb["07-客户出席回访"].sheet_properties.tabColor = "1F7A4D"
wb["10-指标"].sheet_properties.tabColor = "1F7A4D"

out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "deliverables")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "泰隆银行客户参访与高端小局_2026年10-12月_执行表.xlsx")
wb.save(out_path)
print(f"已生成：{out_path}（共 {len(wb.sheetnames)} 个工作表）")
for name in wb.sheetnames:
    print(f"  - {name}")
