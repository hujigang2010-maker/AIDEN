#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 2026 年 10 月 31 日王德峰讲座的每人收费评估（内部 Word）。"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"

NAVY = RGBColor(0x1A, 0x3A, 0x5C)
NAVY_HEX = "1A3A5C"
GOLD = RGBColor(0x8B, 0x69, 0x14)
INK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0x8B, 0x1E, 0x1E)
RED_HEX = "8B1E1E"
CREAM_HEX = "F7F3EA"
SOFT_HEX = "EEF3F8"
AMBER_HEX = "F8F1E3"
GREEN_HEX = "E8F2EA"


def set_run_font(run, *, name="微软雅黑", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_table_widths(table, widths_cm):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths_cm[idx] * 567)))
            tc_w.set(qn("w:type"), "dxa")
            v_align = OxmlElement("w:vAlign")
            v_align.set(qn("w:val"), "center")
            tc_pr.append(v_align)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=9, color=MUTED)


def setup_section(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run("内部评估  ·  不直接发给学员、赞助商或合作方")
    set_run_font(run, size=9, color=RED, bold=True)
    p_pr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), RED_HEX)
    pbdr.append(bottom)
    p_pr.append(pbdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = fp.add_run("10月31日讲座每人收费评估  ·  第 ")
    set_run_font(left, size=9, color=MUTED)
    add_page_number(fp)
    right = fp.add_run(" 页")
    set_run_font(right, size=9, color=MUTED)


def add_para(
    doc,
    text,
    *,
    size=11,
    bold=False,
    color=INK,
    align=None,
    space_before=0,
    space_after=6,
    first_line=False,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.25
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_line:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_h1(doc, text):
    return add_para(doc, text, size=14, bold=True, color=NAVY, space_before=12, space_after=6)


def add_h2(doc, text):
    return add_para(doc, text, size=12, bold=True, color=GOLD, space_before=8, space_after=4)


def set_cell_text(cell, text, *, bold=False, size=9.5, color=INK, fill=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True,
            size=9,
            color=WHITE,
            fill=NAVY_HEX,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for r, row in enumerate(rows):
        fill = "F5F7FA" if r % 2 else "FFFFFF"
        for c, value in enumerate(row):
            set_cell_text(table.rows[r + 1].cells[c], value, size=9, fill=fill)
    set_table_widths(table, widths)
    for row in table.rows:
        prevent_row_split(row)
    add_para(doc, "", size=4, space_after=2)
    return table


def build():
    doc = Document()
    setup_section(doc)

    add_para(
        doc,
        "2026年10月31日讲座",
        size=12,
        bold=True,
        color=GOLD,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    add_para(
        doc,
        "每人收费评估",
        size=22,
        bold=True,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=2,
    )
    add_para(
        doc,
        "《传习录》与阳明心学  ·  一场上午专题讲座",
        size=11,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    add_para(
        doc,
        "评估日期：2026年10月10日",
        size=10,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )

    add_para(doc, "结论", size=14, bold=True, color=NAVY, space_before=4, space_after=4)
    add_para(
        doc,
        "公开报名按每人 499 元收取。海报、公众号和朋友圈只出现这一个价格。",
        size=12,
        bold=True,
        color=RED,
        space_after=6,
    )
    add_para(
        doc,
        "本场实际交付是 2026 年 10 月 31 日上午 09:00—12:00、约 3 小时的线下讲座。听众买到的是一次入场听讲。499 元对应该产品。",
        first_line=True,
    )
    add_para(
        doc,
        "此前内部测算使用过“12 人 × 2,399 元”。该口径接近全日小班的单日锚点，当时也写明不是已到账数字。本场议程只有上午一场，公开售价改按讲座收取，定为 499 元。2,399 元只保留给已经按这个价格作出书面承诺的人，并单独写清他们多出来的权益。",
        first_line=True,
    )

    add_h1(doc, "一、按实际交付定价")
    add_para(
        doc,
        "海报精简稿里的“《传习录》与阳明心学（上）”和“（下）”写在同一个上午时段里，是一场讲座的两段内容。活动信息也写明具体授课时间以最终议程为准，目前能够对外销售的只有这一场。",
        first_line=True,
    )
    add_para(
        doc,
        "因此对外只卖一种票：标准席 499 元，权益是上午讲座入场。报名页在赠书、资料、问答得到书面确认之前，不增加前区席或前排席，也不单设“半天专题席位”。议程本身就是半天里的一场讲座，再拆出一个更低的半天价，会让 499 元看起来被拆售。",
        first_line=True,
    )
    add_table(
        doc,
        ["项目", "本场可以对外说的事实"],
        [
            ["日期与时段", "2026年10月31日（星期六）09:00—12:00"],
            ["时长", "约 3 小时，一场"],
            ["地点", "上海，具体地址待公布"],
            ["主讲", "王德峰教授"],
            ["已确认权益", "现场听讲"],
            ["尚未确认、暂不写入票价", "签名书、学习资料、专属交流、优先报名"],
            ["规划容量", "沿用内部讨论的 200 人，场地未定"],
        ],
        [4.2, 12.2],
    )

    add_h1(doc, "二、499 元怎么来的")
    add_h2(doc, "2.1 对照城镇居民的月收入")
    add_para(
        doc,
        "国家统计局 2026 年 7 月 15 日发布：2026 年上半年城镇居民人均可支配收入中位数 26,389 元。按 6 个月折算，月均约 4,398 元。来源：国家统计局《2026年上半年居民收入和消费支出情况》。",
        first_line=True,
    )
    add_table(
        doc,
        ["票价", "约占月均可支配收入中位数", "适合的卖法"],
        [
            ["499 元", "约 11.3%", "公开讲座的标准价，可印在海报上"],
            ["699 元", "约 15.9%", "前区座位；须先确认资料或赠书"],
            ["999 元", "约 22.7%", "前排交流；须先获得教授授权"],
            ["2,399 元", "约 54.5%", "已有书面承诺的历史报价，不作为公开价"],
            ["5,000 元", "约 1.1 个月", "第三方页面上的全日小班标价，不是本场价格"],
        ],
        [3.2, 5.6, 7.6],
    )
    add_para(
        doc,
        "499 元需要听众认真决定，仍然落在一次文化活动的决策范围内。2,399 元大约相当于城镇居民半个多月的可支配收入中位数，决策周期更长，和目前这套朋友圈、公众号、海报扫码的传播方式不匹配。",
        first_line=True,
    )

    add_h2(doc, "2.2 和市面上同类课程分开")
    add_para(
        doc,
        "第三方招生页面上，与王德峰相关的课程有另一套标价：全日小班可见 5,000 元/人·1 天；《传习录》系统课可见 19,800 元/人、8 个月、每月 1 天，折合每个授课日约 2,475 元。这些页面不是本场主办方或教授本人的报价，产品里通常还有小班、教材、茶歇和连续学制。本场公开卖的是 3 小时听讲，使用讲座价格。",
        first_line=True,
    )
    add_para(
        doc,
        "2,399 元与“每个授课日约 2,475 元”非常接近，是把系统课的单日锚点拿来用了。本场没有 8 个月学制，也没有写成全日，公开价不沿用这个锚点。",
        first_line=True,
    )

    add_h2(doc, "2.3 距离活动日还有 21 天")
    add_para(
        doc,
        "评估日是 10 月 10 日，活动日是 10 月 31 日。场地和报名二维码确定后，真正能卖票的天数还要再短一些。499 元可以在机构社群和朋友圈里当天做决定。2,399 元更依赖一对一邀约，三周内难以靠公开海报铺开。",
        first_line=True,
    )
    add_table(
        doc,
        ["卖法", "人数假设", "估算实收", "读取方式"],
        [
            ["公开讲座", "150 人 × 499 元", "约 7.04 万元", "三家机构渠道有机会达到"],
            ["高价邀约", "30 人 × 2,399 元", "约 6.77 万元", "要在三周内完成 30 个高价承诺"],
        ],
        [3.2, 4.4, 3.4, 5.4],
    )
    add_para(
        doc,
        "实收按票面的 94% 估算，留出支付手续费、少量退款和票务成本。150 × 499 × 94% = 70,359 元；30 × 2,399 × 94% = 67,652 元。公开讲座用更多听众走到相近甚至更高的票款，也和已经写好的大众传播文案一致。",
        size=10,
        color=MUTED,
    )

    add_h1(doc, "三、对外就这一个价")
    add_table(
        doc,
        ["名称", "价格", "哪里出现", "包含什么"],
        [
            ["标准席", "499 元/人", "海报、公众号、朋友圈、报名页", "10月31日上午讲座入场"],
        ],
        [3.2, 3.0, 5.6, 4.6],
    )
    add_para(
        doc,
        "主海报不印早鸟价，也不印多个席位。开售后如果 72 小时未售出 40 张，再内部决定是否限量放出 399 元；放出前改报名页，不回头改已经发出的主海报。",
        first_line=True,
    )
    add_para(
        doc,
        "赠书、学习资料或现场交流一旦拿到书面确认，可以加开升级席，价格使用 699 元（前区，含已确认资料）和 999 元（前排，含已授权交流）。未确认前，这两档不出现在任何对外材料里。签名书尤其如此：原海报提到过，新版必须等首乾书院及相关权利方确认。",
        first_line=True,
    )
    add_para(
        doc,
        "同一机构购买多张，仍按 499 元/人。需要品牌露出的企业，走已经讨论过的赞助方案，赞助席位包含入场，不与门票重复计收入。新的赞助表述使用“含讲座入场”，不再写“与 2,399 元学员权益相同”。",
        first_line=True,
    )

    add_h1(doc, "四、499 元能不能覆盖成本")
    add_para(
        doc,
        "下面按保守客单估算：全部按 499 元出售，实收按 94% 计算，人均净贡献约 469 元。为了留余地，保本人数按人均 450 元测算。容量 200 人来自此前内部讨论，场地未定。师资 10 万元、赞助 6 万至 8 万元同样是内部口径，不是已签约或已到账数字。",
        first_line=True,
    )
    add_table(
        doc,
        ["成本情景", "现金成本", "保本人数", "对 200 座的含义"],
        [
            ["师资按分成，场地与执行 5 万元", "5 万元", "112 人", "售出 112 人可覆盖现金成本"],
            ["师资现金 6 万元，场地与执行 4 万元", "10 万元", "223 人", "坐满 200 座仍差一截"],
            ["师资现金 10 万元，场地与执行 5 万元", "15 万元", "334 人", "门票单独覆盖不了"],
        ],
        [6.4, 2.6, 2.4, 4.0],
    )
    add_para(
        doc,
        "200 座售出七成，即 140 张票，按人均 450 元计，净票款 6.3 万元。再计入 6 万元赞助，合计 12.3 万元：高于中等情景的 10 万元，低于重情景的 15 万元。赞助席和工作人员席还要占座，可售票少于 200 张，实际票款会低于 6.3 万元。",
        first_line=True,
    )
    add_para(
        doc,
        "499 元是听众侧的合理价格。活动要打平，靠的是把师资现金支出谈成分成或压到与座位匹配的水平，并落实赞助。票价维持 499 元，成本按场次规模来配。",
        first_line=True,
        bold=True,
    )

    add_h1(doc, "五、已经报过 2,399 元的人")
    add_para(
        doc,
        "内部材料里的 12 人、每人 2,399 元是测算口径，不是本评估确认的到账名单。处理规则按承诺状态区分。",
        first_line=True,
    )
    add_table(
        doc,
        ["状态", "收费", "权益写法"],
        [
            ["尚无报价、尚未收款", "499 元", "上午讲座入场"],
            ["口头提过 2,399 元，没有书面承诺，也未收款", "改为 499 元", "按标准席重新确认"],
            ["已有书面承诺或已收款", "维持 2,399 元", "入场，加上当时写明的前排、资料、交流或后续优先"],
        ],
        [5.6, 3.2, 7.6],
    )
    add_para(
        doc,
        "已按 2,399 元承诺的权益，只兑现当时写明的项目。新海报不把未确认的签名书、交流机会追加进这一档，也不把这一档写成和 499 元完全相同。两档差价对应多出来的、已经答应的服务。",
        first_line=True,
    )

    add_h1(doc, "六、可以直接贴进物料的句子")
    add_h2(doc, "海报票务区")
    add_para(doc, "标准席  499元/人", size=12, bold=True, color=NAVY, space_after=2)
    add_para(doc, "10月31日上午  09:00—12:00", space_after=2)
    add_para(doc, "扫码报名", space_after=2)
    add_para(
        doc,
        "课程安排及退费规则，以报名页面为准。",
        size=10,
        color=MUTED,
    )
    add_h2(doc, "报名页补充")
    add_para(
        doc,
        "标准席 499 元/人，含 2026 年 10 月 31 日上午《传习录》与阳明心学专题讲座入场。建议退费规则：10 月 24 日（含）前可全额退款；10 月 25 日至 29 日退还票价的 50%；10 月 30 日起可以转让，不退款。正式规则以报名页公布的文字为准。",
        first_line=True,
    )
    add_h2(doc, "朋友圈文案末尾加一句")
    add_para(doc, "标准席 499 元/人，活动报名及详情请见海报二维码。", bold=True)

    add_h1(doc, "七、开售前核对")
    add_para(doc, "1. 三家联合主办机构同意名称和标识用于这场收费活动。", space_after=3)
    add_para(doc, "2. 王德峰教授已授权使用姓名、肖像和本场课程信息。", space_after=3)
    add_para(
        doc,
        "3. 场地容量和师资付款方式写入同一张预算表。师资若按约 10 万元现金支付，同时赞助尚未落实，先把可售座位和预算对齐，再对外放票。",
        space_after=3,
    )
    add_para(
        doc,
        "4. 海报正文保持一场讲座。精简稿里的“（上）”“（下）”并成一段课程介绍，避免读者理解成两场课程、一张门票。",
        space_after=8,
    )
    add_para(
        doc,
        "本评估用于内部定价。对外票价、权益和退费，以报名页及主办方正式通知为准。",
        size=10,
        color=MUTED,
    )

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "10月31日讲座_每人收费评估.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build())
